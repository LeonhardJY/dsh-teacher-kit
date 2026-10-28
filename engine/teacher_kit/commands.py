"""命令表：CLI 与测试都从这里进入，避免两套入口漂移。"""

import copy
import os
import re

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .errors import EngineError, require
from .docx_build import build_docx
from .docx_inspect import inspect_docx
from .paper import MARGINS, PAPERS
from .pptx_build import SLIDE_SIZES, build_pptx
from .pptx_inspect import inspect_pptx
from .placeholders import replace_text
from .preflight import preflight
from .template_model import build_profile, diff_profiles, page_from_profile, templatize
from .exam import build_exam, build_exam_pair, build_exam_teacher, exam_parse
from .lesson import build_deck, plan_deck
from .lesson_html import build_deck_html
from .handout import build_dictation
from .audio_script import build_script_audio
from .dictation_audio import build_dictation_audio
from .question_bank import bank_build, bank_inspect
from .plan import build_plan, parse_plan
from .files import docx_clean, docx_extract, docx_merge
from . import profile as profile_store
from .docflow import docflow_build
from .wechat import wechat_check, wechat_render
from .paths import RE_BAD


def _path(payload, key="path"):
    value = payload.get(key)
    require(bool(value), "path-required", "没有给出文件路径", "点「浏览…」选一份文件，或者把完整路径粘进输入框")
    full = os.path.abspath(value)
    require(os.path.isfile(full), "file-not-found", "找不到这个文件：" + full, "确认文件还在原位置；从 U 盘/微信里拷到电脑上再选")
    return full


def selfcheck(_payload):
    """体检：库在不在、版本多少。宿主 status 页面与用户排障都用它。"""
    report = {"python": ".".join(str(v) for v in __import__("sys").version_info[:3])}
    for module_name, key in (("docx", "python-docx"), ("pptx", "python-pptx"), ("openpyxl", "openpyxl"), ("PIL", "Pillow")):
        try:
            module = __import__(module_name)
            report[key] = getattr(module, "__version__", "ok")
        except Exception as error:
            report[key] = None
            report.setdefault("missing", []).append({"module": key, "error": str(error)})
    report["ok"] = "missing" not in report
    return report


def inspect(payload):
    """按扩展名分发：docx / pptx 各给一份版式事实。"""
    full = _path(payload)
    lower = full.lower()
    if lower.endswith(".docx") or lower.endswith(".docm"):
        return inspect_docx(full, preview=int(payload.get("preview", 12)))
    if lower.endswith(".pptx") or lower.endswith(".pptm"):
        return inspect_pptx(full, preview=int(payload.get("preview", 6)))
    raise EngineError(
        "unsupported-format",
        "这个格式还不支持检查：" + os.path.splitext(lower)[1],
        "目前支持 .docx / .docm / .pptx / .pptm",
    )


def docx_new(payload):
    """从结构化内容造一个 docx（可以套学校模板）。"""
    output = payload.get("outputPath")
    require(bool(output), "output-required", "没有指定输出文件路径", "在面板上选一个保存位置，或者把 outputPath 写成完整路径")
    return build_docx(payload)


def pptx_new(payload):
    """从"一节课的若干页"造一个 pptx（默认 16:9）。"""
    output = payload.get("outputPath")
    require(bool(output), "output-required", "没有指定输出文件路径", "在面板上选一个保存位置，或者把 outputPath 写成完整路径")
    return build_pptx(payload)


def papers(_payload):
    """给面板下拉框用：教师看得懂的纸名，不是 pt 数字。"""
    return {
        "papers": [
            {
                "size": name,
                "widthMm": size[0],
                "heightMm": size[1],
                "defaultMarginsMm": MARGINS.get(name),
            }
            for name, size in PAPERS.items()
        ],
        "defaultPaper": "A4",
        "slideSizes": [{"name": name, "widthMm": size[0], "heightMm": size[1]} for name, size in SLIDE_SIZES.items()],
        "defaultSlideSize": "16:9",
    }


def template_model(payload):
    """从老师自己的 1—3 份文件反推版式规则；多份不一致会列成冲突。"""
    paths = payload.get("paths")
    if isinstance(paths, str):
        paths = [paths]
    require(bool(paths), "paths-required", "请先选 1—3 份你自己的文件（上学期的成品，或学校的空白模板）", "把文件从「文件管理器」拖进面板的虚线框，或者点它选文件")
    checked = []
    for path in paths:
        full = os.path.abspath(path)
        require(os.path.isfile(full), "file-not-found", "找不到这个文件：" + full, "确认文件还在原位置；从 U 盘/微信里拷到电脑上再选")
        checked.append(full)
    result = build_profile(checked, payload.get("options"))
    profile = result["profile"]
    profile["buildPage"] = page_from_profile(profile)
    return profile


def template_diff(payload):
    """两份 profile 的差异：这学期和上学期，版式到底哪里变了。"""
    require(bool(payload.get("before")) and bool(payload.get("after")), "profiles-required", "需要 before 与 after 两份 profile")
    return diff_profiles(payload["before"], payload["after"])


def _sectpr_list(document):
    """按文档顺序列出所有 sectPr：段落 pPr 里的那些 + body 末尾那个。

    多节文档里，靠前的节（标题单栏节、正文两栏节）的版面存在**段落**的 pPr/sectPr 里，
    只有最后一节存在 body 末尾。清正文时只留 body 末尾那个，就会把 8 开两栏的版面一起删掉。
    """
    body = document.element.body
    found = []
    for child in body:
        if child.tag == qn("w:p"):
            ppr = child.find(qn("w:pPr"))
            if ppr is not None:
                sect = ppr.find(qn("w:sectPr"))
                if sect is not None:
                    found.append(sect)
    final = body.find(qn("w:sectPr"))
    if final is not None:
        found.append(final)
    return found


def _twips_to_mm(value):
    return round(int(value) * 25.4 / 1440.0, 1)


def _sectpr_paper(sect_pr):
    """从一个 sectPr 里读出纸型、方向、栏数、页边距（内部一律毫米）。"""
    facts = {"size": None, "widthMm": None, "heightMm": None, "orientation": None,
             "columns": 1, "marginMm": {}}
    if sect_pr is None:
        return facts
    pgsz = sect_pr.find(qn("w:pgSz"))
    if pgsz is not None:
        width, height = pgsz.get(qn("w:w")), pgsz.get(qn("w:h"))
        if width and height:
            width_mm, height_mm = _twips_to_mm(width), _twips_to_mm(height)
            facts["widthMm"], facts["heightMm"] = width_mm, height_mm
            facts["orientation"] = "landscape" if width_mm > height_mm else "portrait"
            long_side, short_side = max(width_mm, height_mm), min(width_mm, height_mm)
            for name, (paper_w, paper_h) in PAPERS.items():
                if abs(max(paper_w, paper_h) - long_side) <= 2 and abs(min(paper_w, paper_h) - short_side) <= 2:
                    facts["size"] = name
                    break
    cols = sect_pr.find(qn("w:cols"))
    if cols is not None:
        number = cols.get(qn("w:num"))
        facts["columns"] = int(number) if number else 1
    pgmar = sect_pr.find(qn("w:pgMar"))
    if pgmar is not None:
        for key, attr in (("top", "w:top"), ("bottom", "w:bottom"), ("left", "w:left"), ("right", "w:right")):
            value = pgmar.get(qn(attr))
            if value:
                facts["marginMm"][key] = _twips_to_mm(value)
    return facts


def _pick_body_sectpr(document, profile):
    """挑出"题目那一节"的版面：纸型最接近、栏数最对得上的那节。"""
    sect_prs = _sectpr_list(document)
    if not sect_prs:
        return None
    paper = profile.get("paper") or {}
    width = float(paper.get("widthMm") or 0)
    height = float(paper.get("heightMm") or 0)
    columns = int(paper.get("columns") or 1)
    best, best_score = None, None
    for sect_pr in sect_prs:
        facts = _sectpr_paper(sect_pr)
        score = abs(float(facts["widthMm"] or 0) - width) + abs(float(facts["heightMm"] or 0) - height)
        if columns > 1 and int(facts["columns"] or 1) != columns:
            score += 50
        if best_score is None or score < best_score:
            best, best_score = sect_pr, score
    return best


def _install_sectpr(final, source):
    """把 final 这个 sectPr 的内容换成 source 的（不带 w:type，也不带起始页码）。"""
    if final is None or source is None:
        return False
    # 先深拷贝：单节文档里 final 与 source 是同一个元素，先清后拷会把版面清空
    keep = [copy.deepcopy(child) for child in source
            if child.tag not in (qn("w:type"), qn("w:pgNumType"))]
    for child in list(final):
        final.remove(child)
    for child in keep:
        final.append(child)
    return True


RE_FOOTER_CLASS = re.compile(r"(使用\s*班级\s*[:：]\s*)([^\s　]+)")
RE_FOOTER_TEACHER = re.compile(r"((?:出卷|命题|制卷)(?:\s*老师|\s*教师)?\s*[:：]\s*)([^\s　]+)")


def _footer_placeholders(text):
    """页脚里的"使用班级：七(3)班 / 出卷老师：王老师"换成 {{className}} / {{teacherName}}。

    否则它们会被当成无名下划线换成 {{blank}}，老师在面板上填的班级和姓名就印不上去
    （模板套用时页脚永远是横线）。
    """
    text = RE_FOOTER_CLASS.sub(lambda m: m.group(1) + "{{className}}", text)
    text = RE_FOOTER_TEACHER.sub(lambda m: m.group(1) + "{{teacherName}}", text)
    return text


RE_BARE_CLASS = re.compile(
    r"(?:[一二三四五六七八九]|初[一二三]|高[一二三])\s*年级?\s*[（(]?\s*(?:\d{1,2}|[一二三四五六七八九十]{1,3})\s*[)）]?\s*班"
)


def _header_placeholders(text):
    """页眉里的「XX中学七年三班 英语」把班级抽成 {{className}}。

    页脚有「使用班级：」这种标签可以照着换，页眉往往是光秃秃一个班级名；
    不抽出来的话，下学期出新卷，页眉还印着上一届的「七年三班」。
    """
    value = _footer_placeholders(text)
    value = RE_BARE_CLASS.sub("{{className}}", value)
    return value


def _donor_sectpr(document, preferred):
    """找一份带页眉页脚引用（w:headerReference / w:footerReference）的节当捐赠者。

    真实卷子里只有第一节写引用，后面的节靠"链接到上一节"继承；把多节压成一节后
    继承链断了，模板就没有页眉页脚了（老师套用时页脚页码、学校抬头全没）。
    """
    candidates = [preferred] + _sectpr_list(document)
    for sect_pr in candidates:
        if sect_pr is None:
            continue
        if sect_pr.find(qn("w:headerReference")) is not None or sect_pr.find(qn("w:footerReference")) is not None:
            return sect_pr
    return None


def _adopt_header_footer_refs(final, donor):
    """final 自己没有页眉页脚引用时，把 donor 的引用元素抄过来（rId 在同一份包里有效）。"""
    if final is None or donor is None or final is donor:
        return 0
    copied = 0
    for tag in ("w:headerReference", "w:footerReference"):
        if final.find(qn(tag)) is not None:
            continue
        for reference in donor.findall(qn(tag)):
            final.append(copy.deepcopy(reference))
            copied += 1
    return copied


def _measured_paper(path):
    """保存后重新读回来实测纸型：返回的就是老师拿到手那份文件的真实版面。"""
    try:
        document = docx.Document(path)
    except Exception:
        return {}
    facts = [_sectpr_paper(sect_pr) for sect_pr in _sectpr_list(document)]
    if not facts:
        return {}
    # 同一张纸上栏数最多的那节就是正文节（标题节是单栏）
    return max(facts, key=lambda item: (int(item.get("columns") or 1), float(item.get("widthMm") or 0)))


def template_make(payload):
    """把老师手上"上学期那份卷子"变成下学期还能用的空白模板。

    读出版式规则 → 页眉页脚里的具体值改成占位符（学年/页码/班级/出卷人）
    → 正文留标题与说明行，题目位置空出来给老师粘。版式一点不动。
    """
    source = payload.get("source")
    require(bool(source), "source-required", "请先指定一份你自己的文件（上学期做好的那份）", "点「浏览…」选一份 .docx：越接近你要的成品，学出来的版式越准")
    full = os.path.abspath(source)
    require(os.path.isfile(full), "file-not-found", "找不到这个文件：" + full, "确认文件还在原位置；从 U 盘/微信里拷到电脑上再选")

    output = payload.get("outputPath")
    if not output:
        base = os.path.splitext(full)[0]
        output = base + "-模板.docx"
    output = os.path.abspath(output)
    if os.path.exists(output) and not payload.get("overwrite", False):
        raise EngineError("output-exists", "目标文件已存在，不覆盖：" + output, "换一个文件名，或显式允许覆盖")

    profile = template_model({"paths": [full]})
    profile["buildPage"] = page_from_profile(profile)

    document = docx.Document(full)
    # 模板本身不该带"连续分节"：套用时会和正文的两栏节并块，印出来变通栏。
    for section in document.sections:
        try:
            section.start_type = WD_SECTION.NEW_PAGE
        except Exception:
            pass
    # 先从原卷正文里把"满分/时长"读出来：模板里留的是占位符，但面板要给出像样的默认值
    source_texts = [(paragraph.text or "").strip() for paragraph in document.paragraphs]
    source_total = _first_number(source_texts, r"满分\s*(\d{1,4}(?:\.\d+)?)\s*分")
    source_minutes = _first_number(source_texts, r"时间\s*(\d{1,3})\s*分钟")

    replaced = 0
    for section in document.sections:
        for name in ("header", "first_page_header", "even_page_header",
                     "footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, name, None)
            if part is None:
                continue
            is_footer = "footer" in name
            for paragraph in part.paragraphs:
                original = paragraph.text or ""
                if not original.strip():
                    continue
                next_text = _footer_placeholders(original) if is_footer else _header_placeholders(original)
                next_text = templatize(next_text)
                if next_text != original:
                    replace_text(paragraph, next_text)
                    replaced += 1

    # 正文清空，只留标题与一行说明：题目由老师自己粘（v0.1 的边界）
    # 清之前先把"题目那一节"的版面记下来：8 开两栏/页边距/方向都存在段落 pPr 的 sectPr 里，
    # 跟着段落一起删，模板就退回 body 末尾那节（往往是 A4 答题卡）——套用时整份卷子变 A4 单栏。
    body_sectpr = _pick_body_sectpr(document, profile)
    donor = _donor_sectpr(document, body_sectpr)
    body = document.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)
    final_sectpr = body.find(qn("w:sectPr"))
    _install_sectpr(final_sectpr, body_sectpr)
    # 多节压成一节后"链接到上一节"的页眉页脚会断：把引用抄到自己身上
    _adopt_header_footer_refs(final_sectpr, donor)

    roles = profile.get("roles") or {}
    title_role = roles.get("title") or {}
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(payload.get("titlePlaceholder") or "{{examTitle}}")
    _style_like(run, title_role, default_size=16, default_bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = payload.get("notePlaceholder") or "（满分 {{totalScore}} 分，时间 {{minutes}} 分钟）"
    _style_like(subtitle.add_run(note), roles.get("question") or {}, default_size=11)

    parent = os.path.dirname(output)
    if parent and not os.path.isdir(parent):
        os.makedirs(parent, exist_ok=True)
    document.save(output)

    # 回报的纸型必须是**产出文件的实测值**，不是源文件的画像：两者不一致时老师会照着
    # 一个假的"8 开两栏"去套模板，印出来却是 A4 单栏。
    measured = _measured_paper(output)
    paper_out = dict(profile["paper"])
    if measured:
        for key in ("size", "widthMm", "heightMm", "orientation", "columns"):
            if measured.get(key):
                paper_out[key] = measured[key]
        if measured.get("marginMm"):
            paper_out["marginMm"] = dict(measured["marginMm"])
    warnings = []
    if measured and profile.get("paper"):
        expected = profile["paper"]
        drift = abs(float(measured.get("widthMm") or 0) - float(expected.get("widthMm") or 0))
        drift += abs(float(measured.get("heightMm") or 0) - float(expected.get("heightMm") or 0))
        if drift > 3:
            warnings.append({
                "code": "template-paper-drift",
                "message": "做出来的模板是 " + str(measured.get("widthMm")) + "×" + str(measured.get("heightMm"))
                           + "mm，和原卷的 " + str(expected.get("widthMm")) + "×" + str(expected.get("heightMm"))
                           + "mm 对不上；套用前请先用 Word 看一眼纸张大小",
            })

    return {
        "outputPath": output,
        "bytes": os.path.getsize(output),
        "source": full,
        "placeholderLines": replaced,
        "paper": paper_out,
        "warnings": warnings,
        "buildPage": page_from_profile({"paper": paper_out}),
        "confidence": profile["confidence"],
        "conflicts": profile["conflicts"],
        "values": {
            "examTitle": "七年级英语期中测试卷",
            "totalScore": source_total or "100",
            "minutes": source_minutes or "90",
            # 学年度是「2026—2027」这种年份区间，不是页眉原文：这里以前直接把
            # header.template（页眉那句"XX中学七年三班 英语"）塞进来，套出来新卷的
            # 标题行就印着上学期的班级。只有画像里真有像样的年份才用它。
            "schoolYear": _school_year_value(profile),
            "className": "{{blank}}",
        },
    }


def _school_year_value(profile):
    raw = (profile.get("header") or {}).get("schoolYear")
    if isinstance(raw, (list, tuple)):
        raw = " ".join(str(item) for item in raw)
    text = str(raw or "").strip()
    if re.search(r"20\d{2}", text):
        return text
    return "{{schoolYear}}"


def _first_number(texts, pattern):
    """在正文前 15 行里找第一个匹配的数字（满分 / 考试时长）。"""
    import re as _re
    regex = _re.compile(pattern)
    for text in texts[:15]:
        hit = regex.search(text)
        if hit:
            value = hit.group(1)
            return value[:-2] if value.endswith(".0") else value
    return None


def _style_like(run, role, default_size=12, default_bold=False):
    """按反推出来的角色样式给新文字上格式（模板没这段话，就按它的规矩来）。"""
    from docx.shared import Pt
    run.font.size = Pt(role.get("sizePt") or default_size)
    run.font.name = role.get("font") or "宋体"
    run.bold = role.get("bold") if role.get("bold") is not None else default_bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), run.font.name)


def profile_get(_payload=None):
    """读画像：家里这位老师是什么学校、什么学段学科、习惯什么版式。"""
    return profile_store.summary()


def profile_set(payload):
    """把这次对话里得知的画像并进去（learn=True 时顺带记一条学到了什么）。"""
    patch = payload.get("patch") or payload.get("profile") or {}
    require(isinstance(patch, dict), "patch-invalid", "patch 要给一份画像片段（JSON 对象）")
    saved = profile_store.update(patch, note=payload.get("note"), task=payload.get("task"),
                                 learn=bool(payload.get("learn", True)))
    result = profile_store.summary(saved["profile"])
    result["ignored"] = saved["ignored"]
    result["saved"] = saved["saved"]
    # 认得的节名但形状不对（比如 paper 直接给了字符串）会被洗掉，别让老师以为存上了
    dropped = [name for name in patch if name in profile_store.SECTIONS and name not in result["saved"]]
    if dropped and not result["saved"]:
        raise EngineError(
            "patch-invalid",
            "画像片段形状不对，这些节一个字都没存上：" + "、".join(dropped),
            "每一节要写成对象，例如 {school:{name:'临河三中',subject:'英语',grades:'七年级'},"
            " paper:{key:'8K-2'}}；别把字符串直接塞进节名")
    result["dropped"] = dropped
    result["path"] = saved["path"]
    return result


def profile_reset(_payload=None):
    """清空画像（只清记忆，不动任何已经生成的文件）。

    以前这里只 return summary()，于是「忘掉我学校的设置」回 ok:true 而画像一字未动，
    下一次 defaults.apply 又把旧学校、旧纸型兜底注入进来（真实踩过：换学校后卷面还是老的）。
    """
    profile_store.reset()
    result = profile_store.summary()
    result["cleared"] = True
    return result


def _raw_basename(path):
    return str(path or "").replace("\\", "/").rstrip("/").rsplit("/", 1)[-1]


def _guard_raw_paths(value, depth=0):
    """在 abspath 之前把关文件名。

    Windows 的 abspath 会把 "D:/x/a:b.docx" 里的 a: 当成盘符，算成 "A:\b.docx"——
    等调用方拿到 abspath 结果再查，半角冒号已经看不出来了（实测 makedirs 直接报
    盘符被算歪）。所以入口处先按原始字符串的 basename 检查一遍。
    """
    if depth > 6:
        return
    if isinstance(value, dict):
        for key, item in value.items():
            name = str(key).lower()
            if isinstance(item, str) and ("path" in name or name.endswith("dir")):
                base = _raw_basename(item)
                bad = RE_BAD.search(base) if base else None
                if bad:
                    raise EngineError(
                        "bad-output-path",
                        "文件名里有不能用的符号「" + bad.group(0) + "」：" + base,
                        "去掉 \\ / : * ? \" < > | 这些符号再存（用「-」或全角「：」都行）")
            else:
                _guard_raw_paths(item, depth + 1)
    elif isinstance(value, (list, tuple)):
        for item in value:
            _guard_raw_paths(item, depth + 1)


BAD_OPTION_KEYS = ("options", "page", "values", "properties")


def _guard_shapes(payload):
    """几个固定形状的字段先查一遍：options 写成字符串会让 .get 直接崩成英文 TypeError。"""
    for key in BAD_OPTION_KEYS:
        value = payload.get(key)
        if value is None or isinstance(value, dict):
            continue
        if key == "page" and isinstance(value, str):
            continue  # 纸型允许直接写 "A4"
        examples = {
            "options": '{"options": {"answerCard": "card", "minutes": 45}}',
            "page": '{"page": {"size": "8K", "orientation": "landscape"}}',
            "values": '{"values": {"className": "七年三班", "teacherName": "李老师"}}',
            "properties": '{"properties": {"title": "七年级英语阶段测试卷"}}',
        }
        raise EngineError(
            "bad-" + key,
            "「" + key + "」要写成对象（{...}），这次收到的是 " + type(value).__name__,
            "例如 " + examples.get(key, '{"' + key + '": {}}')
            + "；不确定就整段删掉，我按学校画像的默认值来")


def _guarded(function):
    def wrapper(payload=None, *args, **kwargs):
        if isinstance(payload, dict):
            _guard_raw_paths(payload)
            _guard_shapes(payload)
        return function(payload, *args, **kwargs)

    wrapper.__name__ = getattr(function, "__name__", "command")
    wrapper.__doc__ = getattr(function, "__doc__", None)
    return wrapper


TABLE = {
    "profile:get": profile_get,
    "profile:set": profile_set,
    "profile:reset": profile_reset,
    "wechat:check": wechat_check,
    "wechat:render": wechat_render,
    "docflow:build": docflow_build,
    "exam:build": build_exam,
    "exam:build-ab": build_exam_pair,
    "exam:build-teacher": build_exam_teacher,
    "handout:build": build_dictation,
    "handout:audio": build_dictation_audio,
    "audio:script": build_script_audio,
    "bank:build": bank_build,
    "bank:inspect": bank_inspect,
    "plan:build": build_plan,
    "plan:parse": parse_plan,
    "exam:parse": exam_parse,
    "lesson:deck": build_deck,
    "lesson:outline": plan_deck,
    "lesson:html": build_deck_html,
    "docx:clean": docx_clean,
    "docx:merge": docx_merge,
    "docx:extract": docx_extract,
    "preflight": preflight,
    "template:make": template_make,
    "template:model": template_model,
    "template:diff": template_diff,
    "selfcheck": selfcheck,
    "inspect": inspect,
    "papers": papers,
    "docx:new": docx_new,
    "pptx:new": pptx_new,
}

# 所有命令走同一道文件名关：调用方各自 os.path.abspath 之前先查（见 _guard_raw_paths）
TABLE = {name: _guarded(function) for name, function in TABLE.items()}
