"""造 docx：从一份"结构化内容"生成 Word 文档。

这里不负责设计版式 —— 版式来自学校模板（模板反推在 template_model.py）。
这里只保证：结构化内容 → 段落/表格/分页/分栏，且每一块都能追回源头题目。
"""

import copy
import datetime
import os

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Twips

from .errors import EngineError, require
from .paper import resolve as resolve_paper
from .paths import guard_output_path
from .placeholders import fill_document

CN_BODY_FONT = "宋体"
CN_BODY_SIZE_PT = 12.0  # 小四
CN_LINE_SPACING = 1.5   # 教案/试卷的常规行距

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def school_year(today=None):
    """学年度：8 月开学 → "2026—2027"。模板里的 {{schoolYear}} 用它兜底。"""
    today = today or datetime.date.today()
    start = today.year if today.month >= 8 else today.year - 1
    return str(start) + "—" + str(start + 1)


def _default_values(values):
    """占位符兜底值。

    模板反推会把「使用班级：____」变成 {{blank}}，而出卷时老师没有填班级的入口，
    早先页脚就直接印出「使用班级：{{blank}}」（persona 评测里最难看的一条）。
    这里把没人填的空白位渲染成下划线、学年按当前时间算，剩下的真占位符仍会被
    印前检查的 placeholders-left 拦下。
    """
    values = dict(values or {})
    values.setdefault("blank", "______")
    values.setdefault("schoolYear", school_year())
    for alias, target in (("className", "class"), ("teacher", "teacherName"),
                          ("subject", "course"), ("title", "examTitle")):
        if values.get(alias) and not values.get(target):
            values[target] = values[alias]
    # 页脚「使用班级 / 出卷老师」：面板填了就印真名，没填就留手写横线（别名之后再兜底，
    # 否则 "teacher" 这类写法会被下面的默认值挡住）
    values.setdefault("className", values.get("class") or "__________")
    values.setdefault("teacherName", values.get("teacher") or "__________")
    values.setdefault("class", values["className"])
    values.setdefault("teacher", values["teacherName"])
    return values


def _ensure_parent(path):
    parent = os.path.dirname(os.path.abspath(path))
    if parent and not os.path.isdir(parent):
        try:
            os.makedirs(parent, exist_ok=True)
        except OSError as error:
            # U 盘拔了/盘符不存在这种真事：以前把 WinError 3 的英文原文端给老师
            raise EngineError(
                "output-dir-unwritable",
                "这个位置存不进去：" + parent,
                "盘符或文件夹可能被移走了（比如 U 盘拔了、网盘掉了）；换到「桌面」或「文档」再试（" + str(error)[:60] + "）")


def _apply_chinese_defaults(document):
    """新文档不留 Calibri：正文回到宋体小四、1.5 倍行距（中文本该如此）。"""
    style = document.styles["Normal"]
    style.font.name = CN_BODY_FONT
    style.font.size = Pt(CN_BODY_SIZE_PT)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), CN_BODY_FONT)
    try:
        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing = CN_LINE_SPACING
    except Exception:
        pass  # 样式表被模板锁住时不影响生成


def _normalize_doc_defaults(document):
    """默认字体写成明确的字体名，不留主题字体（asciiTheme/minorHAnsi）。

    留主题字体时，打开文档的机器按自己的主题渲染，换台电脑版式就变；
    东亚语言也顺手设成 zh-CN，否则 Word 会把中文按 en-US 排（标点、行距都不对）。
    """
    doc_defaults = document.styles.element.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    rpr = rpr_default.find(qn("w:rPr")) if rpr_default is not None else None
    if rpr is None:
        return
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), CN_BODY_FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = rpr.makeelement(qn("w:lang"), {})
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def _neutralize_template_fonts(document):
    """python-docx 的默认模板里 MacroText 等样式写着 Courier。

    正文一个字都没用它，但导出/预览会报"本机缺字体 Courier"，老师看着像出了错。
    换成宋体，这条假警报就没了。只动默认模板，不动老师自己的模板。
    """
    changed = 0
    for style in document.styles:
        try:
            name = style.font.name
        except Exception:
            continue
        if name and name.lower().startswith("courier"):
            style.font.name = CN_BODY_FONT
            changed += 1
    for part in document.part.package.parts:
        if "stylesWithEffects" not in str(part.partname):
            continue
        blob = part.blob
        if b"Courier" not in blob:
            continue
        blob = blob.replace(b"Courier New", CN_BODY_FONT.encode("utf-8"))
        # 只替完整字体名：整体替 "Courier" 会把「Courier New」变成「宋体 New」这种不存在的字体
        part._blob = blob
        changed += 1
    return changed


# 班级/老师走占位符，由 _default_values 兜底成手写横线：填了真名就印真名（老师不用每张都手写）
DEFAULT_FOOTER = "第 {{page}} 页 共 {{pages}} 页　　使用班级：{{className}}　　出卷老师：{{teacherName}}"


def _default_footer(document, values=None, text=None):
    """没有学校模板时也得有页码页脚。

    老师手打「第 1 页」，打印出来整本都是第 1 页 —— 印前检查会拦成 error，
    可这页脚是工具自己该给的，不该让老师去补。{{page}}/{{pages}} 由
    fill_document 换成真域代码（PAGE / NUMPAGES），页码才会自己走。
    """
    section = document.sections[0]
    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text or DEFAULT_FOOTER)
    run.font.size = Pt(9)


def _apply_page(document, page, preserve):
    """page = 纸张规格；preserve=True（模板已定版式）时只在显式给了 page 才改。"""
    if preserve and not page:
        return None
    layout = resolve_paper(page)
    for section in document.sections:
        section.orientation = WD_ORIENT.LANDSCAPE if layout["orientation"] == "landscape" else WD_ORIENT.PORTRAIT
        section.page_width = Mm(layout["widthMm"])
        section.page_height = Mm(layout["heightMm"])
        section.top_margin = Mm(layout["marginMm"]["top"])
        section.bottom_margin = Mm(layout["marginMm"]["bottom"])
        section.left_margin = Mm(layout["marginMm"]["left"])
        section.right_margin = Mm(layout["marginMm"]["right"])
        sect_pr = section._sectPr
        cols = sect_pr.find(qn("w:cols"))
        if cols is None:
            cols = sect_pr.makeelement(qn("w:cols"), {})
            sect_pr.append(cols)
        cols.set(qn("w:num"), str(layout["columns"]))
    return layout


def _flatten_sections(document):
    """把模板里的多节压成一节：拆掉段落级 w:sectPr，只保留 body 末尾那个页面设置。"""
    body = document.element.body
    for paragraph in list(body.findall(qn("w:p"))):
        ppr = paragraph.find(qn("w:pPr"))
        if ppr is None:
            continue
        section = ppr.find(qn("w:sectPr"))
        if section is None:
            continue
        ppr.remove(section)
        if len(ppr) == 0:
            paragraph.remove(ppr)
    final = body.find(qn("w:sectPr"))
    if final is not None:
        kind = final.find(qn("w:type"))
        if kind is not None:
            final.remove(kind)


def _end_section_with_columns(document, count):
    """在当前位置结束这一节，并指定它的栏数（后续内容进新起的一节）。

    用途：8 开横排两栏的卷子，标题/说明/班级行要**跨整页**，正文才分两栏。
    不这么切，标题会被挤在左栏里（老师一眼看出不像学校的卷子）。
    分节符用「连续」，所以标题和正文还在同一页上。
    """
    document.add_section(WD_SECTION.CONTINUOUS)
    sections = document.sections
    ended = sections[-2] if len(sections) >= 2 else sections[-1]
    sect_pr = ended._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = sect_pr.makeelement(qn("w:cols"), {})
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(max(1, int(count))))
    return ended


def _loose_break(document):
    """在分节符之后垫一个不带 keep 的极小空段落。

    正文第一节常以「大题标题(keepWithNext) + 题干(keepTogether) + 选项」这样的 keep 链开头；
    LibreOffice 遇到"连续分节 + 改栏数"的断点，如果链在页内放不下，会把**整节**推到下一页
    —— 实测 8 开两栏、题干 14 行的卷子，第 1 页只剩卷头三行，白印一张纸。
    断点落在一个可以自由折断的段落上就不会整节搬走。空段落压到 1 磅，肉眼看不出。
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    run = paragraph.add_run("")
    run.font.size = Pt(1)
    return paragraph


_BAD_COLORS = set()


def _normalize_hex(value, length=6):
    """'#C00000' / 'c00000' / 'f00' → 'C00000' / 'FF0000'；认不出就返回空串。"""
    text = str(value or "").strip().lstrip("#")
    if len(text) == 3 and all(char in "0123456789abcdefABCDEF" for char in text):
        text = "".join(char * 2 for char in text)
    if len(text) != length or any(char not in "0123456789abcdefABCDEF" for char in text):
        return ""
    return text.upper()


def _number_or_raise(value, label):
    """把「行距/段后/缩进」这类数字字段转成 float。

    空值留给调用方按默认处理；写了却不是数字（"abc"、{"x":1}）以前一路传到 python-docx，
    报成英文 ValueError，老师只会看到「引擎内部错误」。这里换成一句人话。
    """
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raise EngineError(
            "bad-spacing",
            label + "要填数字：这次收到的是「" + str(value)[:20] + "」",
            "写成 1.5、18、28 这样的数字（行距按磅：1.5 倍 ≈ 27、单倍 ≈ 18）")


def _write_paragraph(document, block, container=None):
    # container 为空时写正文；传 header/footer 时写进页眉页脚（卷头就放在第一页页眉里）
    paragraph = (container or document).add_paragraph()
    style = block.get("style")
    if style:
        try:
            paragraph.style = style
        except Exception:
            pass  # 模板里没有这个样式就用默认的，不因为样式名不对整篇失败
    align = block.get("align")
    if align in ALIGN:
        paragraph.alignment = ALIGN[align]
    if block.get("keepWithNext"):
        # 题干和它的选项不能被栏/页切开（切开后选项会跑到上一栏顶上，学生看得一头雾水）
        paragraph.paragraph_format.keep_with_next = True
    if block.get("keepTogether"):
        paragraph.paragraph_format.keep_together = True
    pieces = block.get("runs")
    if pieces is None:
        pieces = [{"text": block.get("text", "")}]
    elif isinstance(pieces, str):
        # 把 runs 直接写成一句话（常见手误）：按纯文本处理，不要抛 'str' object has no attribute 'get'
        pieces = [{"text": pieces}]
    elif not isinstance(pieces, list):
        raise EngineError(
            "invalid-blocks",
            "段落里的 runs 要写成列表，这次收到的是 " + type(pieces).__name__,
            "写成 [{\"text\": \"第一段\"}]；只想写一句话就直接给 text")
    for position, piece in enumerate(pieces):
        if isinstance(piece, str):
            piece = {"text": piece}
        elif not isinstance(piece, dict):
            raise EngineError(
                "invalid-blocks",
                "第 " + str(position + 1) + " 个 run 不是对象，是 " + type(piece).__name__,
                "每个 run 写成 {\"text\": \"…\"}；只写文字就直接给 text")
        run = paragraph.add_run(piece.get("text", ""))
        if piece.get("bold"):
            run.bold = True
        if piece.get("italic"):
            run.italic = True
        if piece.get("sizePt") not in (None, ""):
            try:
                run.font.size = Pt(float(piece["sizePt"]))
            except (TypeError, ValueError):
                raise EngineError(
                    "bad-size",
                    "字号写得不对：" + str(piece.get("sizePt")),
                    "用数字（单位磅：三号 16、小四 12、二号 22）")
        if piece.get("font"):
            set_run_font(run, piece["font"], piece.get("sizePt"))
        if piece.get("color"):
            hex_text = _normalize_hex(piece.get("color"))
            if hex_text:
                from docx.shared import RGBColor
                run.font.color.rgb = RGBColor.from_string(hex_text)
            else:
                # 颜色写错不该让整篇失败（红头退回黑色照印），但得让老师知道字色没生效
                _BAD_COLORS.add(str(piece.get("color")))
    # 公文与公众号稿的排版靠这几个量：行距、段后、首行缩进（按字符数给，跟着字号走）
    spacing = _number_or_raise(block.get("lineSpacingPt"), "行距（lineSpacingPt）")
    if spacing:
        paragraph.paragraph_format.line_spacing = Pt(spacing)
    after = _number_or_raise(block.get("spaceAfterPt"), "段后间距（spaceAfterPt）")
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    indent = _number_or_raise(block.get("firstLineIndentChars"), "首行缩进字数（firstLineIndentChars）")
    if indent:
        size = _number_or_raise(block.get("sizePt") or block.get("fontSizePt") or 12, "字号（sizePt）") or 12
        paragraph.paragraph_format.first_line_indent = Pt(indent * size)
    if block.get("pageBreakBefore"):
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def set_run_font(run, name, size_pt=None):
    """给一段文字指定中英文字体：不留主题字体，东亚字符也吃到这个字体。

    只设 run.font.name 时 Word 会拿主题字体排中文，换台电脑版式就变（公文尤其要命）。
    """
    run.font.name = name
    if size_pt:
        try:
            run.font.size = Pt(float(size_pt))
        except (TypeError, ValueError):
            pass
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    return run


EMU_PER_TWIP = 635  # 1 twip = 1/1440 英寸 = 635 EMU
DEFAULT_COLUMN_SPACE_TWIPS = 720  # Word 分栏默认栏间距 0.5 英寸


def column_width_twips(section):
    """这一节里"一栏"有多宽（twips）。

    双栏卷子里表格必须按栏宽画 —— 按整版心宽画会横跨两栏，是打印出来最扎眼的一种坏。
    """
    page = section.page_width
    if page is None:
        return None
    text_emu = int(page) - int(section.left_margin or 0) - int(section.right_margin or 0)
    if text_emu <= 0:
        return None
    num = 1
    space = DEFAULT_COLUMN_SPACE_TWIPS
    cols = section._sectPr.find(qn("w:cols"))
    if cols is not None:
        try:
            num = max(1, int(cols.get(qn("w:num")) or 1))
        except (TypeError, ValueError):
            num = 1
        try:
            space = max(0, int(cols.get(qn("w:space")) or DEFAULT_COLUMN_SPACE_TWIPS))
        except (TypeError, ValueError):
            space = DEFAULT_COLUMN_SPACE_TWIPS
    width_emu = (text_emu - space * (num - 1) * EMU_PER_TWIP) // num
    twips = int(round(width_emu / EMU_PER_TWIP))
    return twips if twips > 0 else None


def fit_table(table, width_twips):
    """把表格钉在一栏之内：tblW + tblGrid + 每个 tcW 都写成固定值，别让 Word 自己摊开。"""
    if not width_twips or width_twips <= 0:
        return None
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = tbl_pr.makeelement(qn("w:tblW"), {})
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(width_twips)))
    try:
        table.autofit = False  # 关掉自动伸缩，写入 <w:tblLayout w:type="fixed"/>
    except Exception:
        pass
    count = max(1, len(table.columns))
    base = int(width_twips) // count
    rest = int(width_twips) - base * count
    grid = tbl.tblGrid
    if grid is not None:
        for index, grid_col in enumerate(grid.findall(qn("w:gridCol"))):
            grid_col.set(qn("w:w"), str(base + (rest if index == count - 1 else 0)))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            try:
                cell.width = Twips(base + (rest if index == count - 1 else 0))
            except Exception:
                continue  # 单格设不上就跳过这一格，别让后面列的宽度全错位
    return int(width_twips)


def _write_table(document, block):
    rows = block.get("rows") or []
    require(len(rows) > 0, "empty-table", "表格没有任何一行", "表格至少要有一行内容；只想排版就用段落块")
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    style = block.get("tableStyle")
    if style:
        try:
            table.style = style
        except Exception:
            pass
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            table.cell(r_index, c_index).text = "" if value is None else str(value)
    if block.get("repeatHeader") and len(table.rows) > 1:
        # 教学过程表一跨页，第二页就只剩光秃秃的格子 —— 老师们手写的教案都写「续表」
        trPr = table.rows[0]._tr.get_or_add_trPr()
        flag = OxmlElement("w:tblHeader")
        flag.set(qn("w:val"), "true")
        trPr.append(flag)
    if block.get("noRowSplit"):
        # 一行跨页会被从中间切断：续页上第 1 列（环节）整格空白，老师看到的是
        # 「一页表格中间突然断了」。w:cantSplit 让整行挪到下一页。
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
    height_mm = block.get("rowHeightMm")
    if height_mm:
        # 作文作答横线那种「一行一格」的表：行高要够写字（默认行高挤成一坨）
        try:
            height = Mm(float(height_mm))
        except (TypeError, ValueError):
            height = None
        if height is not None:
            for row in table.rows:
                row.height = height
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    width = block.get("widthTwips") or column_width_twips(document.sections[-1])
    fit_table(table, width)
    return table


def _clear_paragraphs(part):
    for paragraph in list(part.paragraphs):
        element = paragraph._element
        if element.getparent() is not None:
            element.getparent().remove(element)


def _mirror_footer(section):
    """把默认页脚抄一份给「首页页脚」。

    打开 different_first_page_header_footer 后，第 1 页用的是 first_page_footer；
    不补的话第 1 页就没有页码/学校页脚（老师一看就以为丢了）。
    """
    source = section.footer
    target = section.first_page_footer
    try:
        target.is_linked_to_previous = False
    except Exception:
        pass
    _clear_paragraphs(target)
    for paragraph in source.paragraphs:
        target._element.append(copy.deepcopy(paragraph._element))


def _write_first_page_header(document, blocks, template):
    """卷头（卷名/分值说明/班级行）写进第一页页眉，正文才能整页两栏。

    为什么非得放页眉：8 开两栏卷子的卷头要跨整页宽。以前的做法是「标题单栏节 + 连续分节切
    两栏」，LibreOffice 遇到这种"正文中途改栏数"的分节，会把正文整节推到下一页（实测多印一张
    8 开纸），或者把正文塞进右栏、左半页全空。放页眉里两个排版引擎都不会错：
    页眉本来就跨整页宽，正文从页眉下面按两栏排。
    """
    section = document.sections[0]
    section.different_first_page_header_footer = True
    header = section.first_page_header
    try:
        header.is_linked_to_previous = False
    except Exception:
        pass
    if not template:
        # 空白文档：清掉自带的空段，免得卷头上方多一行
        _clear_paragraphs(header)
    else:
        # 套模板时模板的校名/学年在**默认页眉**里。这里开了「首页不同」，第 1 页会改用
        # 一份全新的空首页页眉 —— 不把默认页眉搬过来，学校校名在首页就整块丢了
        # （老师会觉得"套了模板怎么校名没了"）。先搬过来，再把卷头写在下面。
        source = section.header
        empty = all(not paragraph.text.strip() for paragraph in header.paragraphs)
        if empty and not header.tables:
            for paragraph in list(header.paragraphs):
                element = paragraph._element
                if element.getparent() is not None:
                    element.getparent().remove(element)
            for paragraph in source.paragraphs:
                header._element.append(copy.deepcopy(paragraph._element))
            for table in source.tables:
                header._element.append(copy.deepcopy(table._element))
    for block in blocks:
        _write_paragraph(document, block, container=header)
    _mirror_footer(section)


def _mark_update_fields(document):
    """让 Word/WPS 打开文件时刷新域（页码、总页数）。

    套了学校模板时，模板页脚里的 PAGE/NUMPAGES 域带着上一份的缓存值（「第 3 页 共 3 页」），
    不刷新就原样印到新卷上 —— 一份 8 页的卷子页脚写着「共 3 页」。
    """
    try:
        settings = document.settings.element
    except Exception:
        return
    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    settings.append(element)


def _text_block_caption(block):
    for key in ("caption", "alt", "text"):
        value = block.get(key)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _write_picture(document, block):
    """插一张图（公众号稿、宣传稿要用）。

    宽度按版心算，不按像素原尺寸贴——手机拍的图动辄 4000px 宽，按原尺寸会把版面撑爆。
    图找不到时报清楚是哪个文件，不让它悄悄变成一张空白。
    """
    path = block.get("path") or block.get("src")
    if not path:
        raise EngineError("picture-path-required", "这段配图没给文件路径")
    path = os.path.abspath(str(path))
    if not os.path.isfile(path):
        raise EngineError(
            "picture-missing",
            "配图找不到：" + os.path.basename(path),
            "确认图片还在原位置；手机拍的图先存到电脑上再选")
    raw_width = block.get("widthMm")
    width = None
    if raw_width not in (None, ""):
        try:
            width = Mm(float(raw_width))
        except (TypeError, ValueError):
            width = None
        if width is not None and not (10 <= float(raw_width) <= 400):
            # 负数/0 宽度以前一路传到 python-docx，报成「这张图打不开」，让人白换图
            raise EngineError(
                "bad-picture-width",
                "配图宽度不合理：" + str(raw_width) + " 毫米",
                "按毫米给 10—400 之间的数（默认 140；整页宽约 160，半栏约 75）")
    if width is None:
        width = Mm(140)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    try:
        run.add_picture(path, width=width)
    except Exception as error:
        detail = str(error)
        if "mode" in detail.lower() or "cmyk" in detail.lower() or "color" in detail.lower():
            hint = "这张图是 CMYK 印刷色/WEBP 之类的特殊格式，先另存成 RGB 的 jpg 或 png 再放进来（" + detail[:60] + "）"
        elif "range" in detail.lower():
            hint = "图片尺寸参数不合适：把「宽度(mm)」改成 10—400 之间的数再试（" + detail[:60] + "）"
        elif path.lower().endswith((".jpg", ".jpeg")):
            # 印刷用的 CMYK jpg：python-docx 抛的异常常常一个字都没有，原来提示里只剩一对空括号，
            # 还叫老师「换成 jpg」——它本来就是 jpg。
            hint = "这张 jpg 多半是印刷用的 CMYK 色彩，Word 认不了：用「画图」打开，另存为一张新的 jpg 再试"
        else:
            hint = "换成 jpg / png 格式再试" + ("（" + detail[:60] + "）" if detail.strip() else "")
        raise EngineError(
            "picture-unreadable",
            "这张图打不开：" + os.path.basename(path),
            hint)
    caption = _text_block_caption(block)
    if caption:
        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.space_after = Pt(10)
        note.add_run(caption).font.size = Pt(9)


def build_docx(spec):
    """spec = {outputPath, baseTemplate?, blocks:[{type:'paragraph'|'table'|'pageBreak', ...}], properties?}

    给了 baseTemplate 就从模板开始写（保留页面设置、页眉页脚、样式），否则用空白文档。
    blocks 里可以有 {type:"columns", count:1}：在那一处插一个（连续）分节符，
    让前面的标题块单栏排版、后面的正文按 page.columns 分栏。
    """
    output = spec.get("outputPath")
    require(bool(output), "output-required", "没有指定输出文件路径", "在面板上选一个保存位置，或者把 outputPath 写成完整路径")
    if not isinstance(output, str):
        # 传 12345 这种：以前 os.path.abspath 直接抛 TypeError，报成英文内部错误
        raise EngineError(
            "output-required",
            "输出路径要写成一段文字（例如 D:/文档/试卷.docx），这次收到的是 " + type(output).__name__,
            "把 outputPath 写成字符串；不要传数字或对象")
    blocks = spec.get("blocks")
    if blocks is not None and not isinstance(blocks, list):
        raise EngineError(
            "invalid-blocks",
            "正文内容（blocks）要写成一段一段的列表，这次收到的是 " + type(blocks).__name__,
            "每一段写成 {\"type\": \"paragraph\", \"text\": \"…\"}，整份是一个 [ … ] 列表")
    for position, block in enumerate(blocks or []):
        if not isinstance(block, dict):
            raise EngineError(
                "invalid-blocks",
                "第 " + str(position + 1) + " 段不是对象，是 " + type(block).__name__,
                "每一段写成 {\"type\": \"paragraph\", \"text\": \"…\"} 这样的对象")
    _BAD_COLORS.clear()
    guard_output_path(output)
    # 先查再建：以前检查放在 document.save 之前，可整份文件已经排完了（还留过半截 docx）
    if os.path.isdir(output):
        # 老师挑了个文件夹当文件名：以前报「目标文件已存在」，看着像撞名
        raise EngineError(
            "output-is-folder",
            "这个路径是个文件夹，不是文件名：" + (os.path.basename(output) or output),
            "在路径后面补一个文件名，比如 …/" + (os.path.basename(output) or "文件") + ".docx")
    if os.path.exists(output) and not spec.get("overwrite", False):
        raise EngineError(
            "output-exists",
            "目标文件已存在，我不覆盖：" + os.path.basename(output),
            "换一个文件名，或者加 overwrite:true 明确允许覆盖")
    _ensure_parent(output)

    if not isinstance(spec.get("values"), dict):
        spec["values"] = {}
    template = spec.get("baseTemplate")
    if template:
        require(os.path.isfile(template), "template-missing", f"模板文件不存在：{template}")
        document = Document(template)
        # 从模板开始写时，先清空正文占位内容（保留节、页眉页脚、样式定义）
        if spec.get("clearBody", True):
            body = document.element.body
            for child in list(body):
                if child.tag.endswith("}sectPr"):
                    continue
                body.remove(child)
        # 模板里残留的"分节符"会把新写的正文分到别的节里：LibreOffice 会把它和两栏节并成
        # 一块，印出来直接变通栏（实测：用自己出的 8 开两栏卷反推的模板出卷，通栏）。
        # 模板的正文本来就是空的，所以把段落级分节符全拆掉、只留一个页面设置，再按本次的
        # 标题/正文分栏重新分节 —— 版面按模板，结构按这次的内容。
        _flatten_sections(document)
        try:
            if document.sections:
                document.sections[0].start_type = WD_SECTION.NEW_PAGE
        except Exception:
            pass
    else:
        document = Document()
        _apply_chinese_defaults(document)
        _normalize_doc_defaults(document)
        _neutralize_template_fonts(document)
        footer_text = spec.get("footer")
        if footer_text is None and spec.get("defaultFooter", True):
            footer_text = DEFAULT_FOOTER
        if footer_text:
            _default_footer(document, text=str(footer_text))

    _mark_update_fields(document)
    layout = _apply_page(document, spec.get("page"), preserve=bool(template))

    # 模板里留的 {{schoolYear}}/{{page}} 之类，造新卷时在这里一次性填掉
    values = _default_values(spec.get("values"))
    spec["values"] = values

    header_blocks = spec.get("firstPageHeaderBlocks") or []
    if header_blocks:
        _write_first_page_header(document, header_blocks, template)

    previous_is_table = False
    for block in spec.get("blocks") or []:
        kind = block.get("type", "paragraph")
        if kind == "table" and previous_is_table:
            # 两张表紧挨着写，Word / LibreOffice 会当成一张表：后面那张的「标题行重复」
            # 就失效了（标在中间的 w:tblHeader 不算数），跨页时表头不再重现 ——
            # 教案「教学过程」表紧跟教学重点/难点表时就是这样（0.18 审查 L-01）。
            # 垫一个 1 磅空段落把两张表分开，肉眼看不出来。
            _loose_break(document)
        if kind == "paragraph":
            _write_paragraph(document, block)
        elif kind == "table":
            _write_table(document, block)
        elif kind == "picture":
            _write_picture(document, block)
        elif kind == "pageBreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "columns":
            # 结束当前节、给它设栏数；后面的内容进新一节（栏数仍按 page 设置）
            _end_section_with_columns(document, block.get("count", 1))
            _loose_break(document)
        else:
            raise EngineError("unknown-block", f"不认识的内容块类型：{kind}",
                          "type 只认 paragraph / table / pageBreak / columns / picture")
        previous_is_table = kind == "table"

    # 占位符最后统一填：以前在写 blocks 之前填，正文里写的 {{className}} 永远换不成真名
    filled = fill_document(document, values)

    properties = spec.get("properties") or {}
    if properties:
        props = document.core_properties
        for field in ("title", "subject", "author", "comments"):
            if field in properties:
                setattr(props, field, properties[field])

    document.save(output)

    warnings = list(spec.get("_warnings") or [])
    if filled["unfilled"]:
        warnings.append({
            "code": "unfilled-placeholder",
            "message": "这份模板里还有没填的占位符：" + "、".join("{{" + name + "}}" for name in filled["unfilled"]),
            "placeholders": filled["unfilled"],
        })
    if _BAD_COLORS:
        warnings.append({
            "code": "color-fallback",
            "message": "这些颜色写得不认识，按黑色印了：" + "、".join(sorted(_BAD_COLORS))
                       + "（颜色用 6 位十六进制，比如 C00000、2563EB）",
        })
    return {
        "outputPath": os.path.abspath(output),
        "blocks": len(spec.get("blocks") or []),
        "bytes": os.path.getsize(output),
        "fromTemplate": bool(template),
        "page": layout,
        "values": values,
        "placeholdersFilled": filled["filled"],
        "warnings": warnings,
    }