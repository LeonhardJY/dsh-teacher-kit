"""文件加工：老师手里那些"从网上下载/同事转发/去年剩的"文档，先弄干净再谈排版。

三件事，覆盖每周真实会遇到的重复劳动：
- clean  净化：去掉网页粘来的超链接、彩色字、高亮、底纹、大段间距、批注，统一字体
- merge  合并：把几份 docx 拼成一份（练习 + 知识点 + 答案），图片、自动编号、样式一并带过去
- extract 提纯：把文档里的话抽成纯文本，喂给 AI 或贴进微信群

不做的：不改内容、不猜格式意图、不动页眉页脚里的学校抬头（那是学校的，不是脏数据）；
只把页眉页脚里"www.xxjyw.com 免费下载 请勿商用"这类网站广告行清掉。
"""

import os
import posixpath
import re
import shutil
import zipfile
from copy import deepcopy

import docx
from docx.enum.section import WD_SECTION
from docx.opc.exceptions import PackageNotFoundError
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn

from .docx_build import _apply_page, _ensure_parent
from .paths import default_output
from .errors import EngineError, require

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_IMAGE = R_NS + "/image"
RT_HYPERLINK = R_NS + "/hyperlink"
# 批注是一整套部件（comments / commentsExtended / commentsIds / commentsExtensible / people）。
# 正文里留一个批注锚点、包里却没有对应部件，Word 打开就报"文件可能已经损坏"——所以锚点和部件要一起处理。
COMMENT_RELATIONSHIPS = (
    R_NS + "/comments",
    R_NS + "/commentsExtended",
    R_NS + "/commentsIds",
    R_NS + "/commentsExtensible",
    R_NS + "/people",
)
# python-docx 的 qn() 不认 VML 前缀（v:），这两个标签要手写命名空间
TAG_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
TAG_VML_IMAGE = "{urn:schemas-microsoft-com:vml}imagedata"
KEEP_TAGS = ("w:drawing", "w:object", "w:pict", "w:fldSimple", "w:fldChar", "w:instrText",
             "w:footnoteReference", "w:endnoteReference", "w:sym")

# 页眉页脚里的网站广告行：命中这些字样就整段删掉。只删"一眼是广告"的行，
# 校名 / 教研组 / 年级 / 「第 N 页」这类真正属于学校的抬头一律保留（见 _looks_like_ad）。
AD_MARKERS = ("免费下载", "请勿商用", "禁止商用", "请勿转载", "禁止转载", "资源网", "课件",
              "下载地址", "扫码关注", "微信公众号", "客服电话", "QQ群", "题库网",
              # 下载站/代写铺子的页眉常见字眼：只认这些具体词，不认「微信」单个词（学校通知里
              # 也会写「微信扫描班级群」，那种行还带校名，后面 _looks_like_ad 会放过）
              "代写", "加群", "公众号", "扫码加", "加微信", "联系QQ")
URL_PATTERN = re.compile(r"(https?://|www\.)", re.I)
SCHOOL_MARKERS = ("中学", "小学", "学校", "学院", "大学", "教研", "教务处", "年级", "班级",
                  "校区", "附属")

# 超链接文字的蓝色不是写在 run 上的，而是来自字符样式（Hyperlink / 超链接）：
# 只删 w:color 文字还是蓝的，得把 w:rStyle 一起摘掉（见 _hyperlink_style_ids）。
LINK_STYLE_IDS = {"hyperlink", "followedhyperlink", "internetlink"}
LINK_STYLE_NAMES = {"hyperlink", "followedhyperlink", "internetlink", "超链接", "已访问的超链接"}

# w:rPr 的子元素也有顺序（CT_RPr）。w:sz / w:szCs 插错位置 Word 会说"文件需要修复"。
SZ_SUCCESSORS = ("w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
                 "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                 "w:specVanish", "w:oMath")
SZCS_SUCCESSORS = SZ_SUCCESSORS[1:]

UNSUPPORTED_HINT = "用 Word/WPS 打开这个文件，点「文件 → 另存为」选 .docx 再试；如果设了密码，先去密码"

# w:pPr 的子元素是有顺序的（CT_PPrBase）。顺序插错的段落 Word 打开时会提示"需要修复"，
# 所以凡是往 pPr 里加东西，都按这些后继标签插到它该在的位置。
PAGE_BREAK_SUCCESSORS = (
    "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
SPACING_SUCCESSORS = (
    "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _ppr_insert(ppr, tag_name, successors):
    """按 CT_PPrBase 的顺序把 tag_name 插进 pPr（已有就原样返回）。"""
    element = ppr.find(qn(tag_name))
    if element is not None:
        return element
    element = ppr.makeelement(qn(tag_name), {})
    sequence = getattr(ppr, "_tag_seq", ())
    if tag_name in sequence:
        successors = tuple(sequence[list(sequence).index(tag_name) + 1:])
    ppr.insert_element_before(element, *successors)
    return element


def _open_document(path, kind="文件"):
    """打开 docx，把"不是 docx / 加密 / 下载不完整"统一翻成老师看得懂的话。

    .doc、Office 加密文件、半截 zip 都不是 OOXML 包，python-docx 抛的是
    PackageNotFoundError / BadZipFile（英文异常名），直接漏到界面上就是"引擎内部错误"。
    """
    try:
        return docx.Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile) as error:
        raise EngineError(
            "unsupported-file",
            "这个" + str(kind) + "打不开，它可能不是 .docx，也可能已经损坏或设了密码：" + os.path.basename(str(path)),
            UNSUPPORTED_HINT,
        ) from error
    except (KeyError, ValueError, SyntaxError) as error:
        # 包里缺 [Content_Types].xml / document.xml，或者 XML 本身坏了
        raise EngineError(
            "unsupported-file",
            "这个" + str(kind) + "的结构不完整，Word 也打不开：" + os.path.basename(str(path)),
            UNSUPPORTED_HINT,
        ) from error


def _header_footer_parts(document):
    """每一节的页眉页脚部件（同一个部件被多节共用时只给一次）。

    is_linked_to_previous 的部件内容和上一节共用，内容在上一节那份里，跳过免得同一段落被处理两遍。
    """
    seen = set()
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header, section.even_page_footer):
            try:
                if part.is_linked_to_previous:
                    continue
            except Exception:
                continue
            element = getattr(part, "_element", None)
            if element is None or id(element) in seen:
                continue
            seen.add(id(element))
            yield part, element


def _iter_paragraphs(document):
    """正文 + 页眉页脚的所有 w:p，返回 (段落, 是否在正文里)。

    页眉页脚里的彩色字一样碍眼，所以也要去色；但字号字体、段后留白只按正文的口径调，
    免得把学校抬头的版式改样。
    """
    for paragraph in document.element.body.iter(qn("w:p")):
        yield paragraph, True
    for _part, root in _header_footer_parts(document):
        for paragraph in root.iter(qn("w:p")):
            yield paragraph, False


def _paragraph_text(paragraph_element):
    return "".join(node.text or "" for node in paragraph_element.iter(qn("w:t"))).strip()


def _looks_like_ad(text):
    """页眉页脚里这行是不是网站广告（"www.xxjyw.com 免费下载 请勿商用"）。

    判据保守：命中明显广告词就删；只有一条网址时，带校名/教研组字样的留下（学校官网也是学校的），
    整段就是一行网址又没有任何学校信息的，才算广告。
    """
    text = (text or "").strip()
    if not text:
        return False
    if any(marker in text for marker in AD_MARKERS):
        return True
    if not URL_PATTERN.search(text):
        return False
    if any(marker in text for marker in SCHOOL_MARKERS):
        return False
    return len(text) <= 60


def _strip_ad_paragraphs(part_element):
    """删掉页眉/页脚里的广告段，返回删了几段。

    删空了要补一个空段：w:hdr / w:ftr 至少得有一个块级元素，空页眉 Word 会说"文件需要修复"。
    """
    removed = 0
    for paragraph in list(part_element.iter(qn("w:p"))):
        if paragraph.getparent() is None:
            continue
        if not _looks_like_ad(_paragraph_text(paragraph)):
            continue
        # 「第 PAGE 页 共 NUMPAGES 页 www.1ppt.com 免费下载」这种整条页脚也长得像广告：
        # 整段删掉页码域就没了（实测页脚只剩空白）。带域的段只摘广告 run，域留着。
        has_field = paragraph.find(qn("w:fldSimple")) is not None or any(
            run.find(qn("w:fldChar")) is not None or run.find(qn("w:instrText")) is not None
            for run in paragraph.iter(qn("w:r")))
        if has_field:
            trimmed = 0
            for run in list(paragraph.iter(qn("w:r"))):
                if run.find(qn("w:fldChar")) is not None or run.find(qn("w:instrText")) is not None:
                    continue
                run_text = "".join(node.text or "" for node in run.iter(qn("w:t")))
                if URL_PATTERN.search(run_text) or _looks_like_ad(run_text):
                    holder = run.getparent()
                    if holder is not None:
                        holder.remove(run)
                        trimmed += 1
            if trimmed:
                removed += 1
            continue
        parent = paragraph.getparent()
        parent.remove(paragraph)
        removed += 1
    if removed and part_element.find(qn("w:p")) is None:
        part_element.append(part_element.makeelement(qn("w:p"), {}))
    return removed


def _unwrap_hyperlink(hyperlink):
    """把一个 w:hyperlink 拆掉、里面的文字按原位留下。返回是否拆成功。"""
    parent = hyperlink.getparent()
    if parent is None:
        return False
    index = list(parent).index(hyperlink)
    for child in list(hyperlink):
        parent.insert(index, child)
        index += 1
    parent.remove(hyperlink)
    return True


def _unwrap_hyperlinks(root, stats):
    """w:hyperlink 包着的文字要留下，链接本身要拆掉——不然点一下就跳到别的网站。"""
    for hyperlink in list(root.iter(qn("w:hyperlink"))):
        if _unwrap_hyperlink(hyperlink):
            stats["links"] += 1


def _strip_comment_marks(root, ids=None):
    """拔掉批注锚点（w:commentRangeStart/End/w:commentReference）。

    ids 收集批注 id，用来数"清掉了几处批注"（一条批注 = 起止锚点 + 一个引用，只能算一处）。
    """
    if ids is None:
        ids = set()
    removed = 0
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for element in list(root.iter(qn(tag))):
            value = element.get(qn("w:id"))
            if value is not None:
                ids.add(value)
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                removed += 1
    for reference in list(root.iter(qn("w:commentReference"))):
        value = reference.get(qn("w:id"))
        if value is not None:
            ids.add(value)
        run = reference.getparent()
        if run is None:
            continue
        run.remove(reference)
        removed += 1
        # 只剩 rPr 的空 run 留在段落里没意义
        if run.tag == qn("w:r") and all(child.tag == qn("w:rPr") for child in list(run)):
            holder = run.getparent()
            if holder is not None:
                holder.remove(run)
    return removed


def _xml_parts(document):
    for part in document.part.package.iter_parts():
        element = getattr(part, "element", None)
        if element is not None:
            yield part, element


def _strip_comments_everywhere(document):
    """正文、页眉页脚、脚注里都可能挂批注锚点，一处不拔就等于没清干净。"""
    ids = set()
    removed = 0
    for _part, root in _xml_parts(document):
        removed += _strip_comment_marks(root, ids)
    return ids, removed


def _drop_comment_parts(document):
    """把批注相关部件的关系摘掉；python-docx 保存时只写有关系的部件，部件自然就没了。"""
    dropped = 0
    for rel_id, relationship in list(document.part.rels.items()):
        if relationship.is_external:
            continue
        if relationship.reltype in COMMENT_RELATIONSHIPS:
            document.part.drop_rel(rel_id)
            dropped += 1
    return dropped


RPR_FONT_SUCCESSORS = (
    "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike", "w:outline",
    "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
    "w:color", "w:spacing", "w:w", "w:kern", "w:position",
) + SZ_SUCCESSORS + ("w:eastAsianLayout", "w:specVanish", "w:oMath")


def _rpr_insert(rpr, tag_name, successors):
    """按 CT_RPr 的顺序把 tag_name 插进 rPr（w:rStyle 永远是第一个孩子）。"""
    element = rpr.find(qn(tag_name))
    if element is not None:
        return element
    element = rpr.makeelement(qn(tag_name), {})
    rpr.insert_element_before(element, *successors)
    return element


def _hyperlink_style_ids(document):
    """字符样式里哪些是"超链接样式"（Hyperlink / 超链接）。

    中文版 Word 的 styleId 可能不是 Hyperlink（而是 a3 之类），所以按 styles.xml 里的
    w:name 一起认；认不出来时至少认得 Hyperlink 这个标准 id。
    """
    ids = set(LINK_STYLE_IDS)
    try:
        root = document.styles.element
    except (AttributeError, KeyError, NotImplementedError):
        return ids
    for style in root.findall(qn("w:style")):
        style_id = (style.get(qn("w:styleId")) or "").strip()
        name_element = style.find(qn("w:name"))
        name = ((name_element.get(qn("w:val")) if name_element is not None else "") or "").strip()
        if style_id.lower() in LINK_STYLE_IDS or name.lower() in LINK_STYLE_NAMES:
            if style_id:
                ids.add(style_id.lower())
    return ids


def _strip_run_format(run_element, options, stats, link_style_ids=(), text_style=True):
    """去掉一个 run 上的网页痕迹；只有真的改到了才计数（docx_clean 的 changed 靠 stats 判定）。

    link_style_ids：超链接字符样式。蓝色和"带下划线"是那个样式给的，不摘掉的话
    w:color 删了文字还是蓝的（下划线按 keepUnderline 的口径另说）。
    text_style=False 用于页眉页脚：只去色，不改字体字号（学校抬头的版式别动）。
    """
    has_text = run_element.find(qn("w:t")) is not None
    rpr = run_element.find(qn("w:rPr"))

    def ensure_rpr():
        nonlocal rpr
        if rpr is None:
            rpr = run_element.makeelement(qn("w:rPr"), {})
            run_element.insert(0, rpr)
        return rpr

    if options.get("stripColor", True) and rpr is not None:
        removed = 0
        for tag in ("w:color", "w:highlight", "w:shd"):
            element = rpr.find(qn(tag))
            if element is not None:
                rpr.remove(element)
                removed += 1
        # 下划线是语义（英语"划线部分读音不同"、填空横线、语文"画出中心句"），默认保留；
        # 只有老师明确传 keepUnderline=false 才去掉，并且 warnings 里要说清后果（见 docx_clean）。
        if not options.get("keepUnderline", True):
            element = rpr.find(qn("w:u"))
            if element is not None:
                rpr.remove(element)
                stats["underlines"] += 1
        for style in list(rpr.findall(qn("w:rStyle"))):
            if (style.get(qn("w:val")) or "").strip().lower() in link_style_ids:
                rpr.remove(style)
                removed += 1
        stats["colors"] += removed

    if not text_style or not has_text:
        return
    font = options.get("font")
    if font:
        attributes = ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs")
        fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if fonts is None or any(fonts.get(qn(attribute)) != font for attribute in attributes):
            if fonts is None:
                fonts = _rpr_insert(ensure_rpr(), "w:rFonts", RPR_FONT_SUCCESSORS)
            for attribute in attributes:
                fonts.set(qn(attribute), font)
            stats["fonts"] += 1
    size = options.get("sizePt")
    if size:
        value = str(int(float(size) * 2))
        for tag, successors in (("w:sz", SZ_SUCCESSORS), ("w:szCs", SZCS_SUCCESSORS)):
            element = rpr.find(qn(tag)) if rpr is not None else None
            if element is not None and element.get(qn("w:val")) == value:
                continue
            if element is None:
                element = _rpr_insert(ensure_rpr(), tag, successors)
            element.set(qn("w:val"), value)
            stats["sizes"] += 1


def _paragraph_ppr(paragraph_element, ppr=None):
    """拿到（必要时补上）段落的 w:pPr —— 只有真的要写东西时才补，别给页眉段塞空 pPr。"""
    if ppr is None:
        ppr = paragraph_element.find(qn("w:pPr"))
    if ppr is None:
        ppr = paragraph_element.makeelement(qn("w:pPr"), {})
        paragraph_element.insert(0, ppr)
    return ppr


def _normalize_spacing(paragraph_element, stats, keep=False, body=True):
    """统一段距，两种口径。

    keep=True（默认，为的是"别比原件更挤"）：
      - 原有的段前/段后间距原样保留，不清零；
      - 行距只把"单倍/不到 1.5 倍"或"固定磅值(exact/atLeast)"的改成 1.5 倍（360 twips, auto）；
      - 正文段落压根没写段距时补 6 磅段后（120 twips），删掉空段后相邻题目不贴死。
    keep=False：旧口径，段前段后一律清零（老师明确要求"排紧一点"时用）。
    """
    ppr = paragraph_element.find(qn("w:pPr"))
    spacing = ppr.find(qn("w:spacing")) if ppr is not None else None
    if spacing is None:
        if not keep:
            spacing = _ppr_insert(_paragraph_ppr(paragraph_element, ppr), "w:spacing",
                                  SPACING_SUCCESSORS)
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            stats["spacing"] += 1
        elif body:
            spacing = _ppr_insert(_paragraph_ppr(paragraph_element, ppr), "w:spacing",
                                  SPACING_SUCCESSORS)
            spacing.set(qn("w:after"), "120")
            stats["spacing"] += 1
        return
    if not keep:
        modified = False
        for attribute in ("w:before", "w:after"):
            if spacing.get(qn(attribute)) != "0":
                spacing.set(qn(attribute), "0")
                modified = True
        if modified:
            stats["spacing"] += 1
        return
    rule = spacing.get(qn("w:lineRule"))
    try:
        value = int(spacing.get(qn("w:line")))
    except (TypeError, ValueError):
        value = None
    if rule not in (None, "auto"):
        # 固定磅值的行距会把汉字切头切脚，统一成 1.5 倍
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")
        stats["spacing"] += 1
    elif value is not None and value < 360:
        spacing.set(qn("w:line"), "360")
        stats["spacing"] += 1


def _paragraph_is_empty(paragraph_element):
    if _paragraph_text(paragraph_element):
        return False
    for tag in KEEP_TAGS:
        if paragraph_element.find(".//" + qn(tag)) is not None:
            return False
    return True


def _collapse_empty_paragraphs(body, stats, strip_all=False):
    """空段处理：默认只把"连续空段"压成 1 个（留出呼吸感），strip_all=True 才全删。

    只动正文（body 的直接子段落）；页眉页脚里的空段删掉会让章名贴边。
    """
    removed = 0
    run_length = 0
    for paragraph in list(body):
        if paragraph.tag != qn("w:p"):
            run_length = 0
            continue
        if not _paragraph_is_empty(paragraph):
            run_length = 0
            continue
        run_length += 1
        if strip_all or run_length > 1:
            body.remove(paragraph)
            removed += 1
    stats["emptyParagraphs"] += removed
    return removed


def _same_path(left, right):
    return os.path.normcase(os.path.abspath(str(left))) == os.path.normcase(os.path.abspath(str(right)))


def _default_clean_output(path):
    """净化默认另存：同目录 <原名>_净化后.docx —— 原件必须还在，老师要能对得上。"""
    folder, filename = os.path.split(os.path.abspath(str(path)))
    stem, extension = os.path.splitext(filename)
    return os.path.join(folder, stem + "_净化后" + (extension or ".docx"))


def _backup_path(path):
    """原地覆盖前的备份名：<原名>_原件备份.docx（重名就 _原件备份2、3…）。"""
    folder, filename = os.path.split(os.path.abspath(str(path)))
    stem, extension = os.path.splitext(filename)
    extension = extension or ".docx"
    candidate = os.path.join(folder, stem + "_原件备份" + extension)
    counter = 2
    while os.path.exists(candidate):
        candidate = os.path.join(folder, stem + "_原件备份" + str(counter) + extension)
        counter += 1
    return candidate


def _require_docx(path, action="处理"):
    """文件得存在、得是 .docx；出错要说人话（老师看的是这句提示）。

    真实撞车场景：① 路径只写了一半（"第三单元.docx"）② 拖进来的是 PDF
    ③ 学校电脑上存的是 .doc。
    """
    text = str(path)
    require(os.path.exists(text), "file-not-found", "找不到这个文件：" + text,
            "路径要写完整（例如 D:/教学/第三单元.docx），也可以点「浏览…」挑")
    lowered = text.lower()
    if lowered.endswith(".docx"):
        return
    ext = os.path.splitext(lowered)[1]
    if ext == ".pdf":
        raise EngineError(
            "unsupported-source", "这是 PDF，我只能" + action + " .docx 文件",
            "用 Word/WPS 打开它，「另存为」选 .docx 再试一次；PDF 里的文字我不保证取得全")
    if ext in (".doc", ".wps", ".dot"):
        raise EngineError(
            "unsupported-source", "这是旧格式（" + ext + "），先另存为 .docx 再来",
            "在 Word/WPS 里「另存为」→ 文件类型选「Word 文档 (*.docx)」")
    raise EngineError(
        "unsupported-source", "只能" + action + " .docx 文件（这个是 " + (ext or "没扩展名") + "）",
        "支持 .docx；旧格式先在 Word/WPS 里另存为 .docx")


def docx_clean(payload):
    """payload = {path | inputPath, outputPath?, overwrite?, options?{font, sizePt, stripColor,
    stripLinks, stripComments, keepUnderline, keepSpacing, normalizeSpacing, removeEmptyParagraphs,
    stripEmpty, cleanHeader, cleanFooter}}

    输出口径（P0-5）：
      - 默认另存同目录「<原名>_净化后.docx」，原件一个字不动；note 里写明「原文件已保留：…」。
      - 只有显式 overwrite:true 才原地写；原地写之前先备份成「<原名>_原件备份.docx」，回执带 backup。
      - outputPath 写的就是原文件、又没给 overwrite:true → 报 overwrite-required，原件一个字不动。
    changed 说的是「内容真的变了没有」（不是"输出路径和输入路径不同"，那个 false 会骗人）。

    默认不把版面弄得更挤（P1-9）：
      - keepSpacing=true：段前/段后原样保留，只把行距统一到 1.5 倍；正文没写段距的补 6 磅段后。
      - removeEmptyParagraphs=true：只把「连续空段」压成 1 个（stripEmpty=true 才全删）。
      - cleanHeader/cleanFooter=true：清页眉页脚里的网站广告行；学校的抬头、校名、页码照留。
    """
    path = payload.get("path") or payload.get("inputPath")
    require(bool(path), "path-required", "没有指定要清理的文件", "点「浏览…」选一份 .docx，或把路径粘进输入框")
    _require_docx(path, "清理")
    source = os.path.abspath(str(path))
    raw_options = dict(payload.get("options") or {})
    overwrite = bool(payload.get("overwrite") or payload.get("inPlace") or raw_options.get("inPlace"))
    requested = payload.get("outputPath")
    if requested:
        output = os.path.abspath(str(requested))
        in_place = _same_path(output, source)
        if in_place and not overwrite:
            raise EngineError(
                "overwrite-required",
                "输出的就是原文件，这么写会把原件换掉：" + os.path.basename(source),
                "把输出文件名改一下（默认就是「" + os.path.basename(_default_clean_output(source)) +
                "」），或者加 overwrite:true —— 那样我会先把原件备份成「" +
                os.path.basename(_backup_path(source)) + "」再写")
    elif overwrite:
        output = source
        in_place = True
    else:
        output = _default_clean_output(source)
        in_place = False
    # 输出指向别的已有文件时以前直接覆盖（实测把 36 字节的哨兵换成了 36KB），
    # 这里跟原地写一个口径：没有 overwrite:true 就不动别人的文件。
    if not in_place and not overwrite and os.path.exists(output):
        raise EngineError(
            "output-exists",
            "这个位置已经有同名文件了，我不动它：" + os.path.basename(output),
            "换个输出文件名，或者加 overwrite:true 明确允许覆盖")

    options = raw_options
    options.setdefault("font", "宋体")
    options.setdefault("stripColor", True)
    options.setdefault("stripLinks", True)
    options.setdefault("stripComments", True)
    options.setdefault("keepUnderline", True)
    options.setdefault("keepSpacing", True)
    options.setdefault("normalizeSpacing", True)
    options.setdefault("removeEmptyParagraphs", True)
    options.setdefault("stripEmpty", False)

    document = _open_document(path, "文件")
    stats = {"links": 0, "colors": 0, "underlines": 0, "comments": 0, "ads": 0,
             "fonts": 0, "sizes": 0, "spacing": 0, "emptyParagraphs": 0}
    comment_ids = set()
    comment_parts = 0
    if options.get("stripComments", True):
        comment_ids, _marks = _strip_comments_everywhere(document)
        comment_parts = _drop_comment_parts(document)
        stats["comments"] = len(comment_ids)

    header_parts = list(_header_footer_parts(document))
    if options.get("stripLinks", True):
        _unwrap_hyperlinks(document.element.body, stats)
        # 页眉里的网址也是链接，一样拆掉（文字留下）
        for _part, root in header_parts:
            _unwrap_hyperlinks(root, stats)
    # 页眉页脚里那句「www.xxjyw.com 免费下载 请勿商用」——下载站塞进去的，
    # 印在卷子上等于替它做广告。带校名/年级/「第 N 页」的一律不动。
    # 老参数名 stripAds / headerAds 一律认，cleanHeader / cleanFooter 分开管。
    fallback = options.get("headerAds", options.get("stripAds", True))
    clean_header = bool(options.get("cleanHeader", fallback))
    clean_footer = bool(options.get("cleanFooter", fallback))
    for _part, root in header_parts:
        if root.tag == qn("w:ftr"):
            if not clean_footer:
                continue
        elif not clean_header:
            continue
        stats["ads"] += _strip_ad_paragraphs(root)

    # 正文和页眉页脚都去色（页眉里的红字一样碍眼），但字体字号只按正文口径调，
    # 免得把学校抬头的版式改样。
    link_style_ids = _hyperlink_style_ids(document)
    for run_element in document.element.body.iter(qn("w:r")):
        _strip_run_format(run_element, options, stats, link_style_ids, text_style=True)
    for _part, root in header_parts:
        for run_element in root.iter(qn("w:r")):
            _strip_run_format(run_element, options, stats, link_style_ids, text_style=False)

    if options.get("normalizeSpacing", True):
        # keep 默认 True：保留老师原本的段距，别净化完比原件还挤（persona P1-9：段距全归零）
        keep_spacing = bool(options.get("keepSpacing", True))
        for paragraph_element, in_body in _iter_paragraphs(document):
            _normalize_spacing(paragraph_element, stats, keep=keep_spacing, body=in_body)
    if options.get("removeEmptyParagraphs", True):
        # 只动正文：页眉页脚的空段删掉会让章名贴边
        _collapse_empty_paragraphs(document.element.body, stats,
                                   strip_all=bool(options.get("stripEmpty")))

    backup = None
    if in_place:
        backup = _backup_path(source)
        shutil.copy2(source, backup)
    _ensure_parent(output)
    document.save(output)

    warnings = []
    if stats["links"]:
        warnings.append({"code": "links-removed",
                         "message": "拆掉了 " + str(stats["links"]) + " 个超链接（文字留下、跳转去掉）"})
    if stats["ads"]:
        warnings.append({"code": "header-ad-removed",
                         "message": "页眉页脚里去掉了 " + str(stats["ads"]) + " 行网站广告"
                                    "（「www.…」「免费下载 请勿商用」这类）；校名、年级、页码没动"})
    if stats["comments"]:
        warnings.append({"code": "comments-removed",
                         "message": "删掉了 " + str(stats["comments"]) + " 处批注："
                                    "备课批注（「答案改成 B」「这题去年错一半」）不会跟着文件发给学生/家长了，正文内容不受影响"})
    elif comment_parts:
        warnings.append({"code": "comments-removed",
                         "message": "清掉了文件里的批注部件：原来那些批注不会跟着文件发出去，正文内容不受影响"})
    if stats["underlines"]:
        warnings.append({"code": "underline-removed",
                         "message": "去掉了 " + str(stats["underlines"]) + " 处下划线：英语「划线部分读音不同」、"
                                    "填空横线、语文「画出中心句」这类标注会看不出来；"
                                    "下划线默认是保留的（这次传了 keepUnderline=false），下次不传这个选项就行，"
                                    "或者净化前先留一份副本"})
    if in_place:
        warnings.append({"code": "original-backed-up",
                         "message": "原件已经备份成「" + os.path.basename(backup) + "」，"
                                    "净化后的内容写回了原文件"})

    # changed 说的是「内容真的变了没有」（早先它算的是"输出路径和输入路径不同"，
    # 原地覆盖时永远是 false，老师以为净化没生效）。stats 里每项都是"真的改到了才加一"。
    changed = any(value for value in stats.values()) or bool(comment_parts)
    if in_place:
        note = "已原地净化，原件备份在：" + os.path.basename(backup)
    else:
        note = "原文件已保留：" + os.path.basename(source) + "；净化后的副本：" + os.path.basename(output)
    return {
        "outputPath": os.path.abspath(output),
        "originalPath": source,
        "changed": changed,
        "overwritten": in_place,
        "backup": os.path.abspath(backup) if backup else None,
        "stats": stats,
        "warnings": warnings,
        "note": note,
    }


def _part_with_unique_name(target, part, used_names):
    """图片/媒体部件搬进目标包时要改名，否则两份文档的 image1.png 会撞成重名条目（包就坏了）。"""
    name = str(part.partname)
    if name not in used_names:
        used_names.add(name)
        return part
    folder, filename = posixpath.split(name)
    stem, dot, suffix = filename.rpartition(".")
    if not dot:
        stem, suffix = filename, ""
    counter = 2
    while True:
        candidate = stem + "_" + str(counter) + ("." + suffix if suffix else "")
        full = posixpath.join(folder, candidate) if folder else candidate
        if full not in used_names:
            break
        counter += 1
    used_names.add(full)
    return Part(PackURI(full), part.content_type, part.blob, target.part.package)


def _copy_media(source, target, subtree, used_names):
    """把复制过来的图片关系重新挂到目标文档上，否则合并后是"无法显示的红叉"。"""
    copied = 0
    for element in subtree.iter():
        if element.tag not in (TAG_BLIP, TAG_VML_IMAGE):
            continue
        attribute = qn("r:embed") if element.tag == TAG_BLIP else qn("r:id")
        rel_id = element.get(attribute)
        if not rel_id:
            continue
        try:
            relationship = source.part.rels[rel_id]
        except KeyError:
            continue
        if relationship.is_external:
            continue
        part = _part_with_unique_name(target, relationship.target_part, used_names)
        new_id = target.part.relate_to(part, relationship.reltype)
        element.set(attribute, new_id)
        copied += 1
    return copied


def _relink_hyperlinks(source, target, root):
    """超链接要么把关系重建到目标包上，要么拆成普通文字——绝不能留一个源包里的 rId。

    源包里的 rIdLink9 搬到目标包里没人认识，Word 会判"文件已损坏"。
    """
    rebuilt = 0
    flattened = 0
    for hyperlink in list(root.iter(qn("w:hyperlink"))):
        rel_id = hyperlink.get(qn("r:id"))
        if not rel_id:
            continue  # 只有 w:anchor 的站内跳转不需要关系
        new_id = None
        try:
            relationship = source.part.rels[rel_id]
        except KeyError:
            relationship = None
        if relationship is not None and relationship.is_external and relationship.reltype == RT_HYPERLINK:
            new_id = target.part.relate_to(relationship.target_ref, RT_HYPERLINK, is_external=True)
        if new_id:
            hyperlink.set(qn("r:id"), new_id)
            rebuilt += 1
        elif _unwrap_hyperlink(hyperlink):
            flattened += 1
    return rebuilt, flattened


def _strip_note_references(root):
    """合并是"拼正文"，目标包里没有脚注/尾注部件，留一个引用就是指向空处的坏引用。"""
    removed = 0
    for tag in ("w:footnoteReference", "w:endnoteReference"):
        for reference in list(root.iter(qn(tag))):
            parent = reference.getparent()
            if parent is None:
                continue
            parent.remove(reference)
            removed += 1
            if parent.tag == qn("w:r") and all(child.tag == qn("w:rPr") for child in list(parent)):
                holder = parent.getparent()
                if holder is not None:
                    holder.remove(parent)
    return removed


def _remap_numbering(root, mapping):
    """把复制内容里的 w:numId 换成目标包里的新编号；换不到说明源里就没有定义，摘掉 numPr。"""
    dropped = 0
    for num_id in list(root.iter(qn("w:numId"))):
        value = num_id.get(qn("w:val"))
        if value in mapping:
            num_id.set(qn("w:val"), mapping[value])
            continue
        num_pr = num_id.getparent()
        if num_pr is not None and num_pr.tag == qn("w:numPr"):
            holder = num_pr.getparent()
            if holder is not None:
                holder.remove(num_pr)
                dropped += 1
    return dropped


def _add_abstract_num(target_root, element):
    first_num = target_root.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(element)
    else:
        target_root.append(element)


def _add_num(target_root, element):
    cleanup = target_root.find(qn("w:numIdMacAtCleanup"))
    if cleanup is not None:
        cleanup.addprevious(element)
    else:
        target_root.append(element)


def _prepare_numbering(source, target):
    """把源的 numbering 定义按新 id 搬进目标，返回 {源 numId: 目标 numId}。

    自动编号题的题号（Word 的 1. 2. 3. 列表）全靠 numbering.xml 里的定义；
    只搬正文不搬定义，合并后题号就整片消失。
    """
    mapping = {}
    try:
        source_root = source.part.numbering_part.element
        target_root = target.part.numbering_part.element
    except (KeyError, AttributeError, NotImplementedError):
        return mapping
    used_abstract = set()
    for element in target_root.findall(qn("w:abstractNum")):
        try:
            used_abstract.add(int(element.get(qn("w:abstractNumId")) or 0))
        except (TypeError, ValueError):
            continue
    used_num = set()
    for element in target_root.findall(qn("w:num")):
        try:
            used_num.add(int(element.get(qn("w:numId")) or 0))
        except (TypeError, ValueError):
            continue
    abstract_map = {}
    next_abstract = max(used_abstract) + 1 if used_abstract else 1
    for abstract in source_root.findall(qn("w:abstractNum")):
        old_id = abstract.get(qn("w:abstractNumId"))
        if old_id is None:
            continue
        new_id = next_abstract
        next_abstract += 1
        used_abstract.add(new_id)
        abstract_map[old_id] = str(new_id)
        copied = deepcopy(abstract)
        copied.set(qn("w:abstractNumId"), str(new_id))
        # 图片项目符号（w:numPicBullet）不跟着搬，留着 w:lvlPicBulletId 就是指向空处的坏引用
        for bullet_id in list(copied.iter(qn("w:lvlPicBulletId"))):
            holder = bullet_id.getparent()
            if holder is not None:
                holder.remove(bullet_id)
        _add_abstract_num(target_root, copied)
    for num in source_root.findall(qn("w:num")):
        old_id = num.get(qn("w:numId"))
        if old_id is None:
            continue
        abstract_id = num.find(qn("w:abstractNumId"))
        if abstract_id is None or abstract_id.get(qn("w:val")) not in abstract_map:
            continue  # 定义不全的编号不搬，正文里对应的 numPr 会被摘掉
        copied = deepcopy(num)
        new_id = max(used_num) + 1 if used_num else 1
        used_num.add(new_id)
        copied.set(qn("w:numId"), str(new_id))
        abstract_id = copied.find(qn("w:abstractNumId"))
        abstract_id.set(qn("w:val"), abstract_map[abstract_id.get(qn("w:val"))])
        _add_num(target_root, copied)
        mapping[old_id] = str(new_id)
    return mapping


def _used_style_ids(root):
    ids = set()
    for element in root.iter():
        if element.tag in (qn("w:pStyle"), qn("w:rStyle"), qn("w:tblStyle"), qn("w:tableStyle")):
            value = element.get(qn("w:val"))
            if value:
                ids.add(value)
    return ids


def _merge_styles(source, target, style_ids, numbering_map):
    """源文档用到的样式定义按需搬进目标（含 basedOn 链），否则合并后标题/列表全变正文样式。"""
    report = {"imported": 0, "missing": []}
    if not style_ids:
        return report
    try:
        source_root = source.styles.element
        target_root = target.styles.element
    except (AttributeError, KeyError):
        report["missing"] = sorted(style_ids)
        return report
    source_styles = {}
    for element in source_root.findall(qn("w:style")):
        style_id = element.get(qn("w:styleId"))
        if style_id:
            source_styles[style_id] = element
    target_styles = {}
    for element in target_root.findall(qn("w:style")):
        style_id = element.get(qn("w:styleId"))
        if style_id:
            target_styles[style_id] = element
    pending = list(style_ids)
    seen = set()
    while pending:
        style_id = pending.pop()
        if not style_id or style_id in seen:
            continue
        seen.add(style_id)
        source_element = source_styles.get(style_id)
        if source_element is None:
            report["missing"].append(style_id)
            continue
        if style_id in target_styles:
            continue  # 目标里已有同名样式（内置样式），沿用目标定义
        copied = deepcopy(source_element)
        _remap_numbering(copied, numbering_map)
        target_root.append(copied)
        target_styles[style_id] = copied
        report["imported"] += 1
        for link_tag in ("w:basedOn", "w:next", "w:link"):
            linked = copied.find(qn(link_tag))
            value = linked.get(qn("w:val")) if linked is not None else None
            if value:
                pending.append(value)
    return report


def _containing(target_element, tags):
    for tag in tags:
        candidate = target_element
        while candidate is not None:
            if candidate.tag == qn(tag):
                return candidate
            candidate = candidate.getparent()
    return None


def _sanitize_rids(root, part, tags):
    """兜底：正文里任何一个 rId 只要目标包不认，就整块摘掉（或拆成文字），绝不带出去。"""
    known = {relationship.rId for relationship in part.rels.values()}
    removed = 0
    for element in list(root.iter()):
        if element.getparent() is None:
            continue
        unresolved = [key for key, value in element.attrib.items()
                      if key.startswith("{" + R_NS + "}") and value not in known]
        if not unresolved:
            continue
        if element.tag == qn("w:hyperlink"):
            if _unwrap_hyperlink(element):
                removed += 1
            continue
        container = _containing(element, tags) or element
        holder = container.getparent()
        if holder is not None:
            holder.remove(container)
            removed += 1
    return removed


def _page_break_before(paragraph_element):
    """在新一份的第一段上写 pageBreakBefore：分页但不多出一个空行。"""
    ppr = paragraph_element.find(qn("w:pPr"))
    if ppr is None:
        ppr = paragraph_element.makeelement(qn("w:pPr"), {})
        paragraph_element.insert(0, ppr)
    _ppr_insert(ppr, "w:pageBreakBefore", PAGE_BREAK_SUCCESSORS)


def _page_break_paragraph(target):
    """第一份内容不是段落（比如直接以表格开头）时，用一个只含分页符的段落分页。"""
    paragraph = target.element.body.makeelement(qn("w:p"), {})
    run = paragraph.makeelement(qn("w:r"), {})
    br = run.makeelement(qn("w:br"), {})
    br.set(qn("w:type"), "page")
    run.append(br)
    paragraph.append(run)
    return paragraph


def _append_body(target, element, sect_pr):
    if sect_pr is not None and sect_pr.getparent() is not None:
        sect_pr.addprevious(element)
    else:
        target.element.body.append(element)


HEADER_FOOTER_NAMES = ("header", "first_page_header", "even_page_header",
                       "footer", "first_page_footer", "even_page_footer")


def _container_has_text(container):
    return bool("".join(node.text or "" for node in container.iter(qn("w:t"))).strip())


def _inherit_header_footer(source, target):
    """把第一份文件的页眉页脚搬到合并稿上：搬文字、表格、页码域，图片不带（rId 会悬空）。"""
    copied = 0
    if not source.sections:
        return copied
    source_section = source.sections[0]
    target_section = target.sections[0]
    for name in HEADER_FOOTER_NAMES:
        try:
            src = getattr(source_section, name)
            dst = getattr(target_section, name)
        except Exception:
            continue
        try:
            if src.is_linked_to_previous:
                continue
            container = src._element
            if container is None or not _container_has_text(container):
                continue
            dst.is_linked_to_previous = False
            dst_element = dst._element
            for child in list(dst_element):
                if child.tag == qn("w:p"):
                    dst_element.remove(child)
            for child in container:
                if child.tag == qn("w:sectPr"):
                    continue
                copy = deepcopy(child)
                _sanitize_rids(copy, target.part, ("w:drawing", "w:object", "w:pict", "w:altChunk"))
                dst_element.append(copy)
            copied += 1
        except Exception:
            continue
    return copied


# 「页眉页脚怎么带」三档：first 沿用第一份（老行为）/ each 每份各带自己的 / none 一个都不带
HEADER_MODE_FIRST = "first"
HEADER_MODE_EACH = "each"
HEADER_MODE_NONE = "none"
EACH_MODE_ALIASES = ("each", "per-source", "per_source", "perfile", "per-file", "every", "section",
                     "逐份", "各自", "每份")
NONE_MODE_ALIASES = ("none", "off", "no", "false", "0", "不带", "无", "不要")
# 页眉页脚里带 rId 的容器：图片/嵌入对象/表单。搬到别的包里关系接不上，就是悬空引用
HEADER_IMAGE_CONTAINERS = ("w:drawing", "w:object", "w:pict", "w:altChunk")
# 认得出来的页码类域。别的域（REF / TOC / DATE）原样留着，不碰
PAGE_FIELD_NAMES = ("PAGE", "NUMPAGES", "SECTIONPAGES", "SECTION")


def _merge_header_mode(payload):
    """解析「页眉页脚怎么带」：mergeHeaderMode 优先，其次 mergeHeaders（默认 True）。

    老参数 mergeHeaders 的语义一个字不能变：true 或不传 = 沿用第一份；
    只有显式写 'each'（或「逐份」「每份」这类等价说法）才给每份文件套各自的抬头。
    """
    value = payload.get("mergeHeaderMode")
    if value is None:
        value = payload.get("mergeHeaders", True)
    if isinstance(value, bool):
        return HEADER_MODE_FIRST if value else HEADER_MODE_NONE
    if value is None:
        return HEADER_MODE_FIRST
    text = str(value).strip().lower()
    if text in EACH_MODE_ALIASES:
        return HEADER_MODE_EACH
    if text in NONE_MODE_ALIASES:
        return HEADER_MODE_NONE
    return HEADER_MODE_FIRST


def _field_tokens(paragraph):
    """按文档顺序把段落里的域记号列出来：[(run, 'begin'|'separate'|'end'|'instr', 文本)]。"""
    tokens = []
    for run in paragraph.iter(qn("w:r")):
        for node in run:
            if node.tag == qn("w:fldChar"):
                kind = (node.get(qn("w:fldCharType")) or "").strip()
                if kind in ("begin", "separate", "end"):
                    tokens.append((run, kind, ""))
            elif node.tag == qn("w:instrText"):
                tokens.append((run, "instr", node.text or ""))
    return tokens


def _page_field_name(instruction):
    r"""域代码里第一个词就是域名：` PAGE  * MERGEFORMAT ` → PAGE。"""
    text = (instruction or "").replace("\\", " ").strip()
    if not text:
        return ""
    return text.split()[0].upper()


def _run_is_empty(run):
    """只剩 rPr 的 run 没有意义，删掉——否则页脚里会攒一串空 run。"""
    return all(child.tag == qn("w:rPr") for child in list(run))


def _strip_field_markers(run):
    for node in list(run):
        if node.tag in (qn("w:fldChar"), qn("w:instrText")):
            run.remove(node)


def _field_result_run(paragraph, text):
    run = paragraph.makeelement(qn("w:r"), {})
    node = paragraph.makeelement(qn("w:t"), {})
    node.text = text
    run.append(node)
    return run


def _simplify_complex_fields(paragraph):
    """把 Word 造的复杂域（begin / instrText / separate / 结果 / end 散在好几个 run 里）
    压成一个 w:fldSimple —— 只有一个标签，跨文档搬最不容易散架。

    页码类的域单独认：合并时源页脚整段搬过来，复杂域里的记号在别的包里容易配不上，
    Word 就显示「错误！未定义书签。」；w:fldSimple 没这个问题。
    非页码域（REF / TOC / DATE）一律不动，那是正文自己的事。
    """
    tokens = _field_tokens(paragraph)
    index = 0
    while index < len(tokens):
        if tokens[index][1] != "begin":
            index += 1
            continue
        depth = 0
        end_at = None
        separate_at = None
        instruction = ""
        for probe in range(index, len(tokens)):
            probe_kind = tokens[probe][1]
            if probe_kind == "begin":
                depth += 1
            elif probe_kind == "end":
                depth -= 1
                if depth == 0:
                    end_at = probe
                    break
            elif depth == 1:
                if probe_kind == "separate":
                    if separate_at is None:
                        separate_at = probe
                elif probe_kind == "instr" and separate_at is None:
                    instruction += tokens[probe][2]
        if end_at is None:
            # 配不平的域原样留着：宁可页码显示旧值，也不能把用户正文切坏
            index += 1
            continue
        name = _page_field_name(instruction)
        if name not in PAGE_FIELD_NAMES:
            index = end_at + 1
            continue
        span = []
        for probe in range(index, end_at + 1):
            holder = tokens[probe][0]
            if holder not in span:
                span.append(holder)
        for holder in span:
            _strip_field_markers(holder)
        field = paragraph.makeelement(qn("w:fldSimple"), {})
        field.set(qn("w:instr"), " " + name + " ")
        result_from = separate_at + 1 if separate_at is not None else end_at
        kept = False
        for probe in range(result_from, end_at):
            holder = tokens[probe][0]
            if holder.getparent() is not paragraph or holder not in span:
                continue
            if holder.find(qn("w:t")) is None:
                continue
            # 上一次算出来的页码留着当缓存值：Word 打开先显示它，按 F9 再更新
            paragraph.remove(holder)
            field.append(holder)
            kept = True
        anchor = span[0]
        parent = anchor.getparent()
        if parent is None:
            index = end_at + 1
            continue
        anchor.addprevious(field)
        for holder in span:
            if holder.getparent() is parent and _run_is_empty(holder):
                parent.remove(holder)
        if not kept:
            # 一个结果都没留下时补个「1」：w:fldSimple 里空着，Word 有时只显示空白
            field.append(_field_result_run(paragraph, "1"))
        index = end_at + 1


def _normalize_page_fields(root):
    """页眉页脚里所有段落都过一遍域整理（容器里可能不止一个段落）。"""
    for paragraph in list(root.iter(qn("w:p"))):
        _simplify_complex_fields(paragraph)


def _container_carries_content(container):
    """这个页眉/页脚容器里有没有值得印出来的东西：文字、表格、页码域都算。"""
    if container is None:
        return False
    if _container_has_text(container):
        return True
    if container.find(qn("w:tbl")) is not None:
        return True
    for tag in ("w:fldSimple", "w:fldChar"):
        if container.find(".//" + qn(tag)) is not None:
            return True
    return False


def _fill_header_footer_container(dst_element, container, part):
    """把一份源文件的页眉（或页脚）内容搬进目标节：文字、表格、页码域带，图片不带。

    part 必须是页眉/页脚部件自己的 HeaderPart/FooterPart —— 图片关系挂在那个部件上，
    传 target.part 会把「这个 rId 还在」判成真，搬过去就是悬空引用。
    """
    children = []
    for child in container:
        if child.tag == qn("w:sectPr"):
            continue
        copy = deepcopy(child)
        # 图片、嵌入对象整块丢掉：它们的关系在源文件里，搬过来接不上
        _sanitize_rids(copy, part, HEADER_IMAGE_CONTAINERS)
        _normalize_page_fields(copy)
        children.append(copy)
    for child in list(dst_element):
        dst_element.remove(child)
    if not children:
        dst_element.append(dst_element.makeelement(qn("w:p"), {}))
        return False
    for child in children:
        dst_element.append(child)
    # w:hdr / w:ftr 至少要有一个块级元素，否则 Word 说「文件需要修复」
    if dst_element.find(qn("w:p")) is None and dst_element.find(qn("w:tbl")) is None:
        dst_element.append(dst_element.makeelement(qn("w:p"), {}))
    return _container_carries_content(dst_element)


def _has_title_page(section):
    try:
        return section._sectPr.find(qn("w:titlePg")) is not None
    except Exception:
        return False


def _copy_title_page(source_section, target_section):
    """源文件设了「首页不同」时把 w:titlePg 也带过来，否则首页那套抬头不生效。"""
    if not _has_title_page(source_section):
        return False
    try:
        target_pr = target_section._sectPr
    except Exception:
        return False
    title = source_section._sectPr.find(qn("w:titlePg"))
    if target_pr.find(qn("w:titlePg")) is not None:
        return True
    # CT_SectPr 的子元素也有顺序，插错位置 Word 会提示「需要修复」
    target_pr.insert_element_before(deepcopy(title), "w:textDirection", "w:bidi", "w:rtlGutter",
                                    "w:docGrid", "w:printerSettings", "w:sectPrChange")
    return True


def _apply_section_header_footer(source, target):
    """把当前这一节的抬头/页脚换成 source 自己的那一套，返回这一节是否真带出了内容。

    缺抬头的源文件也照样把 is_linked_to_previous 置 False：不置的话这一节会继承上一份，
    等于给 B 学校的卷子印了 A 学校的抬头。所以「没有」也要显式写成空白。
    """
    if not source.sections or not target.sections:
        return False
    source_section = source.sections[0]
    target_section = target.sections[-1]
    title_page = _copy_title_page(source_section, target_section)
    carried = 0
    for name in HEADER_FOOTER_NAMES:
        try:
            src = getattr(source_section, name)
            dst = getattr(target_section, name)
        except Exception:
            continue
        try:
            blank = name in ("header", "footer")
            if src.is_linked_to_previous:
                if blank:
                    dst.is_linked_to_previous = False
                continue
            container = src._element
            if container is None or not _container_carries_content(container):
                if blank:
                    dst.is_linked_to_previous = False
                continue
            dst.is_linked_to_previous = False
            dst_element = dst._element
            if dst_element is None:
                continue
            if _fill_header_footer_container(dst_element, container, dst.part):
                carried += 1
        except Exception:
            continue
    if title_page:
        # 带了「首页不同」，首页那套也要显式定下来，不然这一节首页会继承上一份
        for name in ("first_page_header", "first_page_footer"):
            try:
                getattr(target_section, name).is_linked_to_previous = False
            except Exception:
                continue
    return carried > 0


def _copy_page_setup(source, target):
    """把首份文件的纸张/方向/四边距抄到合并稿的每一节。

    以前合并稿从 docx.Document() 的默认模板起手（Letter 215.9×279.4、边距 25.4/31.8），
    两份 A4 拼完变成 Letter，打印时被缩放（实测）。
    """
    for attr in ("page_width", "page_height", "orientation", "left_margin", "right_margin",
                 "top_margin", "bottom_margin", "header_distance", "footer_distance", "gutter"):
        try:
            setattr(target, attr, getattr(source, attr))
        except Exception:
            pass


def docx_merge(payload):
    """payload = {inputs:[path...] | paths:[path...] | paths:多行字符串, outputPath?, title?, overwrite?,
    pageBreak?, mergeHeaders? | mergeHeaderMode?, page?}

    参数名同时认 inputs 和 paths（老说明书上写的是 paths，两种都给，免得老师照着文档写却被拒）。

    页眉页脚两种口径（mergeHeaders / mergeHeaderMode 都认，后者优先）：
      - 默认 true / "first"：沿用第一份文件的抬头，其余文件的带不过来（老行为，一个字没改）；
      - "each"：每份文件的正文各占一节（分页符换成节分隔符），每节套各自文件自己的
        页眉页脚文字、表格和页码。图片抬头带不过来（关系在源包里，搬过来就是悬空引用），
        那一节留空；页码按合并稿连续编号，不在每份文件处重新从 1 数。
      - false / "none"：一份抬头都不带。
    """
    raw = payload.get("inputs")
    if not raw:
        raw = payload.get("paths")
    if isinstance(raw, str):
        raw = [line.strip().strip('"').strip("'") for line in raw.replace("\r\n", "\n").split("\n")]
    inputs = [item for item in (raw or []) if str(item).strip()]
    require(len(inputs) > 0, "inputs-required", "没有收到要合并的文件",
            "把要合并的 .docx 一行一个列出来（顺序就是合并顺序），或点「浏览…」逐个加")
    output = payload.get("outputPath") or default_output("文件加工", "合并结果", ".docx")
    for path in inputs:
        _require_docx(path, "合并")
    page_break = payload.get("pageBreak", True)
    header_mode = _merge_header_mode(payload)

    target = docx.Document()
    # 默认空文档自带一个空段和一个 sectPr，空段要去掉（否则合并后的首页顶上多一行）
    body = target.element.body
    for paragraph in list(body.iter(qn("w:p"))):
        body.remove(paragraph)
    sect_pr = body.find(qn("w:sectPr"))

    images = 0
    sources = []
    source_headers = 0
    comments = 0
    notes_dropped = 0
    links_flattened = 0
    numbering_dropped = 0
    objects_dropped = 0
    styles_imported = 0
    pages = 0
    header_sections = 0
    comment_ids = set()
    used_part_names = {str(part.partname) for part in target.part.package.iter_parts()}
    drop_containers = ("w:drawing", "w:object", "w:pict", "w:altChunk")
    first_document = None
    for index, path in enumerate(inputs):
        source = _open_document(path, "文件")
        if index == 0:
            first_document = source
        sources.append(os.path.basename(path))
        if index > 0 and header_mode == HEADER_MODE_EACH:
            # 每份文件各占一节：节分隔符自带下一页，后面不再加 pageBreakBefore，
            # 否则每份文件前面会多出一张空白页
            target.add_section(WD_SECTION.NEW_PAGE)
            pages += 1
        for section in source.sections:
            try:
                if not section.header.is_linked_to_previous and section.header.paragraphs[0].text.strip():
                    source_headers += 1
                    break
            except Exception:
                continue
        numbering_map = _prepare_numbering(source, target)
        styles_imported += _merge_styles(source, target, _used_style_ids(source.element.body),
                                         numbering_map)["imported"]
        first_in_document = True
        for element in list(source.element.body):
            if element.tag == qn("w:sectPr"):
                continue
            copied = deepcopy(element)
            _strip_comment_marks(copied, comment_ids)
            notes_dropped += _strip_note_references(copied)
            _rebuilt, flattened = _relink_hyperlinks(source, target, copied)
            links_flattened += flattened
            images += _copy_media(source, target, copied, used_part_names)
            numbering_dropped += _remap_numbering(copied, numbering_map)
            objects_dropped += _sanitize_rids(copied, target.part, drop_containers)
            if first_in_document:
                first_in_document = False
                if index > 0 and page_break and header_mode != HEADER_MODE_EACH:
                    if copied.tag == qn("w:p"):
                        _page_break_before(copied)
                    else:
                        _append_body(target, _page_break_paragraph(target), sect_pr)
                    pages += 1
            _append_body(target, copied, sect_pr)
        if header_mode == HEADER_MODE_EACH:
            # 这一节的正文已经就位，把 source 自己的抬头搬进刚分出来的这一节
            if _apply_section_header_footer(source, target):
                header_sections += 1
    comments = len(comment_ids)

    # 学校抬头/页码在第一份文件里：默认继承过来，别让合并稿变成没有抬头的白纸。
    # 逐份模式（each）下每节各管各的，这里不搬。
    # 纸张/边距跟第一份文件走（含 each 模式）——不抄就退回 python-docx 默认的 Letter
    if first_document is not None and first_document.sections:
        for section in target.sections:
            _copy_page_setup(first_document.sections[0], section)

    headers_inherited = 0
    if header_mode == HEADER_MODE_FIRST and first_document is not None:
        headers_inherited = _inherit_header_footer(first_document, target)

    page = payload.get("page")
    if page:
        _apply_page(target, page, True)

    title = payload.get("title")
    if title:
        target.core_properties.title = str(title)

    _ensure_parent(output)
    if os.path.exists(output) and not payload.get("overwrite", False):
        raise EngineError("output-exists", "目标文件已存在，不覆盖：" + output, "换一个文件名，或显式允许覆盖")
    target.save(output)
    warnings = []
    if comments:
        warnings.append({"code": "comments-dropped",
                         "message": str(comments) + " 处批注没有跟过来：合并稿里已移除批注锚点"
                                    "（否则 Word 打开会报「文件可能已经损坏」），批注仍留在各自的原文里"})
    if notes_dropped:
        warnings.append({"code": "notes-dropped",
                         "message": "去掉了 " + str(notes_dropped) + " 处脚注/尾注标记：合并只搬正文，"
                                    "注释内容和编号不会跟过来"})
    if links_flattened:
        warnings.append({"code": "links-flattened",
                         "message": str(links_flattened) + " 处超链接的目标地址找不到，按普通文字保留了"})
    if numbering_dropped:
        warnings.append({"code": "numbering-dropped",
                         "message": str(numbering_dropped) + " 处自动编号在原文件里没有对应定义，"
                                    "合并后按普通段落排列（题号可能需要手打）"})
    if objects_dropped:
        warnings.append({"code": "objects-dropped",
                         "message": "有 " + str(objects_dropped) + " 个图表/嵌入对象没能带过来"
                                    "（合并只搬文字、表格、图片；图表请单独截图）"})
    if headers_inherited:
        warnings.append({"code": "headers-inherited",
                         "message": "沿用第一份文件的页眉页脚（学校抬头、页码都在）；"
                                    "其余文件的页眉页脚带不过来"})
    elif source_headers and header_mode == HEADER_MODE_NONE:
        warnings.append({"code": "headers-not-merged",
                         "message": str(source_headers) + " 份文件原来有自己的页眉页脚，"
                                    "但你这次选了「不带页眉页脚」，合并稿里没有抬头"})
    if header_mode == HEADER_MODE_EACH:
        if header_sections:
            warnings.append({"code": "headers-per-section",
                             "message": "每份文件保留了自己的抬头：一共分成 " + str(len(target.sections)) +
                                        " 节，其中 " + str(header_sections) + " 节带出了自己的页眉页脚"
                                        "（每份文件取第一节的抬头）。页眉里的图片带不过来，"
                                        "那一节留空；页码按合并稿连续编号，不在每份文件处重新从 1 数；"
                                        "要删掉某一节的抬头，直接双击那一页的页眉改"})
        else:
            warnings.append({"code": "headers-per-section",
                             "message": "每份文件各占一节：一共分成 " + str(len(target.sections)) +
                                        " 节，但这些文件都没有自己的页眉页脚，合并稿里没有抬头"})
    return {
        "outputPath": os.path.abspath(output),
        "inputs": sources,
        "images": images,
        "comments": comments,
        "stylesMerged": styles_imported,
        "pageBreaks": pages,
        "headersInherited": headers_inherited,
        "headersMode": header_mode,
        "headerSections": header_sections,
        "bytes": os.path.getsize(output),
        "warnings": warnings,
    }


def docx_extract(payload):
    """payload = {path, tables?} → {text, paragraphs, tables}"""
    path = payload.get("path") or payload.get("inputPath")
    require(bool(path), "path-required", "没有指定要提取的文件",
            "点「浏览…」选一份 .docx，或把路径粘进输入框")
    _require_docx(path, "读文字")
    document = _open_document(path, "文件")
    want_tables = bool(payload.get("tables", True))
    # 按 body 里的真实顺序走一遍：以前先收全部段落、再把表格整块贴到文末，
    # 教案那种「段 → 表 → 段」会变成「段段表」（实测）。
    from docx.table import Table

    lines = []
    paragraph_count = 0
    tables = []
    for child in list(document.element.body):
        if child.tag == qn("w:p"):
            text = _paragraph_text(child).strip()
            if text:
                lines.append(text)
                paragraph_count += 1
        elif child.tag == qn("w:tbl"):
            try:
                table = Table(child, document)
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            except Exception:
                continue
            tables.append(rows)
            if want_tables:
                lines.append("")
                for row in rows:
                    lines.append(" | ".join(row))
    return {
        "text": "\n".join(lines).strip(),
        "paragraphCount": paragraph_count,
        "tables": tables if want_tables else [],
    }
