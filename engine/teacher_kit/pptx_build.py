"""造 pptx：把"一节课的若干页"写成真正的 PowerPoint 文件。

不追求设计感 —— 版式留给模板与后续预设（edupt 预设、HTML 演示都在 v0.2）。
这里只保证：默认 16:9（教室一体机就是 16:9）、中文不落到 Calibri、
每页标题与要点可编辑（不是整页图片）、文字太多会给出溢出提示而不是静默挤压。
"""

import os

from pptx import Presentation
from pptx.util import Emu, Mm, Pt

from .errors import EngineError, require

SLIDE_SIZES = {
    "16:9": (338.7, 190.5),  # 教室一体机/投影仪标准，绝不给 4:3 默认
    "4:3": (254.0, 190.5),
    "A4横向": (297.0, 210.0),
}

# 版式索引（默认模板 python-pptx 自带）：0 封面 1 标题+内容 5 仅标题 6 空白
LAYOUTS = {
    "title": 0,
    "titleAndContent": 1,
    "section": 2,
    "titleOnly": 5,
    "blank": 6,
}

DEFAULT_FONT = "微软雅黑"
CJK_MARKERS = ("宋体", "黑体", "微软雅黑", "等线", "楷体", "仿宋", "思源", "苹方", "华文")


# p14（PowerPoint 2010 扩展）不在 python-pptx 的 nsmap 里，qn("p14:xxx") 会 KeyError，
# 这里直接用 Clark 记法。
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


def _local_name(element):
    tag = element.tag if element is not None else ""
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _drop_empty_ext(ext):
    """把一个被清空的 p:ext 从 extLst 里摘掉；extLst 空了也摘掉。"""
    if ext is None or _local_name(ext) != "ext" or len(ext):
        return
    ext_lst = ext.getparent()
    if ext_lst is None:
        return
    ext_lst.remove(ext)
    if _local_name(ext_lst) == "extLst" and len(ext_lst) == 0:
        parent = ext_lst.getparent()
        if parent is not None:
            parent.remove(ext_lst)


def _prune_slide_refs(presentation, remaining_ids, remaining_rel_ids):
    """清掉分节（p14:sectionLst）和自定义放映（p:custShowLst）里指向已删页的引用。

    这两处存的不是关系，而是演示文稿级 sldId（分节用 id、自定义放映用 r:id），
    drop_rel 不会碰它们；留着就是悬空引用 —— 老师打开会看到模板自带的节名落在
    新页上，或者自定义放映指向不存在的页。只保留还能对上现存页的条目。
    """
    from pptx.oxml.ns import qn

    pres = presentation.part._element

    # 分节：<p14:section name="..."><p14:sldIdLst><p14:sldId id="256"/></...>
    for section_lst in list(pres.iter("{%s}sectionLst" % NS_P14)):
        for section in list(section_lst):
            if _local_name(section) != "section":
                continue
            kept = 0
            for sld_id_lst in list(section):
                if _local_name(sld_id_lst) != "sldIdLst":
                    continue
                for node in list(sld_id_lst):
                    if _local_name(node) == "sldId" and node.get("id") in remaining_ids:
                        kept += 1
                    else:
                        sld_id_lst.remove(node)
                if len(sld_id_lst) == 0:
                    section.remove(sld_id_lst)
            if kept <= 0:
                section_lst.remove(section)
        if len(section_lst) == 0:
            ext = section_lst.getparent()
            if ext is not None:
                ext.remove(section_lst)
            _drop_empty_ext(ext)

    # 自定义放映：<p:custShow><p:sldLst><p:sld r:id="rId2"/></...>
    cust_show_lst = pres.find(qn("p:custShowLst"))
    if cust_show_lst is not None:
        for cust_show in list(cust_show_lst):
            if _local_name(cust_show) != "custShow":
                continue
            kept = 0
            for sld_lst in list(cust_show):
                if _local_name(sld_lst) != "sldLst":
                    continue
                for node in list(sld_lst):
                    rel_id = node.get(qn("r:id"))
                    current = rel_id is not None and rel_id in remaining_rel_ids
                    if current:
                        kept += 1
                    else:
                        sld_lst.remove(node)
                if len(sld_lst) == 0:
                    cust_show.remove(sld_lst)
            if kept <= 0:
                cust_show_lst.remove(cust_show)
        if len(cust_show_lst) == 0:
            pres.remove(cust_show_lst)


def _clear_slides(presentation):
    """学校模板里常留着示例页（"单击输入标题"），照抄出去会被学生看见，先清干净。

    只从 sldIdLst 里摘掉是不够的：presentation 与这张幻灯片的 rel 还在，保存时旧
    slide 部件仍会写进包，和新页撞同一个 slide1.xml → 文件里出现重名条目（PowerPoint
    能开，但 WPS/校验工具会报损坏）。必须连 rel 一起丢，让旧页变成不可达部件。
    分节（sectionLst）与自定义放映（custShowLst）里还存着 sldId/rId 引用，也要同步
    剪掉，否则模板的旧节名会落到新前几页上、甚至指向不存在的页。
    """
    from pptx.oxml.ns import qn

    id_list = presentation.slides._sldIdLst
    removed = list(id_list)
    removed_ids = {slide_id.get("id") for slide_id in removed}
    removed_rel_ids = {slide_id.get(qn("r:id")) for slide_id in removed if slide_id.get(qn("r:id"))}
    remaining_ids = {slide_id.get("id") for slide_id in id_list} - removed_ids
    remaining_rel_ids = {slide_id.get(qn("r:id")) for slide_id in id_list if slide_id.get(qn("r:id"))} - removed_rel_ids
    # 必须先剪引用再丢 rel：Part.drop_rel 带引用计数，presentation.xml 里
    # custShowLst 还指着这些 rId 时它拒绝删，旧页会继续留在包里和新页重名。
    _prune_slide_refs(presentation, remaining_ids, remaining_rel_ids)
    for slide_id in removed:
        rel_id = slide_id.get(qn("r:id"))
        if rel_id:
            try:
                presentation.part.drop_rel(rel_id)
            except KeyError:
                pass
            if rel_id in presentation.part.rels:
                presentation.part._rels.pop(rel_id)
        id_list.remove(slide_id)
    return len(removed)


def _sync_slide_size_type(presentation, width_mm, height_mm):
    """p:sldSz 的 type 要跟实际尺寸一致。

    python-pptx 改 slide_width/height 时不会同步 type（还是 screen4x3），
    有些播放器/投影按 type 走，16:9 的课件会被拉成 4:3。
    """
    from pptx.oxml.ns import qn

    sld_sz = presentation.part._element.find(qn("p:sldSz"))
    if sld_sz is None:
        return None
    ratio = (width_mm / height_mm) if height_mm else 0
    if abs(ratio - 16.0 / 9.0) < 0.03:
        name = "screen16x9"
    elif abs(ratio - 4.0 / 3.0) < 0.03:
        name = "screen4x3"
    elif abs(width_mm - 297) < 1 and abs(height_mm - 210) < 1:
        name = "A4"
    else:
        name = "custom"
    sld_sz.set("type", name)
    return name


def _fill_theme_east_asian(presentation, font):
    """主题里的中文槽位是空的（<a:ea typeface=""/>）：PowerPoint 会按系统默认中文字体
    渲染，同一份课件在老师电脑上就变形。填成我们实际用的中文字体。"""
    changed = 0
    for part in presentation.part.package.iter_parts():
        name = str(part.partname)
        if "theme" not in name or not name.endswith(".xml"):
            continue
        blob = part.blob
        if b'<a:ea typeface=""/>' not in blob:
            continue
        part._blob = blob.replace(b'<a:ea typeface=""/>',
                                  ('<a:ea typeface="' + font + '"/>').encode("utf-8"))
        changed += 1
    return changed


def _layout_by_name(presentation, candidates, fallback_index):
    """模板的版式名各不相同（"标题幻灯片"/"封面"/"标题和内容"），按名字找，找不到退回索引。"""
    for layout in presentation.slide_layouts:
        name = (layout.name or "").strip()
        if any(candidate in name for candidate in candidates):
            return layout
    try:
        return presentation.slide_layouts[fallback_index]
    except IndexError:
        return presentation.slide_layouts[0]


def _resolve_layout(presentation, layout_name, from_template):
    if not from_template:
        return presentation.slide_layouts[LAYOUTS[layout_name]]
    if layout_name == "title":
        return _layout_by_name(presentation, ("封面", "标题幻灯片", "Title Slide"), 0)
    if layout_name in ("section",):
        return _layout_by_name(presentation, ("节标题", "Section"), 2)
    if layout_name in ("titleOnly",):
        return _layout_by_name(presentation, ("仅标题", "Title Only"), 5)
    if layout_name in ("blank",):
        return _layout_by_name(presentation, ("空白", "Blank"), 6)
    return _layout_by_name(presentation, ("标题和内容", "标题与内容", "Title and Content"), 1)


def _body_placeholder(slide):
    """正文占位符：不能假定一定是 placeholders[1]（模板的 idx 未必连着）。"""
    from pptx.enum.shapes import PP_PLACEHOLDER

    title = slide.shapes.title
    for shape in slide.placeholders:
        if title is not None and shape.shape_id == title.shape_id:
            continue
        if shape.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
                                             PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.CENTER_TITLE):
            if shape.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE:
                continue
            return shape
    for shape in slide.placeholders:
        if title is not None and shape.shape_id == title.shape_id:
            continue
        if shape.has_text_frame:
            return shape
    return None


def _ensure_parent(path):
    parent = os.path.dirname(os.path.abspath(path))
    if parent and not os.path.isdir(parent):
        os.makedirs(parent, exist_ok=True)


def _set_cjk_font(run, name):
    """必须同时写 latin 与 eastAsia，否则中文仍落在默认的 Calibri 上，WPS 打开会变字。"""
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    from pptx.oxml.ns import qn

    for tag in ("a:latin", "a:ea", "a:cs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = rpr.makeelement(qn(tag), {})
            rpr.append(element)
        element.set("typeface", name)


def _style_run(run, font=None, size_pt=None, bold=None, color=None):
    if font:
        _set_cjk_font(run, font)
    if size_pt:
        run.font.size = Pt(float(size_pt))
    if bold is not None:
        run.font.bold = bool(bold)
    if color:
        from pptx.dml.color import RGBColor

        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def _add_bullets(text_frame, bullets, font, size_pt, warnings, index):
    text_frame.clear()
    for position, item in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if position == 0 else text_frame.add_paragraph()
        text = item if isinstance(item, str) else str(item.get("text", ""))
        level = 0 if isinstance(item, str) else int(item.get("level") or 0)
        paragraph.level = min(level, 4)
        run = paragraph.add_run()
        run.text = text
        _style_run(run, font=font, size_pt=size_pt)
        # 粗估溢出：一页正文超过约 9 行 28 字就会挤出边界（真实几何检查在 qc 阶段做）
        if len(text) > 60:
            warnings.append({
                "slide": index,
                "code": "long-bullet",
                "message": "第 " + str(index) + " 页有一条要点超过 60 字，投影上大概率要换行挤在一起",
                "text": text[:40] + "…",
            })


def build_pptx(spec):
    """spec = {outputPath, slideSize?, slides:[{title, subtitle?, bullets?, notes?, layout?}], properties?, overwrite?}"""
    output = spec.get("outputPath")
    require(bool(output), "output-required", "没有指定输出文件路径", "在面板上选一个保存位置，或者把 outputPath 写成完整路径")
    slides = spec.get("slides") or []
    require(len(slides) > 0, "empty-deck", "一份课件至少要有一页")
    _ensure_parent(output)

    template = spec.get("template")
    from_template = bool(template)
    if from_template:
        require(os.path.exists(template), "file-not-found", "找不到课件模板：" + str(template),
                "把学校/自己常用的 .pptx 模板拖进来，或去掉模板直接用默认版式")
        require(str(template).lower().endswith(".pptx"), "unsupported-template", "课件模板必须是 .pptx：" + str(template))
        try:
            presentation = Presentation(template)
        except Exception:  # noqa: BLE001 - python-pptx 的英文异常不能直接端给老师
            raise EngineError("template-broken",
                              "这份课件模板打不开（不是 .pptx，或者文件已损坏）：" + os.path.basename(str(template)),
                              "用 PowerPoint / WPS 打开它另存为 .pptx，再试一次")
        removed_slides = _clear_slides(presentation) if spec.get("clearTemplateSlides", True) else 0
        width_mm = round(presentation.slide_width / 36000.0, 1)
        height_mm = round(presentation.slide_height / 36000.0, 1)
        size_name = "模板 " + str(width_mm) + "×" + str(height_mm) + "mm"
    else:
        removed_slides = 0
        size_name = spec.get("slideSize") or "16:9"
        if size_name not in SLIDE_SIZES:
            raise EngineError("unknown-slide-size", "不认识的画布比例：" + str(size_name), "可用：" + "、".join(SLIDE_SIZES))
        width_mm, height_mm = SLIDE_SIZES[size_name]
        presentation = Presentation()
        presentation.slide_width = Mm(width_mm)
        presentation.slide_height = Mm(height_mm)

    title_font = spec.get("titleFont") or DEFAULT_FONT
    body_font = spec.get("bodyFont") or DEFAULT_FONT
    _sync_slide_size_type(presentation, width_mm, height_mm)
    _fill_theme_east_asian(presentation, body_font)
    title_pt = spec.get("titleSizePt") or 40
    body_pt = spec.get("bodySizePt") or 28
    warnings = []

    for index, item in enumerate(slides, start=1):
        layout_name = item.get("layout") or ("title" if index == 1 and len(slides) > 1 else "titleAndContent")
        if layout_name not in LAYOUTS:
            raise EngineError("unknown-layout", "不认识的版式：" + str(layout_name), "可用：" + "、".join(LAYOUTS))
        slide = presentation.slides.add_slide(_resolve_layout(presentation, layout_name, from_template))

        title = item.get("title")
        title_size = item.get("titleSizePt") or title_pt
        if title and slide.shapes.title is not None:
            title_frame = slide.shapes.title.text_frame
            title_frame.clear()
            run = title_frame.paragraphs[0].add_run()
            run.text = str(title)
            _style_run(run, font=title_font, size_pt=title_size, bold=True)
        elif title:
            # 空白/仅标题版式没有标题占位符：以前标题凭空消失，老师只看到一页要点
            slide_w = presentation.slide_width
            slide_h = presentation.slide_height
            box = slide.shapes.add_textbox(Emu(int(slide_w * 0.06)), Emu(int(slide_h * 0.06)),
                                           Emu(int(slide_w * 0.88)), Emu(int(slide_h * 0.2)))
            box.text_frame.word_wrap = True
            run = box.text_frame.paragraphs[0].add_run()
            run.text = str(title)
            _style_run(run, font=title_font, size_pt=title_size, bold=True)

        body = item.get("bullets")
        subtitle = item.get("subtitle")
        body_shape = _body_placeholder(slide)
        if subtitle and layout_name == "title" and body_shape is not None:
            _add_bullets(body_shape.text_frame, [subtitle], body_font, body_pt, warnings, index)
        elif body and body_shape is not None:
            _add_bullets(body_shape.text_frame, body, body_font, body_pt, warnings, index)
        elif body:
            # 仅标题/空白版式没有内容占位符，自己放一个文本框，避免"要点凭空消失"
            box = slide.shapes.add_textbox(Mm(20), Mm(45), Mm(width_mm - 40), Mm(height_mm - 65))
            _add_bullets(box.text_frame, body, body_font, body_pt, warnings, index)

        notes = item.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)

    properties = spec.get("properties") or {}
    if properties:
        props = presentation.core_properties
        for field in ("title", "subject", "author", "comments"):
            if field in properties:
                setattr(props, field, properties[field])

    if os.path.exists(output) and not spec.get("overwrite", False):
        raise EngineError("output-exists", "目标文件已存在，不覆盖：" + output, "换一个文件名，或显式允许覆盖")
    presentation.save(output)
    return {
        "outputPath": os.path.abspath(output),
        "slides": len(slides),
        "removedTemplateSlides": removed_slides if from_template else 0,
        "slideSize": size_name,
        "sizeMm": {"width": width_mm, "height": height_mm},
        "bytes": os.path.getsize(output),
        "warnings": warnings,
    }