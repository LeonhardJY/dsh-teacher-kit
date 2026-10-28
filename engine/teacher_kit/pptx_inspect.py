"""读 pptx 的版式事实：幻灯片尺寸、每页形状、文字、图片、字体。"""

from pptx import Presentation
from pptx.util import Emu

from .errors import EngineError
from .units import emu_to_mm


def _slide_facts(slide, index):
    texts = []
    images = 0
    tables = 0
    shapes = 0
    for shape in slide.shapes:
        shapes += 1
        if shape.shape_type == 13 or getattr(shape, "image", None) is not None:
            images += 1
        if getattr(shape, "has_table", False):
            tables += 1
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    texts.append(text)
    notes = ""
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        notes = ""
    return {
        "index": index + 1,
        "layout": slide.slide_layout.name,
        "shapes": shapes,
        "images": images,
        "tables": tables,
        "texts": texts[:20],
        "textChars": sum(len(t) for t in texts),
        "notes": notes[:400] or None,
    }


def _fonts(presentation):
    counts = {}
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    name = run.font.name
                    if name:
                        counts[name] = counts.get(name, 0) + 1
    return [{"font": name, "runs": count} for name, count in sorted(counts.items(), key=lambda kv: -kv[1])[:12]]


def _metadata(presentation):
    props = presentation.core_properties
    return {
        "title": props.title or None,
        "subject": props.subject or None,
        "author": props.author or None,
        "comments": props.comments or None,
        "lastModifiedBy": props.last_modified_by or None,
    }


def inspect_pptx(path, preview=6):
    try:
        presentation = Presentation(path)
    except Exception as error:
        raise EngineError("pptx-unreadable", f"打不开这个 PowerPoint 文件：{path}", str(error))

    slides = [_slide_facts(slide, index) for index, slide in enumerate(presentation.slides)]
    return {
        "kind": "pptx",
        "path": path,
        "slideSizeMm": {
            "width": emu_to_mm(presentation.slide_width),
            "height": emu_to_mm(presentation.slide_height),
        },
        "slideAspect": _aspect(presentation),
        "slideCount": len(slides),
        "slides": slides,  # 全量：逐页预览与印前检查都要每一页，不做截断
        "metadata": _metadata(presentation),
        "fonts": _fonts(presentation),
        "emptySlides": [s["index"] for s in slides if s["textChars"] == 0 and s["images"] == 0],
        "textOnlySlides": [s["index"] for s in slides if s["images"] == 0],
    }


def _aspect(presentation):
    try:
        return round(presentation.slide_width / presentation.slide_height, 3)
    except Exception:
        return None
