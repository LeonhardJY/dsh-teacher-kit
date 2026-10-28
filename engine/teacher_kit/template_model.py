"""模板反推：从老师自己的真实文件里读出"这份卷子长什么样"。

为什么不内置一堆"标准模板"：每个学校、每个备课组的版式都不一样，
老师的真实要求就藏在她上学期那份成品里。所以：
  1. 让她把 1—3 份旧成品（或学校空白模板）拖进来；
  2. 我们读出纸张、页边距、分栏、字体字号、行距、页眉页脚句式、题号层级；
  3. 多份之间不一致的地方，明确列成"冲突"，请她自己挑一个 —— 不替她拍板。

输出（profile）同时给三处用：
  · docx_build 的 page 参数（page_from_profile）；
  · 印前检查器的基准（检查器拿 profile 当"应该长什么样"）；
  · 面板里的"已识别/待确认"清单。
"""

import re

import docx
from docx.oxml.ns import qn

from .errors import EngineError, require
from .paper import PAPERS, resolve as resolve_paper
from .units import pt_to_mm

# ── 文本角色分类 ────────────────────────────────────────────────────────────
SECTION_RE = re.compile(r"^\s*([一二三四五六七八九十]+)\s*[、.．]")
QUESTION_RE = re.compile(r"^\s*(\d{1,2})\s*[、.．]")
SUBSECTION_RE = re.compile(r"^\s*[（(]\s*([一二三四五六七八九十]+)\s*[）)]")
SUBQUESTION_RE = re.compile(r"^\s*[（(]\s*(\d{1,2})\s*[）)]")
OPTION_RE = re.compile(r"^\s*([A-H])\s*[、.．]")
TITLE_HINT = ("试卷", "测试卷", "考试", "期中", "期末", "单元", "月考", "模拟", "答题卡", "答题卷", "答案")
BLANK_RE = re.compile(r"[_＿]{2,}")
SCORE_RE = re.compile(r"[（(]\s*\d+\s*分\s*[）)]|共\s*\d+\s*分")

ROLES = ("title", "instruction", "section", "question", "subquestion", "option", "body")


def classify(text, index):
    clean = (text or "").strip()
    if not clean:
        return "blank"
    if index <= 2 and len(clean) <= 40 and any(hint in clean for hint in TITLE_HINT):
        return "title"
    if SECTION_RE.match(clean):
        return "section"
    if SUBSECTION_RE.match(clean):
        return "section"
    if QUESTION_RE.match(clean):
        return "question"
    if SUBQUESTION_RE.match(clean):
        return "subquestion"
    if OPTION_RE.match(clean):
        return "option"
    if SCORE_RE.search(clean) or clean.endswith(("：", ":")) and len(clean) < 40:
        return "instruction"
    return "body"


# ── 读单个文件 ──────────────────────────────────────────────────────────────

def _run_attr(run, attr):
    """读 run 上的显式设置（没写就是 None，交给样式与默认值）。"""
    if attr == "font":
        name = run.font.name
        rpr = run._element.rPr
        east = None
        if rpr is not None:
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is not None:
                east = rfonts.get(qn("w:eastAsia"))
        return east or name
    if attr == "sizePt":
        size = run.font.size
        return None if size is None else round(size.pt, 1)
    if attr == "bold":
        return run.font.bold
    return None


def _style_font(document, style_name):
    try:
        style = document.styles[style_name]
    except Exception:
        return None, None
    name = style.font.name
    east = None
    element = getattr(style, "element", None)
    if element is not None:
        rpr = element.find(qn("w:rPr"))
        if rpr is not None:
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is not None:
                east = rfonts.get(qn("w:eastAsia"))
    size = style.font.size
    return (east or name), (None if size is None else round(size.pt, 1))


def _paragraph_facts(paragraph, document, index):
    text = paragraph.text
    runs = [run for run in paragraph.runs if run.text.strip()]
    fonts = [_run_attr(run, "font") for run in runs]
    sizes = [_run_attr(run, "sizePt") for run in runs]
    bolds = [run.bold for run in runs]
    style_name = paragraph.style.name if paragraph.style is not None else None
    style_font, style_size = _style_font(document, style_name) if style_name else (None, None)
    fallback_font, fallback_size = _style_font(document, "Normal")
    fmt = paragraph.paragraph_format
    spacing = None
    if fmt.line_spacing is not None:
        try:
            spacing = round(float(fmt.line_spacing), 2)
        except Exception:
            spacing = None
    return {
        "index": index,
        "role": classify(text, index),
        "text": text.strip()[:120],
        "style": style_name,
        "align": None if fmt.alignment is None else str(fmt.alignment).split(" ")[0].split(".")[-1],
        "font": _majority([f for f in fonts if f]) or style_font or fallback_font,
        "sizePt": _majority([s for s in sizes if s]) or style_size or fallback_size,
        "bold": _majority([b for b in bolds if b is not None]),
        "lineSpacing": spacing,
        "indentFirstLineMm": None if fmt.first_line_indent is None else pt_to_mm(fmt.first_line_indent.pt),
        "textChars": len(text.strip()),
    }


def _majority(values):
    if not values:
        return None
    counts = {}
    for value in values:
        counts[value] = counts.get(value, 0) + 1
    return sorted(counts.items(), key=lambda kv: (-kv[1], str(kv[0])))[0][0]


def _footer_signature(part, label):
    """页眉页脚：正文 + 有没有页码域（第 X 页 共 Y 页 这类）。"""
    texts = []
    has_page_field = False
    try:
        for paragraph in part.paragraphs:
            if paragraph.text.strip():
                texts.append(paragraph.text.strip())
        for element in part._element.iter():
            if element.tag == qn("w:instrText") and element.text and "PAGE" in element.text.upper():
                has_page_field = True
    except Exception:
        pass
    return {"label": label, "texts": texts, "hasPageField": has_page_field}


def _label_placeholder(text, pattern, token):
    """把"使用班级：七(3)班"换成"使用班级：{{className}}"，已经是占位符就原样保留。"""
    def repl(match):
        existing = re.search(r"\{\{\w+\}\}", match.group(0))
        return match.group(1) + "：" + (existing.group(0) if existing else "{{" + token + "}}")
    return re.sub(pattern, repl, text)


def templatize(text):
    """把"这一份"的具体值抽成占位符，才能套到下个学期去。"""
    value = text or ""
    # 只吞年份区间，「学年」两个字留在占位符外：以前整个「2025—2026 学年」都变成
    # {{schoolYear}}，套模板时 values.schoolYear（2026—2027）后面就丢了「学年」。
    value = re.sub(r"20\d{2}\s*[—\-~至]\s*20\d{2}(?!\d)", "{{schoolYear}}", value)
    value = re.sub(r"第\s*([XxYyＸＹ＃]|\d+)?\s*页", "第{{page}}页", value)
    value = re.sub(r"共\s*([XxYyＸＹ＃]|\d+)?\s*页", "共{{pages}}页", value)
    value = BLANK_RE.sub("{{blank}}", value)
    # 姓名/班级有专用占位符：老师在面板上填了班级和出卷老师，要靠它们印进页脚。
    # 注意别把已经换好的 {{className}} 又吞成 {{blank}}（旧写法 \S* 会把它整个吃掉）。
    value = _label_placeholder(value, r"(使用\s*班级|班级)\s*[：:]\s*(?:\{\{\w+\}\}|\S*)", "className")
    value = _label_placeholder(value, r"(出卷老师|出卷人|命题人|制卷人)\s*[：:]\s*(?:\{\{\w+\}\}|\S*)", "teacherName")
    return value.strip()


def _columns_of(section):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        return 1
    try:
        return int(cols.get(qn("w:num")) or 1)
    except Exception:
        return 1


def _paper_of(section):
    width_mm = round(section.page_width.mm, 1) if section.page_width else None
    height_mm = round(section.page_height.mm, 1) if section.page_height else None
    matches = [
        name for name, size in PAPERS.items()
        if width_mm and height_mm and abs(size[0] - min(width_mm, height_mm)) < 3 and abs(size[1] - max(width_mm, height_mm)) < 3
    ]
    orientation = "landscape" if (width_mm or 0) > (height_mm or 0) else "portrait"
    return {
        "size": matches[0] if matches else None,
        "widthMm": width_mm,
        "heightMm": height_mm,
        "orientation": orientation,
        "marginMm": {
            "top": round(section.top_margin.mm, 1) if section.top_margin else None,
            "bottom": round(section.bottom_margin.mm, 1) if section.bottom_margin else None,
            "left": round(section.left_margin.mm, 1) if section.left_margin else None,
            "right": round(section.right_margin.mm, 1) if section.right_margin else None,
        },
        "columns": _columns_of(section),
    }


def analyse_docx(path):
    """一份 docx → 版式事实（未合并）。"""
    document = docx.Document(path)
    if not document.sections:
        raise EngineError("no-section", "这份 Word 里没有任何页面设置：" + path)
    section = document.sections[0]
    paragraphs = [_paragraph_facts(p, document, i) for i, p in enumerate(document.paragraphs)]
    used = [p for p in paragraphs if p["role"] != "blank"]

    by_role = {}
    for role in ROLES:
        sample = [p for p in used if p["role"] == role]
        if not sample:
            continue
        by_role[role] = {
            "count": len(sample),
            "font": _majority([p["font"] for p in sample if p["font"]]),
            "sizePt": _majority([p["sizePt"] for p in sample if p["sizePt"]]),
            "bold": _majority([p["bold"] for p in sample if p["bold"] is not None]),
            "align": _majority([p["align"] for p in sample if p["align"]]),
            "lineSpacing": _majority([p["lineSpacing"] for p in sample if p["lineSpacing"]]),
            "samples": [p["text"][:60] for p in sample[:3]],
        }

    header = _footer_signature(section.header, "header")
    footer = _footer_signature(section.footer, "footer")
    paper = _paper_of(section)
    # 一份卷子常被切成多节（标题跨整页单栏 + 正文两栏），栏数要看"最大的那节"，
    # 否则拿自己出的卷子反推模板，模板会变成一栏。
    try:
        paper["columns"] = max([_columns_of(item) for item in document.sections] or [paper["columns"]])
    except Exception:
        pass
    return {
        "path": path,
        "paper": paper,
        "roles": by_role,
        "header": {
            "raw": header["texts"],
            "template": [templatize(t) for t in header["texts"]],
        },
        "footer": {
            "raw": footer["texts"],
            "template": [templatize(t) for t in footer["texts"]],
            "hasPageField": footer["hasPageField"],
        },
        "numbering": {
            "sections": sorted({SECTION_RE.match(p["text"]).group(1) for p in used if SECTION_RE.match(p["text"])}),
            "questions": sorted({QUESTION_RE.match(p["text"]).group(1) for p in used if QUESTION_RE.match(p["text"])}, key=int),
            "subquestions": sorted({SUBQUESTION_RE.match(p["text"]).group(1) for p in used if SUBQUESTION_RE.match(p["text"])}, key=int),
            "options": sorted({OPTION_RE.match(p["text"]).group(1) for p in used if OPTION_RE.match(p["text"])}),
        },
        "structure": {
            "paragraphs": len(paragraphs),
            "sections": sum(1 for p in used if p["role"] == "section"),
            "questions": sum(1 for p in used if p["role"] == "question"),
            "tables": len(document.tables),
            "images": len(document.inline_shapes),
        },
        "paragraphs": paragraphs,
    }


# ── 合并多份 + 冲突检测 ────────────────────────────────────────────────────

def _merge_scalar(field, values, tolerance=0.0):
    """values = [(value, source)]；一致→值；不一致→冲突（不替老师拍板）。"""
    present = [(value, source) for value, source in values if value is not None]
    if not present:
        return None, None
    distinct = []
    for value, source in present:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if not any(isinstance(d["value"], (int, float)) and abs(d["value"] - value) <= tolerance for d in distinct):
                distinct.append({"value": value, "sources": [source]})
            else:
                for d in distinct:
                    if isinstance(d["value"], (int, float)) and abs(d["value"] - value) <= tolerance:
                        d["sources"].append(source)
        else:
            hit = next((d for d in distinct if d["value"] == value), None)
            if hit:
                hit["sources"].append(source)
            else:
                distinct.append({"value": value, "sources": [source]})
    if len(distinct) == 1:
        return distinct[0]["value"], None
    winner = sorted(distinct, key=lambda d: -len(d["sources"]))[0]
    return winner["value"], {"field": field, "options": distinct, "hint": "这几份文件不一致，请选一个（默认用了出现最多的）"}


def build_profile(paths, options=None):
    """paths = 1—3 份真实文件；返回合并后的 profile 与冲突清单。"""
    require(bool(paths), "paths-required", "请先选 1—3 份你自己的文件（上学期的成品或学校模板）", "把文件从「文件管理器」拖进面板的虚线框，或者点它选文件")
    require(len(paths) <= 5, "too-many-sources", "一次最多分析 5 份，太多了反而看不出规律")
    sources = []
    for path in paths:
        require(str(path).lower().endswith((".docx", ".docm")), "unsupported-source", "现在只支持 .docx 模板/成品：" + str(path),
            "在 Word / WPS 里打开它，「另存为」.docx 再选；.doc 和 .pdf 读不了版式")
        sources.append(analyse_docx(path))

    conflicts = []

    def pick(field, getter, tolerance=0.0):
        value, conflict = _merge_scalar(field, [(getter(s), s["path"]) for s in sources], tolerance)
        if conflict:
            conflicts.append(conflict)
        return value

    paper = {
        "size": pick("paper.size", lambda s: s["paper"]["size"]),
        "widthMm": pick("paper.widthMm", lambda s: s["paper"]["widthMm"], 1.0),
        "heightMm": pick("paper.heightMm", lambda s: s["paper"]["heightMm"], 1.0),
        "orientation": pick("paper.orientation", lambda s: s["paper"]["orientation"]),
        "columns": pick("paper.columns", lambda s: s["paper"]["columns"]),
        "marginMm": {
            side: pick("paper.margin." + side, lambda s, side=side: s["paper"]["marginMm"][side], 1.0)
            for side in ("top", "bottom", "left", "right")
        },
    }
    roles = {}
    for role in ROLES:
        sample_sources = [s for s in sources if role in s["roles"]]
        if not sample_sources:
            continue
        roles[role] = {
            "count": max(s["roles"][role]["count"] for s in sample_sources),
            "font": pick(role + ".font", lambda s, role=role: s["roles"].get(role, {}).get("font")),
            "sizePt": pick(role + ".sizePt", lambda s, role=role: s["roles"].get(role, {}).get("sizePt"), 0.6),
            "bold": pick(role + ".bold", lambda s, role=role: s["roles"].get(role, {}).get("bold")),
            "align": pick(role + ".align", lambda s, role=role: s["roles"].get(role, {}).get("align")),
            "lineSpacing": pick(role + ".lineSpacing", lambda s, role=role: s["roles"].get(role, {}).get("lineSpacing"), 0.05),
            "samples": sample_sources[0]["roles"][role]["samples"],
        }

    header_templates = sources[0]["header"]["template"]
    footer = {
        "template": pick("footer.template", lambda s: " / ".join(s["footer"]["template"]) or None),
        "hasPageField": pick("footer.hasPageField", lambda s: s["footer"]["hasPageField"]),
    }
    numbering = {
        key: pick("numbering." + key, lambda s, key=key: "、".join(s["numbering"][key]) or None)
        for key in ("sections", "questions", "subquestions", "options")
    }

    profile = {
        "schemaVersion": 1,
        "sources": [{"path": s["path"], "paragraphs": s["structure"]["paragraphs"], "questions": s["structure"]["questions"]} for s in sources],
        "paper": paper,
        "roles": roles,
        "header": {
            "template": header_templates,
            "raw": pick("header.raw", lambda s: " / ".join(s["header"]["raw"]) or None),
        },
        "footer": footer,
        "numbering": numbering,
        "structure": {
            "tables": pick("structure.tables", lambda s: s["structure"]["tables"]),
            "images": pick("structure.images", lambda s: s["structure"]["images"]),
        },
        "conflicts": conflicts,
        "confidence": "high" if len(sources) >= 2 and not conflicts else ("medium" if not conflicts else "low"),
    }
    return {"profile": profile, "analyse": sources}


def page_from_profile(profile):
    """profile → docx_build 的 page 参数（可直接传）。"""
    paper = (profile or {}).get("paper") or {}
    page = {"columns": paper.get("columns") or 1, "orientation": paper.get("orientation") or "portrait"}
    if paper.get("size"):
        page["size"] = paper["size"]
    else:
        page["size"] = _nearest_paper(paper.get("widthMm"), paper.get("heightMm"))
    margins = paper.get("marginMm") or {}
    given = {key: value for key, value in margins.items() if value is not None}
    if given:
        page["marginsMm"] = given
    return page


def _nearest_paper(width_mm, height_mm):
    if not width_mm or not height_mm:
        return "A4"
    short, long = min(width_mm, height_mm), max(width_mm, height_mm)
    best, best_gap = "A4", None
    for name, size in PAPERS.items():
        gap = abs(size[0] - short) + abs(size[1] - long)
        if best_gap is None or gap < best_gap:
            best, best_gap = name, gap
    return best


def diff_profiles(before, after):
    """两份 profile 的差异（用于"这学期和上学期版式哪里变了"）。"""
    require(isinstance(before, dict) and isinstance(after, dict), "bad-profile", "profile 必须是对象")
    changes = []

    def walk(prefix, left, right):
        if isinstance(left, dict) and isinstance(right, dict):
            for key in sorted(set(left) | set(right)):
                walk(prefix + "." + key if prefix else key, left.get(key), right.get(key))
            return
        if left != right:
            changes.append({"field": prefix, "before": left, "after": right})

    walk("paper", before.get("paper") or {}, after.get("paper") or {})
    walk("roles", before.get("roles") or {}, after.get("roles") or {})
    walk("numbering", before.get("numbering") or {}, after.get("numbering") or {})
    return {"changes": changes, "changed": len(changes)}
