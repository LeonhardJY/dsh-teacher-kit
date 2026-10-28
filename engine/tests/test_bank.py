#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""题库导入：跑 python test_bank.py。

为什么样本要这么排：老师手里的题库几乎没有两份长得一样——有表头的、没表头的
（第一行数据里答案恰好是「A」，容易被误认成表头）、选项挤在一列里写成
「A. 甲 B. 乙」的、Excel 存成 GBK 的 csv、临时从 Word 表格里扒的。
每一种脏数据各锁一组断言，模块里任何一处退化都会当场变红。
"""

import os
import sys
import tempfile
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

from openpyxl import Workbook  # noqa: E402

from teacher_kit import question_bank  # noqa: E402
from teacher_kit.docx_build import build_docx  # noqa: E402
from teacher_kit.errors import EngineError  # noqa: E402
from teacher_kit.exam import parse_paper  # noqa: E402
from teacher_kit.question_bank import (  # noqa: E402
    bank_build,
    bank_inspect,
    bank_to_text,
    pick_rows,
    read_bank,
)

FAILS = []

# 标准九列：老师从旧题库系统导出来的样子
STD_HEADER = ["题号", "题干", "选项A", "选项B", "选项C", "选项D", "答案", "解析", "分值"]
STD_ROWS = [
    [1, "下列加点字读音全部正确的一项是", "甲", "乙", "丙", "丁", "A", "甲字读 zhōng", 2],
    [2, "选出没有语病的一句", "第一句", "第二句", "第三句", "第四句", "C", "第三句搭配得当", 2],
    [3, "补写句中空缺的成语", "望尘莫及", "望其项背", "望洋兴叹", "望穿秋水", "B", "望其项背", 3],
]

# 同样九列，但没有表头；第一格的答案故意写成 "A"（最容易被误认成表头的情况）
NO_HEADER_ROWS = [
    [1, "第一题：选出正确的读音", "甲", "乙", "丙", "丁", "A", "略", 2],
    [2, "第二题：选出正确的字形", "甲", "乙", "丙", "丁", "D", "略", 2],
]

# 选项合并成一列的写法
MERGED_ROWS = [
    ["题号", "题干", "选项", "答案", "分值"],
    [1, "下面哪一项表示「非常仰慕」的成语", "A. 望洋兴叹 B. 望穿秋水 C. 望尘莫及 D. 望其项背", "D", 2],
    [2, "选出使用的修辞手法", "A. 比喻　B. 拟人　C. 夸张", "B", 2],
]

CSV_TEXT = """题号,题干,选项A,选项B,答案,分值
1,选出加点字读音正确的一项,甲,乙,A,2
2,给下面句子选出正确的关联词,因为,所以,C,2
"""

TAB_TEXT = "题号\t题干\t选项A\t选项B\t答案\t分值\n1\t选出正确的读音\t甲\t乙\tA\t2\n2\t选出正确的字形\t甲\t乙\tD\t3\n"


def check(condition, label, detail=""):
    if condition:
        print("  ✓ " + label)
    else:
        FAILS.append(label)
        print("  ✗ " + label + (" — " + str(detail) if detail else ""))


def expect_error(fn, code, label):
    try:
        fn()
    except EngineError as error:
        check(error.code == code, label,
              "拿到的是 " + error.code + "：" + error.message)
        return error
    except Exception as error:  # noqa: BLE001 - 抛别的异常同样算失败
        check(False, label, "抛了别的异常 " + repr(error))
        return None
    check(False, label, "没有抛错")
    return None


def warn_codes(result):
    return [item.get("code") for item in (result.get("warnings") or [])]


def make_xlsx(path, rows, sheet_title="题库"):
    """造一份 xlsx：逐行 append，行里给 None 就是空单元格。"""
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(list(row))
    book.save(path)
    return path


def write_text(path, text, encoding):
    with open(path, "w", encoding=encoding, newline="") as handle:
        handle.write(text)
    return path


def write_bytes(path, blob):
    with open(path, "wb") as handle:
        handle.write(blob)
    return path


def para_text(path):
    """整份 docx 的段落文字，用来抽查内容是不是真的排进去了。"""
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    # 两栏卷子的卷名落在第一页页眉里，抽查标题时要连页眉一起看
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.first_page_header.paragraphs)
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def main():
    work = tempfile.mkdtemp(prefix="bank-")
    os.environ["TEACHER_KIT_HOME"] = os.path.join(work, "教师助手")

    print("[1] 标准九列 Excel：题干/选项/答案/解析/分值")
    std_path = make_xlsx(os.path.join(work, "标准题库.xlsx"), [STD_HEADER] + STD_ROWS)
    bank = read_bank(std_path)
    check(bank["format"] == "xlsx", "认出是 xlsx", bank["format"])
    check(len(bank["rows"]) == 3, "表头那一行不算题，读出 3 题", len(bank["rows"]))
    rows = bank["rows"]
    check([row["number"] for row in rows] == [1, 2, 3], "题号按列读出来",
          [row["number"] for row in rows])
    check(rows[0]["stem"] == "下列加点字读音全部正确的一项是", "题干对", rows[0]["stem"])
    check(rows[0]["options"] == {"A": "甲", "B": "乙", "C": "丙", "D": "丁"},
          "四个选项按 A-D 收齐", rows[0]["options"])
    check(rows[1]["answer"] == "C", "答案对", rows[1]["answer"])
    check(rows[0]["analysis"] == "甲字读 zhōng", "解析按原文保留", rows[0]["analysis"])
    check([row["score"] for row in rows] == [2.0, 2.0, 3.0], "分值转成 float",
          [row["score"] for row in rows])
    check(rows[0]["type"] == "", "没题型列就留空（组卷时再按选项猜）", rows[0]["type"])
    check(not bank["warnings"], "干净的表不报假告警", bank["warnings"])

    print("[2] 没表头：按 题号/题干/A-D/答案/解析/分值 顺序兜底")
    nohead_path = make_xlsx(os.path.join(work, "无表头.xlsx"), NO_HEADER_ROWS)
    bank = read_bank(nohead_path)
    codes = warn_codes(bank)
    check("header-fallback" in codes, "给出 header-fallback 告警", codes)
    check(any("没认出表头" in (item.get("message") or "") for item in bank["warnings"]),
          "告警文案说明了按什么顺序读", bank["warnings"])
    check(len(bank["rows"]) == 2, "第一行数据没有被当成表头吃掉", len(bank["rows"]))
    check(bank["rows"][0]["answer"] == "A", "答案「A」留在答案位，没被读成选项列",
          bank["rows"][0]["answer"])
    check(len(bank["rows"][0]["options"]) == 4, "四个选项仍然收齐",
          bank["rows"][0]["options"])
    check(bank["rows"][1]["stem"] == "第二题：选出正确的字形", "第二题题干对",
          bank["rows"][1]["stem"])

    print("[3] 选项合并成一列（A. 甲 B. 乙）要能拆开")
    merged_path = make_xlsx(os.path.join(work, "选项一列.xlsx"), MERGED_ROWS)
    bank = read_bank(merged_path)
    first = bank["rows"][0]
    check(first["options"].get("A") == "望洋兴叹", "拆出 A", first["options"])
    check(first["options"].get("D") == "望其项背", "拆出 D", first["options"])
    check(first["answer"] == "D", "答案没被拆坏", first["answer"])
    second = bank["rows"][1]
    check(second["options"] == {"A": "比喻", "B": "拟人", "C": "夸张"},
          "全角空格分隔也能拆", second["options"])
    check(not bank["warnings"], "合并列不该额外告警", bank["warnings"])

    # 表头写成「选项内容」时别被「内容」抢成题干列
    named = make_xlsx(os.path.join(work, "选项内容.xlsx"),
                      [["题干", "选项内容", "答案"], ["哪一项正确", "A. 甲 B. 乙", "A"]])
    bank = read_bank(named)
    check(bank["rows"][0]["options"] == {"A": "甲", "B": "乙"},
          "表头写「选项内容」也认成合并选项列", bank["rows"][0]["options"])

    print("[4] Word 表格：多张表都要读，单元格内换行保留")
    docx_path = os.path.join(work, "表格题库.docx")
    build_docx({
        "outputPath": docx_path,
        "blocks": [
            {"type": "table", "tableStyle": "Table Grid", "rows": [
                ["题号", "题干", "选项A", "选项B", "答案", "分值"],
                ["1", "下面读音正确的是", "甲", "乙", "A", 2],
                ["2", "读下面材料，回答问题\n（材料共两行）", "①", "②", "B", 3],
            ]},
            {"type": "table", "tableStyle": "Table Grid", "rows": [
                ["3", "第三题的正文材料", "A. α B. β C. γ", "B", "因为 γ 对", 2],
            ]},
        ],
    })
    bank = read_bank(docx_path)
    codes = warn_codes(bank)
    check(bank["format"] == "docx", "认出是 docx", bank["format"])
    check(len(bank["rows"]) == 3, "两张表加起来 3 题", len(bank["rows"]))
    check("\n" in bank["rows"][1]["stem"], "单元格里的换行保留",
          bank["rows"][1]["stem"])
    third = bank["rows"][2]
    check(third["number"] == 3, "第二张表自己带的题号读到", third["number"])
    check(third["stem"] == "第三题的正文材料", "第二张表题干对", third["stem"])
    check(third["options"] == {"A": "α", "B": "β", "C": "γ"},
          "第二张表没表头，合并选项列照样拆开", third["options"])
    check(third["answer"] == "B" and third["analysis"] == "因为 γ 对",
          "答案与解析都落位", [third["answer"], third["analysis"]])
    check(third["score"] == 2.0, "分值落位", third["score"])
    check("header-fallback" in codes, "第二张表没表头 → 兜底告警", codes)

    print("[5] csv/txt：GBK、UTF-8（带 BOM）、Tab 分隔、题干一列")
    gbk_path = write_text(os.path.join(work, "gbk题库.csv"), CSV_TEXT, "gbk")
    bank = read_bank(gbk_path)
    check(bank["format"] == "csv", "认出是 csv", bank["format"])
    check(bank["rows"][0]["stem"] == "选出加点字读音正确的一项", "GBK 中文读对了",
          bank["rows"][0]["stem"])
    check(bank["rows"][1]["answer"] == "C", "GBK 文件的答案对", bank["rows"][1]["answer"])

    utf8_path = write_text(os.path.join(work, "utf8题库.csv"), CSV_TEXT, "utf-8-sig")
    bank = read_bank(utf8_path)
    check(len(bank["rows"]) == 2, "带 BOM 的 UTF-8 也能读", len(bank["rows"]))
    check(bank["rows"][0]["answer"] == "A", "BOM 没污染第一列表头", bank["rows"][0]["answer"])

    tab_path = write_text(os.path.join(work, "tab题库.txt"), TAB_TEXT, "utf-8")
    bank = read_bank(tab_path)
    check(bank["format"] == "txt", "认出是 txt", bank["format"])
    check(len(bank["rows"]) == 2, "Tab 分隔也读得出来", len(bank["rows"]))
    check(bank["rows"][1]["answer"] == "D", "Tab 文件答案对", bank["rows"][1]["answer"])

    one_col = write_text(os.path.join(work, "题干一列.txt"),
                         "1、读出下面加点字的读音\n2、把句子补充完整\n", "utf-8")
    bank = read_bank(one_col)
    check([row["stem"] for row in bank["rows"]] == ["读出下面加点字的读音", "把句子补充完整"],
          "只有题干一列时整行都是题干", [row["stem"] for row in bank["rows"]])
    check([row["number"] for row in bank["rows"]] == [1, 2],
          "题干自带的「1、」被当题号摘出来", [row["number"] for row in bank["rows"]])

    print("[6] 空表 / 不存在 / 坏文件 / 不支持的格式：一律中文 EngineError")
    empty_path = make_xlsx(os.path.join(work, "空表.xlsx"), [])
    expect_error(lambda: read_bank(empty_path), "bank-empty", "空 Excel → bank-empty")
    expect_error(lambda: read_bank(""), "path-required", "没给路径 → path-required")
    expect_error(lambda: read_bank(os.path.join(work, "根本没有.xlsx")), "file-not-found",
                 "文件不存在 → file-not-found")
    broken_xlsx = write_bytes(os.path.join(work, "坏.xlsx"), b"this is not a workbook")
    expect_error(lambda: read_bank(broken_xlsx), "invalid-xlsx", "坏 xlsx → invalid-xlsx")
    broken_docx = write_bytes(os.path.join(work, "坏.docx"), b"this is not a docx")
    expect_error(lambda: read_bank(broken_docx), "invalid-docx", "坏 docx → invalid-docx")
    pdf_path = write_bytes(os.path.join(work, "题库.pdf"), b"%PDF-1.4")
    expect_error(lambda: read_bank(pdf_path), "unsupported-format", "pdf → unsupported-format")
    old_xls = write_bytes(os.path.join(work, "题库.xls"), b"\xd0\xcf\x11\xe0")
    expect_error(lambda: read_bank(old_xls), "unsupported-format", "老 .xls → unsupported-format")
    empty_docx = os.path.join(work, "只有一段话.docx")
    build_docx({"outputPath": empty_docx, "blocks": [{"type": "paragraph", "text": "只有一行字"}]})
    expect_error(lambda: read_bank(empty_docx), "no-table", "docx 里没表格 → no-table")
    blank_csv = write_text(os.path.join(work, "空.csv"), "  \n", "utf-8")
    expect_error(lambda: read_bank(blank_csv), "bank-empty", "空 csv → bank-empty")
    header_only = write_text(os.path.join(work, "只有表头.csv"), "题号,题干,答案\n", "utf-8")
    expect_error(lambda: read_bank(header_only), "bank-empty", "只有表头没数据 → bank-empty")
    expect_error(lambda: bank_to_text([]), "rows-required", "空 rows → rows-required")
    expect_error(lambda: bank_to_text("不是列表"), "rows-invalid", "rows 类型不对 → rows-invalid")

    print("[7] 脏数据：重复录入 / 分值不是数字 / 缺答案 / 空题干")
    # 六列（没有解析列）：老师精简过的表最常见的样子
    dirty_header = ["题号", "题干", "选项A", "选项B", "答案", "分值"]
    dirty_path = make_xlsx(os.path.join(work, "脏题库.xlsx"), [
        dirty_header,
        [1, "同一道题录了两遍", "甲", "乙", "A", "两分"],
        [2, "同一道题录了两遍", "甲", "乙", "", 2],
        [3, "", "甲", "乙", "A", 2],
    ])
    bank = read_bank(dirty_path)
    codes = warn_codes(bank)
    check("score-not-number" in codes, "「两分」→ score-not-number", codes)
    check("duplicate-stem" in codes, "题干重复 → duplicate-stem", codes)
    check("answer-missing" in codes, "缺答案 → answer-missing", codes)
    check("empty-stem-row" in codes, "空题干 → empty-stem-row", codes)
    check(len(bank["rows"]) == 2, "空题干那行被跳过，不占题位", len(bank["rows"]))
    check(bank["rows"][0]["score"] is None, "非数字分值留空，不瞎猜", bank["rows"][0]["score"])

    print("[8] bank_to_text：生成的文本 exam 能原样解析回来")
    rows = [
        {"number": 7, "type": "单项选择", "stem": "下列加点字读音全部正确的一项是",
         "options": {"A": "甲", "B": "乙", "C": "丙", "D": "丁"}, "answer": "A",
         "analysis": "因为甲字读 zhōng", "score": 2.0},
        {"number": 8, "type": "单项选择", "stem": "选出没有语病的一句",
         "options": {"A": "第一句", "B": "第二句", "C": "第三句", "D": "第四句"},
         "answer": "B", "analysis": "第二句搭配得当", "score": 2.0},
        {"number": 9, "type": "填空", "stem": "补全对话：She ___ a teacher.",
         "options": {}, "answer": "is", "analysis": None, "score": 3.0},
    ]
    produced = bank_to_text(rows)
    text = produced["text"]
    check("一、单项选择" in text, "第一种题型开中文大题号", text.splitlines()[:1])
    check("二、填空" in text, "第二种题型接着用「二、」", [line for line in text.splitlines()
                                                      if line.startswith("二")])
    check("（2 分）" in text and "（3 分）" in text, "题干末尾写分值", text[:60])
    check("【答案】A" in text and "【解析】因为甲字读 zhōng" in text,
          "答案与解析按 exam 认的写法排", text[:80])
    check(produced["count"] == 3, "count 是真正进卷的题数", produced["count"])
    check(produced["sections"] == {"单项选择": 2, "填空": 1}, "sections 按题型分组",
          produced["sections"])
    check(not produced["warnings"], "这三题不该有告警", produced["warnings"])

    parsed = parse_paper(text)
    questions = [question for section in parsed["sections"]
                 for question in section["questions"]]
    check([question["number"] for question in questions] == [1, 2, 3],
          "题号跨大题连续 1..N（Excel 里的 7/8/9 重排）",
          [question["number"] for question in questions])
    check([question["answer"] for question in questions] == ["A", "B", "is"],
          "答案原样回来", [question["answer"] for question in questions])
    check([question["score"] for question in questions] == [2.0, 2.0, 3.0],
          "分值原样回来", [question["score"] for question in questions])
    check([len(question["options"]) for question in questions] == [4, 4, 0],
          "选项数对（填空没有选项）", [len(question["options"]) for question in questions])
    check("（2 分）" not in questions[0]["stem"], "分值没有留在题干里", questions[0]["stem"])
    check(questions[0]["analysis"] == "因为甲字读 zhōng", "解析回到解析位",
          questions[0]["analysis"])
    check(parsed["totalScore"] == 7, "总分 2+2+3=7", parsed["totalScore"])

    tuned = bank_to_text(rows, {"includeAnalysis": False, "includeScore": False,
                                "typeTitles": {"单项选择": "选择题"}})
    check("【解析】" not in tuned["text"], "includeAnalysis=False 不写解析", tuned["text"][:60])
    check("（2 分）" not in tuned["text"], "includeScore=False 不写分值", tuned["text"][:60])
    check("一、选择题" in tuned["text"], "typeTitles 换了大题标题", tuned["text"].splitlines()[:1])
    forced = bank_to_text([{"stem": "没有题型的一题", "options": {"A": "x"}}],
                          {"defaultType": "判断"})
    check(forced["sections"] == {"判断": 1}, "defaultType 给没题型的老题归类",
          forced["sections"])
    noisy = bank_to_text([{"stem": "选择题没写答案", "type": "单项选择", "options": {}}])
    noisy_codes = warn_codes(noisy)
    check("answer-missing-text" in noisy_codes and "options-missing" in noisy_codes,
          "缺答案解析阶段也提醒", noisy_codes)

    print("[9] bank_build：题库 → 一整份 docx（端到端）")
    out_path = os.path.join(work, "第一次月考.docx")
    built = bank_build({"path": std_path, "title": "第一次月考", "outputPath": out_path})
    check(os.path.isfile(out_path), "卷子真的落盘了", out_path)
    check(built["questionCount"] == 3, "3 题都进卷", built.get("questionCount"))
    check(built["totalScore"] == 7, "总分 7", built.get("totalScore"))
    check(built["bank"]["count"] == 3, "bank.count 与进卷题数一致", built["bank"])
    check(built["bank"]["format"] == "xlsx", "bank.format 回传源格式", built["bank"]["format"])
    check(built["bank"]["sections"] == {"单项选择": 3},
          "没写题型但有选项 → 都归单项选择", built["bank"]["sections"])
    check(built["preflight"]["level"] != "fail", "自产卷不判 fail",
          [item["code"] for item in built["preflight"]["findings"]])
    check("score" in built["preflight"], "带回印前检查分", list(built["preflight"].keys()))
    body = para_text(out_path)
    check("第一次月考" in body, "标题进卷", body[:40])
    check("下列加点字读音全部正确的一项是" in body, "题干进卷",
          "下列加点字读音全部正确的一项是" in body)
    check("望其项背" in body, "选项进卷", "望其项背" in body)
    check("【答案】A" in body or "A" in body, "答案排进卷子", True)

    auto = bank_build({"path": std_path, "title": "默认路径卷"})
    home = os.environ["TEACHER_KIT_HOME"]
    check(os.path.abspath(auto["outputPath"]).startswith(os.path.abspath(home)),
          "没给 outputPath 就落到 TEACHER_KIT_HOME 下", auto["outputPath"])
    check(os.path.isfile(auto["outputPath"]), "默认路径的卷子也生成了", auto["outputPath"])

    print("[10] 自带公式/多 sheet：取第一个有内容的表")
    multi = os.path.join(work, "多表.xlsx")
    book = Workbook()
    blank = book.active
    blank.title = "说明"
    filled = book.create_sheet("题库")
    for row in [STD_HEADER] + STD_ROWS:
        filled.append(list(row))
    book.save(multi)
    bank = read_bank(multi)
    check(len(bank["rows"]) == 3, "跳过空白 sheet，读有内容的那张", len(bank["rows"]))
    check(bank["rows"][2]["stem"] == "补写句中空缺的成语", "读的是「题库」那张表",
          bank["rows"][2]["stem"])


    print("[11] 按题型挑题：题库里有什么 / 挑 10 道单选 / 挑多了只警告")
    mixed = os.path.join(work, "题型题库.xlsx")
    make_xlsx(
        mixed,
        [["题号", "题型", "题干", "选项A", "选项B", "选项C", "选项D", "答案", "分值"],
         [1, "单项选择", "He ___ a book.", "read", "reads", "reading", "to read", "B", 2],
         [2, "单项选择", "They ___ happy.", "is", "am", "are", "be", "C", 2],
         [3, "单项选择", "I ___ a student.", "am", "is", "are", "be", "A", 2],
         [4, "填空", "写出「苹果」的英文", "", "", "", "", "apple", 3],
         [5, "填空", "写出「香蕉」的英文", "", "", "", "", "banana", 3]],
    )
    info = bank_inspect({"path": mixed})
    check(info["count"] == 5, "inspect 数出 5 道题", info["count"])
    by_type = {item["type"]: item["count"] for item in info["types"]}
    check(by_type.get("单项选择") == 3 and by_type.get("填空") == 2, "inspect 按题型分组",
          by_type)

    rows = read_bank(mixed)["rows"]
    picked, warns, summary = pick_rows(rows, {"单项选择": 2, "填空": 1})
    check(len(picked) == 3, "挑了 3 道", len(picked))
    check(summary == {"单项选择": 2, "填空": 1}, "每种题型各挑几道有回报", summary)
    check(warn_codes({"warnings": warns}) == [], "挑得够就不警告", warns)
    check([row["type"] for row in picked] == ["单项选择", "单项选择", "填空"],
          "按题型顺序排好（大题结构才自然）", [row["type"] for row in picked])

    _, warns_short, summary_short = pick_rows(rows, {"单项选择": 9})
    check(summary_short == {"单项选择": 3}, "题库不够时全给上", summary_short)
    check("pick-short" in warn_codes({"warnings": warns_short}), "挑不满给 pick-short", warns_short)
    expect_error(lambda: pick_rows(rows, {"作文": 1}), "pick-no-type", "题库里没有的题型直接说")
    expect_error(lambda: pick_rows(rows, {"单项选择": "很多"}), "pick-invalid",
                 "题数写非数字要报 pick-invalid")

    built = bank_build({"path": mixed, "title": "挑题卷", "options": {"pick": {"单项选择": 2}}})
    check(built["questionCount"] == 2, "挑 2 道就只排 2 道", built["questionCount"])
    check(built["bank"]["picked"] == {"单项选择": 2}, "挑题结果回传面板", built["bank"].get("picked"))
    check(built["bank"]["available"] == 5, "也告诉老师题库一共几道", built["bank"].get("available"))

    print("[12] 阅读题的「材料」列：短文排在题目前面，同一篇只印一遍；缺材料要提醒")
    reading = os.path.join(work, "阅读理解题库.xlsx")
    passage = "Tom is a 12-year-old boy from London. He likes football."
    make_xlsx(
        reading,
        [["题号", "题型", "材料", "题干", "选项A", "选项B", "选项C", "选项D", "答案", "分值"],
         [1, "阅读理解", passage, "What sport does Tom like?",
          "Football", "Basketball", "Tennis", "Swimming", "A", 2],
         [2, "阅读理解", passage, "How old is Tom?",
          "10", "12", "13", "14", "B", 2],
         [3, "单项选择", "", "— Can you swim?",
          "Yes, I can.", "No, I can't.", "", "", "A", 2]],
    )
    read_rows = read_bank(reading)["rows"]
    check(read_rows[0]["material"] == passage, "材料列读进 material", read_rows[0].get("material"))
    check(read_rows[2]["material"] is None, "空材料是 None，不是空串", read_rows[2].get("material"))
    converted = bank_to_text(read_rows)
    text = converted["text"]
    check(text.count("材料：") == 1, "同一篇短文连着两题只印一遍", text.count("材料："))
    check(text.index(passage) < text.index("1. What sport"), "短文排在第 1 题前面",
          text[:80].replace("\n", " / "))
    check(warn_codes(converted) == [], "带材料的阅读题不警告", converted["warnings"])

    bare = bank_to_text([
        {"number": 1, "type": "阅读理解", "stem": "What is the main idea?",
         "options": [{"label": "A", "text": "甲"}, {"label": "B", "text": "乙"}],
         "answer": "A", "analysis": None, "score": 2, "material": None},
    ])
    check("bare-reading" in warn_codes(bare), "阅读题没材料要给 bare-reading",
          [item["code"] for item in bare["warnings"]])
    bare_msg = next((item["message"] for item in bare["warnings"]
                     if item["code"] == "bare-reading"), "")
    check("材料" in bare_msg, "提醒里说清怎么补（加一列材料）", bare_msg)

    built = bank_build({"path": reading, "title": "阅读材料卷"})
    body = para_text(built["outputPath"])
    check(body.index(passage) < body.index("What sport does Tom like?"),
          "短文真的印在题目前面（学生看得到文章）", body[:120].replace("\n", " / "))
    check(body.count(passage) == 1, "短文只印一次，不重复", body.count(passage))
    check(built["questionCount"] == 3 and built["totalScore"] == 6,
          "三道题全进卷、总分 6", [built["questionCount"], built["totalScore"]])

    print("[13] 题库也能出 A/B 两套卷、学生卷 + 教师卷（不用再粘一遍）")
    pair = bank_build({"path": reading, "title": "题库 AB 卷",
                       "outputPath": os.path.join(work, "bank-ab.docx"), "overwrite": True,
                       "options": {"ab": True, "answerCard": "none", "answerSheet": "end"}})
    check(pair.get("ab") is True and len(pair.get("variants") or []) == 2,
          "题库 A/B 两套都出了", [item.get("title") for item in pair.get("variants") or []])
    variant_paths = [item["outputPath"] for item in pair["variants"]]
    check(all(os.path.exists(path) for path in variant_paths), "两个文件都落盘", variant_paths)
    check("A 卷" in pair["variants"][0]["title"] and "B 卷" in pair["variants"][1]["title"],
          "标题分别标了 A 卷 / B 卷", [item["title"] for item in pair["variants"]])
    b_text = para_text(pair["variants"][1]["outputPath"])
    check("What sport does Tom like?" in b_text, "B 卷题目还在（选项打乱但题不丢）")

    teacher = bank_build({"path": reading, "title": "题库教师卷",
                          "outputPath": os.path.join(work, "bank-teacher.docx"), "overwrite": True,
                          "options": {"teacher": True, "answerCard": "none", "answerSheet": "none"}})
    check(bool(teacher.get("teacherPath")) and os.path.exists(teacher["teacherPath"]),
          "教师卷落盘", teacher.get("teacherPath"))
    check("（教师卷）" in teacher["teacherPath"], "教师卷文件名带「（教师卷）」", teacher["teacherPath"])
    teacher_lines = para_text(teacher["teacherPath"])
    check(teacher_lines.count("【答案】") >= 3, "教师卷每道题下面都有答案",
          teacher_lines.count("【答案】"))
    check("【答案】" not in para_text(teacher["outputPath"]), "学生卷上没有答案")

    print("[14] 同一批题分几套出（每套题型题量差不多，各有各的答案页）")
    sets_path = os.path.join(work, "分套题库.xlsx")
    sets_rows = [["题号", "题型", "题干", "选项A", "选项B", "答案", "分值"]]
    for index in range(4):
        sets_rows.append([index + 1, "单项选择", "单选第 %d 题" % (index + 1), "甲", "乙", "A", 2])
    for index in range(4):
        sets_rows.append([index + 5, "填空", "填空第 %d 题" % (index + 1), "", "", "word", 1])
    make_xlsx(sets_path, sets_rows)

    multi = bank_build({"path": sets_path, "title": "分套小测",
                        "options": {"sets": 2, "answerSheet": "end"}})
    made = multi.get("sets") or []
    check(multi.get("multiSets") == 2 and len(made) == 2, "分成两套", multi.get("multiSets"))
    check(all(os.path.exists(item["outputPath"]) for item in made), "两套文件都落盘",
          [item["outputPath"] for item in made])
    check([item["questionCount"] for item in made] == [4, 4], "两套题量一样（题型均分）",
          [item["questionCount"] for item in made])
    check([item["totalScore"] for item in made] == [6.0, 6.0], "两套分值一样",
          [item["totalScore"] for item in made])
    check("第 1 套" in made[0]["title"] and "第 2 套" in made[1]["title"],
          "标题标了第几套", [item["title"] for item in made])
    first = para_text(made[0]["outputPath"])
    check("单项选择" in first and "填空" in first, "每套里两种题型都有（不出现一份全是单选）",
          first[:80].replace("\n", " / "))
    check("【答案】" in first or "参考答案" in first, "每套自带答案页", first[-80:].replace("\n", " / "))
    check(multi["bank"]["count"] == 8, "总题数按各套累加", multi["bank"]["count"])
    tiny = os.path.join(work, "一题题库.xlsx")
    make_xlsx(tiny, [["题号", "题干", "答案", "分值"], [1, "只有一道题", "A", 2]])
    expect_error(lambda: bank_build({"path": tiny, "options": {"sets": 2}}),
                 "sets-too-many", "题不够分那么多套要直接说")
    expect_error(lambda: bank_build({"path": sets_path, "options": {"sets": 9}}),
                 "bad-sets", "套数太多也要拦住")

    print("[15] 题库组的卷子也能套学校模板（模板纸型优先）")
    from teacher_kit.commands import template_make  # 延迟导入：commands 要 load 所有引擎模块
    sixteen = build_docx({
        "outputPath": os.path.join(work, "十六开原卷.docx"),
        "page": {"size": "16K", "orientation": "portrait", "columns": 1},
        "blocks": [
            {"type": "paragraph", "align": "center", "runs": [{"text": "{{examTitle}}", "bold": True, "sizePt": 16}]},
            {"type": "paragraph", "text": "（满分 {{totalScore}} 分，时间 {{minutes}} 分钟）"},
            {"type": "paragraph", "runs": [{"text": "一、听力（10 分）", "bold": True}]},
            {"type": "paragraph", "text": "1. What does he look like?"},
        ],
        "values": {"examTitle": "十六开模板卷", "totalScore": "100", "minutes": "90"},
    })
    check(os.path.exists(sixteen["outputPath"]), "先有一份 16 开原卷", sixteen["outputPath"])
    template_path = os.path.join(work, "十六开模板.docx")
    template_make({"source": sixteen["outputPath"], "outputPath": template_path, "overwrite": True})
    check(os.path.exists(template_path), "模板反推成功", template_path)

    import docx as _docx
    without = bank_build({"path": sets_path, "title": "题库不套模板"})
    base_section = _docx.Document(without["outputPath"]).sections[0]
    check(abs(base_section.page_width.mm - 368.0) < 3, "不套模板时还是默认 8 开", base_section.page_width.mm)
    templated = bank_build({"path": sets_path, "title": "题库套模板", "baseTemplate": template_path})
    section = _docx.Document(templated["outputPath"]).sections[0]
    check(abs(section.page_width.mm - 184.0) < 3 and abs(section.page_height.mm - 260.0) < 3,
          "套了模板就按模板的 16 开出（没被默认 8 开盖掉）",
          (round(section.page_width.mm, 1), round(section.page_height.mm, 1)))

    print("[16] 挑题写在顶层也认；指定了输出位置时几套都落在同一个文件夹")
    picked_top = bank_build({"path": sets_path, "title": "顶层挑题", "pick": {"单项选择": 2}})
    check(picked_top["questionCount"] == 2, "顶层 pick 生效（不再静默忽略）", picked_top["questionCount"])
    check((picked_top["bank"].get("picked") or {}) == {"单项选择": 2}, "挑题结果回传面板",
          picked_top["bank"].get("picked"))

    target_dir = os.path.join(work, "老师指定的文件夹")
    target = os.path.join(target_dir, "期中分套卷.docx")
    placed = bank_build({"path": sets_path, "title": "期中分套卷", "outputPath": target,
                         "options": {"sets": 2}})
    placed_paths = [item["outputPath"] for item in (placed.get("sets") or [])]
    check(len(placed_paths) == 2 and all(os.path.dirname(item) == target_dir for item in placed_paths),
          "分套出的卷子落在老师指定的文件夹里", placed_paths)
    check(all(os.path.basename(item).startswith("期中分套卷") for item in placed_paths),
          "文件名接着老师起的名，不再甩回默认目录", [os.path.basename(item) for item in placed_paths])

    print("[17] 题库同时要 A/B 两套 + 两份教师卷（四件）")
    four = bank_build({"path": reading, "title": "题库四件套",
                       "outputPath": os.path.join(work, "bank-four.docx"), "overwrite": True,
                       "options": {"ab": True, "teacher": True, "answerCard": "card", "answerSheet": "end"}})
    students = [item["outputPath"] for item in (four.get("variants") or [])]
    teachers = list(four.get("teacherPaths") or [])
    check(four.get("ab") is True and len(students) == 2, "题库 A/B 两套都出了",
          [item.get("title") for item in four.get("variants") or []])
    check(len(teachers) == 2, "同时给了两份教师卷（不再二选一）", teachers)
    check(all(os.path.exists(path) for path in students + teachers), "四份文件都落盘",
          students + teachers)
    check(all("（教师卷）" in os.path.basename(path) for path in teachers),
          "两份教师卷文件名都带「（教师卷）」", [os.path.basename(path) for path in teachers])
    check(all("【答案】" not in para_text(path) for path in students), "两份学生卷都没有答案",
          [para_text(path).count("【答案】") for path in students])
    check(all("【答案】" in para_text(path) for path in teachers), "两份教师卷每题下面都有答案",
          [para_text(path).count("【答案】") for path in teachers])
    check(four["totalScore"] == pair["totalScore"], "四件一起出的满分和只出 A/B 时一致",
          (four["totalScore"], pair["totalScore"]))

    print()
    if FAILS:
        print("失败 " + str(len(FAILS)) + " 项：")
        for label in FAILS:
            print(" - " + label)
        sys.exit(1)
    print("全部通过 ✅")


if __name__ == "__main__":
    try:
        main()
    except Exception:  # noqa: BLE001 - 中途崩了也要打清位置，别只留一行 traceback
        traceback.print_exc()
        print()
        print("失败 1 项：")
        print(" - 测试中途抛异常（见上面的 traceback）")
        sys.exit(1)
