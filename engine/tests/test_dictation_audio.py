"""听写音频（dictation_audio）的回归测试。

这个功能依赖 Windows 自带的语音包，所以：
- 没有 PowerShell / 没装语音的机器上，整组测试会被跳过（打印一条说明），不让它变成
  "别人 clone 下来就红" 的测试；
- 有语音的机器上真跑合成（2 条、停顿 1 秒），再用 wave 模块读回参数与时长核对。
"""

import os
import sys
import tempfile
import wave

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

from teacher_kit.dictation_audio import MAX_ITEMS, build_dictation_audio, list_voices  # noqa: E402
from teacher_kit.errors import EngineError  # noqa: E402

PASSED = 0
FAILED = 0
SKIPPED = []


def check(cond, label, detail=""):
    global PASSED, FAILED
    if cond:
        PASSED += 1
        print("  ✓ " + label)
    else:
        FAILED += 1
        print("  ✗ " + label + ("  ← " + str(detail) if detail else ""))


def expect_error(code, fn, label):
    try:
        fn()
    except EngineError as exc:
        check(exc.code == code, label, "实际 code=" + str(exc.code) + " message=" + str(exc.message))
        check(bool(exc.message), label + "（有中文说明）", exc.message)
        return
    check(False, label, "没有抛 EngineError")


def wav_info(path):
    with wave.open(path, "rb") as fh:
        return {
            "channels": fh.getnchannels(),
            "width": fh.getsampwidth(),
            "rate": fh.getframerate(),
            "seconds": round(fh.getnframes() / float(fh.getframerate()), 1),
        }


def main():
    tmp = tempfile.mkdtemp(prefix="tk-audio-")
    os.environ["TEACHER_KIT_HOME"] = tmp

    print("[1] 这台机器的语音包")
    try:
        voices = list_voices()
        usable = bool(voices)
    except EngineError as exc:
        voices, usable = [], False
        SKIPPED.append("list_voices: " + str(exc.message))
    check(isinstance(voices, list), "能列出已安装语音（返回列表）", type(voices).__name__)
    if not usable:
        print("  这台机器没有可用的语音包，跳过真合成测试（不影响其它测试）")
        print("")
        print("全部通过 ✅（有 " + str(len(SKIPPED)) + " 组跳过）" if FAILED == 0 else "失败 " + str(FAILED) + " 项")
        return 0 if FAILED == 0 else 1

    text = "1. apple 苹果\n2. look after 照顾"

    print("[2] 真合一段（2 条、念 1 遍、停 1 秒）")
    first = build_dictation_audio({
        "text": text,
        "title": "试做一",
        "repeat": 1,
        "pauseSec": 1,
        "leadSec": 0,
        "overwrite": True,
    })
    info = wav_info(first["outputPath"])
    check(os.path.exists(first["outputPath"]), "wav 落盘了", first["outputPath"])
    check(info["channels"] == 1, "单声道", info)
    check(info["width"] == 2, "16 bit", info)
    check(info["rate"] == 22050, "22.05kHz（听写够用、文件不大）", info)
    check(first["count"] == 2, "念了 2 条", first["count"])
    check(1.0 < first["seconds"] < 60, "时长合理（不是空文件也不是半小时）", first["seconds"])
    check(abs(info["seconds"] - round(first["seconds"], 1)) < 0.6, "报的时长与 wav 头一致", (info["seconds"], first["seconds"]))
    check(first["voices"]["en"] or first["voices"]["zh"], "记下了用的哪条语音", first["voices"])
    check(first["bytes"] > 10000, "文件不是空壳", first["bytes"])
    check(first["mode"] == "en2cn", "默认 en2cn", first["mode"])

    print("[3] 念 2 遍 + 停 2 秒 → 明显更长")
    second = build_dictation_audio({
        "text": text,
        "title": "试做二",
        "repeat": 2,
        "pauseSec": 2,
        "leadSec": 0,
        "overwrite": True,
    })
    check(second["seconds"] > first["seconds"] + 1.5, "念两遍比念一遍长", (first["seconds"], second["seconds"]))

    print("[4] 不报题号会更短")
    quiet = build_dictation_audio({
        "text": text,
        "title": "试做三",
        "repeat": 1,
        "pauseSec": 1,
        "leadSec": 0,
        "number": False,
        "overwrite": True,
    })
    check(quiet["seconds"] < first["seconds"], "不报题号时长变短（题号确实念了）", (quiet["seconds"], first["seconds"]))

    print("[5] cn2en：给中文、写英文")
    cn = build_dictation_audio({
        "text": text,
        "title": "试做四",
        "mode": "cn2en",
        "repeat": 1,
        "pauseSec": 1,
        "leadSec": 0,
        "overwrite": True,
    })
    check(cn["mode"] == "cn2en", "模式记成 cn2en", cn["mode"])
    check(cn["seconds"] > 0.5, "也出得来声音", cn["seconds"])

    print("[6] 参数与错误形状")
    expect_error("text-required", lambda: build_dictation_audio({}), "空输入报 text-required")
    expect_error("text-required", lambda: build_dictation_audio({"text": "   \n  \n"}), "全是空白报 text-required")
    expect_error("bad-pause", lambda: build_dictation_audio({"text": text, "pauseSec": 0}), "停顿 0 秒报 bad-pause")
    expect_error("bad-repeat", lambda: build_dictation_audio({"text": text, "repeat": 9}), "念 9 遍报 bad-repeat")
    expect_error("bad-mode", lambda: build_dictation_audio({"text": text, "mode": "xx"}), "模式写错报 bad-mode")
    expect_error("output-exists", lambda: build_dictation_audio({
        "text": text, "outputPath": first["outputPath"], "title": "重复",
    }), "同名音频不覆盖报 output-exists")

    print("[7] 条目上限")
    many = "\n".join("%d. word%d 词%d" % (i, i, i) for i in range(1, MAX_ITEMS + 6))
    capped = build_dictation_audio({
        "text": many, "title": "上限", "repeat": 1, "pauseSec": 1, "leadSec": 0, "overwrite": True,
    })
    check(capped["count"] == MAX_ITEMS, "超过 " + str(MAX_ITEMS) + " 条只念前 " + str(MAX_ITEMS) + " 条", capped["count"])
    check(any(w["code"] == "too-many" for w in capped["warnings"]), "并给了 too-many 提醒",
          [w["code"] for w in capped["warnings"]])

    print("[8] 和听写单同一套词条：标题/落款不念，count 对音频也生效（评测 P0-2）")
    from docx import Document

    from teacher_kit.handout import build_dictation

    words_14 = [("tall", "高的"), ("short", "矮的；短的"), ("thin", "瘦的"), ("heavy", "重的"),
                ("curly", "卷曲的"), ("straight", "直的"), ("medium", "中等的"),
                ("height", "身高；高度"), ("build", "体格；体型"), ("beard", "胡须"),
                ("glasses", "眼镜"), ("blonde", "金黄色的"), ("describe", "描述"),
                ("look", "看起来")]
    word_doc = Document()
    word_doc.add_paragraph("Unit 3 What does he look like? 单词表")
    word_table = word_doc.add_table(rows=len(words_14) + 1, cols=2)
    for row, (en, cn) in zip(word_table.rows, [("英文", "中文")] + words_14):
        row.cells[0].text = en
        row.cells[1].text = cn
    word_doc.add_paragraph("（Unit 3 单词表 · 七年级英语 · 王老师）")
    docx_path = os.path.join(tmp, "Unit3单词表.docx")
    word_doc.save(docx_path)

    paper = build_dictation({"source": docx_path, "outputPath": os.path.join(tmp, "同源听写单.docx")})
    whole = build_dictation_audio({
        "source": docx_path, "title": "词表全部", "repeat": 1, "pauseSec": 1, "leadSec": 0,
        "overwrite": True,
    })
    check(whole["count"] == paper["count"] == 14,
          "音频条数与听写单一致（都是 14 条，标题/落款都不念）", (paper["count"], whole["count"]))
    check(not any(w["code"] == "no-chinese" for w in whole["warnings"]),
          "不再出现「第 15 条只有英文」的警告", whole["warnings"])
    ten = build_dictation_audio({
        "source": docx_path, "title": "词表前十条", "count": 10, "repeat": 1, "pauseSec": 1,
        "leadSec": 0, "overwrite": True,
    })
    check(ten["count"] == 10, "count:10 时音频也只念 10 条", ten["count"])
    check(ten["seconds"] < whole["seconds"], "10 条确实比 14 条短（不是念完了才发现截断）",
          (ten["seconds"], whole["seconds"]))
    expect_error("bad-count", lambda: build_dictation_audio({"source": docx_path, "count": 0}),
                 "count:0 报 bad-count")

    print("")
    if FAILED == 0:
        note = "（有 " + str(len(SKIPPED)) + " 组跳过）" if SKIPPED else ""
        print("全部通过 ✅" + note)
        return 0
    print("失败 " + str(FAILED) + " 项")
    return 1


if __name__ == "__main__":
    sys.exit(main())
