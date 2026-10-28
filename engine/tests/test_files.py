#!/usr/bin/env python
"""文件加工：跑 python test_files.py。脏样本按老师真实遇到的样子造。"""

import io
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from xml.etree import ElementTree

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

from teacher_kit.commands import TABLE  # noqa: E402
from teacher_kit.errors import EngineError  # noqa: E402
from teacher_kit.files import docx_clean, docx_extract, docx_merge  # noqa: E402

FAILS = []


def check(condition, label, detail=""):
    if condition:
        print("  ✓ " + label)
    else:
        FAILS.append(label)
        print("  ✗ " + label + (" — " + str(detail) if detail else ""))


def expect_error(fn, code, label):
    try:
        fn()
    except EngineError as error:
        check(error.code == code, label, "拿到的是 " + error.code)
        return
    check(False, label, "没有抛错")


def dirty_document(path):
    """模拟"从网页复制进 Word 的一段话"：彩色、高亮、下划线、超链接、Calibri、大段间距、空段。"""
    document = docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("七年级英语复习资料")
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.paragraph_format.space_before = Pt(12)

    link_paragraph = document.add_paragraph()
    part = link_paragraph.part
    r_id = part.relate_to("https://example.com/unit3",
                          "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    hyperlink = link_paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    inner = link_paragraph.add_run("点这里看课件")
    hyperlink.append(inner._r)
    link_paragraph._p.append(hyperlink)

    document.add_paragraph("")
    document.add_paragraph("")
    body = document.add_paragraph("tall 高的　　short 矮的")
    for run in body.runs:
        run.font.name = "Calibri"
    document.save(path)
    return path


R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
COMMENTS_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"


def base_document(path, *paragraphs):
    document = docx.Document()
    for text in paragraphs or ("一、单选题　第 1 题",):
        document.add_paragraph(text)
    document.save(path)
    return path


def rewrite_package(source, destination, mutate):
    """老师手里的 docx 是别人机器上存的：直接改 zip 里的 XML，造出评审里的真实样本。"""
    with zipfile.ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    mutate(parts)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as package:
        for name, blob in parts.items():
            package.writestr(name, blob)
    return destination


def first_paragraph_xml(xml):
    return re.search(r"(<w:p\b[^>]*>)(.*?)(</w:p>)", xml, re.S)


def commented_sample(base, destination):
    """Word 里插了一条备课批注的样子：正文锚点 + word/comments.xml + 关系 + 内容类型。"""

    def mutate(parts):
        xml = parts["word/document.xml"].decode("utf-8")
        match = first_paragraph_xml(xml)
        marked = ('<w:commentRangeStart w:id="0"/>' + match.group(2) +
                  '<w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r>')
        parts["word/document.xml"] = (xml[: match.start(2)] + marked + xml[match.end(2):]).encode("utf-8")
        parts["word/comments.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:comments xmlns:w="' + W_NS + '">'
            '<w:comment w:id="0" w:author="张老师" w:date="2026-10-20T10:00:00Z">'
            "<w:p><w:r><w:t>参考答案：B</w:t></w:r></w:p></w:comment></w:comments>").encode("utf-8")
        rels = parts["word/_rels/document.xml.rels"].decode("utf-8")
        parts["word/_rels/document.xml.rels"] = rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdCmt9" Type="' + R_NS + '/comments" Target="comments.xml"/>'
            "</Relationships>").encode("utf-8")
        types = parts["[Content_Types].xml"].decode("utf-8")
        parts["[Content_Types].xml"] = types.replace(
            "</Types>",
            '<Override PartName="/word/comments.xml" ContentType="' + COMMENTS_TYPE + '"/>'
            "</Types>").encode("utf-8")

    return rewrite_package(base, destination, mutate)


def hyperlink_sample(base, destination, url="https://example.com/tiku"):
    """卷子里"（来源：某题库）"那种外部链接：w:hyperlink + 外部关系。"""

    def mutate(parts):
        xml = parts["word/document.xml"].decode("utf-8")
        match = first_paragraph_xml(xml)
        link = '<w:hyperlink r:id="rIdLink9" w:history="1">' + match.group(2) + "</w:hyperlink>"
        parts["word/document.xml"] = (xml[: match.start(2)] + link + xml[match.end(2):]).encode("utf-8")
        rels = parts["word/_rels/document.xml.rels"].decode("utf-8")
        parts["word/_rels/document.xml.rels"] = rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdLink9" Type="' + R_NS + '/hyperlink" Target="' + url +
            '" TargetMode="External"/></Relationships>').encode("utf-8")

    return rewrite_package(base, destination, mutate)


def numbered_sample(base, destination, num_id=77):
    """Word 列表自动编号写题号：正文 w:numPr 指向 numbering.xml 里的自定义 numId。"""

    def mutate(parts):
        xml = parts["word/document.xml"].decode("utf-8")
        match = re.search(r"(<w:p\b[^>]*>)", xml)
        num_pr = ('<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="' + str(num_id) +
                  '"/></w:numPr></w:pPr>')
        parts["word/document.xml"] = (xml[: match.end()] + num_pr + xml[match.end():]).encode("utf-8")
        numbering = parts["word/numbering.xml"].decode("utf-8")
        parts["word/numbering.xml"] = numbering.replace(
            "</w:numbering>",
            '<w:num w:numId="' + str(num_id) + '"><w:abstractNumId w:val="0"/></w:num>'
            "</w:numbering>").encode("utf-8")

    return rewrite_package(base, destination, mutate)


def dangling_rids(path):
    """包内有没有悬空引用：关系指向不存在的部件，或正文用了 rels 里没有的 rId。

    Word 对这两类都是同一句话——「文件可能已经损坏」（S1-1 的根因）。
    """
    problems = []
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        for name in sorted(names):
            if not name.endswith(".rels"):
                continue
            folder = posixpath.dirname(posixpath.dirname(name))
            owner = posixpath.join(folder, posixpath.basename(name)[: -len(".rels")])
            root = ElementTree.fromstring(package.read(name))
            known = set()
            for relationship in root:
                known.add(relationship.get("Id"))
                if relationship.get("TargetMode") == "External":
                    continue
                target = posixpath.normpath(posixpath.join(folder, relationship.get("Target") or ""))
                if target not in names:
                    problems.append(name + " 指向不存在的 " + str(relationship.get("Target")))
            if owner in names:
                xml = package.read(owner).decode("utf-8", "ignore")
                for rel_id in set(re.findall(r'r:(?:id|embed|link)="([^"]+)"', xml)):
                    if rel_id not in known:
                        problems.append(owner + " 用了不存在的 " + rel_id)
    return problems


# w:pPr 的子元素有固定顺序（CT_PPrBase）。顺序错了 Word 会说"文件需要修复"。
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]


def ppr_order_violations(path):
    """挑出 pPr 里子元素顺序放错的段落（分页标记、段间距都得插在合法位置）。"""
    document = docx.Document(path)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    violations = []
    for paragraph in paragraphs:
        ppr = paragraph._p.find(qn("w:pPr"))
        if ppr is None:
            continue
        names = [child.tag.split("}")[1] for child in ppr.iterchildren()]
        positions = [PPR_ORDER.index(name) for name in names if name in PPR_ORDER]
        if positions != sorted(positions):
            violations.append(names)
    return violations


def header_document(path, body_text, header_text=None, footer_text=None, header_image=None):
    """造一份「学校抬头各不相同」的 docx：页眉写校名，页脚写页码。

    header_image 给一张图（IO）时往页眉里塞图 —— 真实卷子的抬头常是扫描的校徽。
    传 None 的项就不建那个部件：缺页眉的文件本来就该是 linked（没有自己的抬头）。
    """
    document = docx.Document()
    section = document.sections[0]
    if header_text is not None or header_image is not None:
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        if header_text is not None:
            paragraph.text = header_text
        if header_image is not None:
            paragraph.add_run().add_picture(header_image)
    if footer_text is not None:
        footer = section.footer
        footer.is_linked_to_previous = False
        footer.paragraphs[0].text = footer_text
    document.add_paragraph(body_text)
    document.save(path)
    return path


def add_complex_page_field(paragraph):
    """手搓 Word 造页码域的样子：begin / instrText / separate / 结果 / end 散在五个 run 里。"""

    def marker(kind):
        run = paragraph.add_run()
        char = run._r.makeelement(qn("w:fldChar"), {})
        char.set(qn("w:fldCharType"), kind)
        run._r.append(char)
        return run

    marker("begin")
    run = paragraph.add_run()
    instruction = run._r.makeelement(qn("w:instrText"), {})
    instruction.text = " PAGE  \\* MERGEFORMAT "
    run._r.append(instruction)
    marker("separate")
    paragraph.add_run("3")
    marker("end")
    return paragraph


def section_header_text(section):
    return "\n".join(paragraph.text for paragraph in section.header.paragraphs).strip()


def section_footer_text(section):
    return "\n".join(paragraph.text for paragraph in section.footer.paragraphs).strip()


def main():
    work = tempfile.mkdtemp(prefix="files-")
    os.environ["TEACHER_KIT_HOME"] = os.path.join(work, "教师助手")

    print("[1] 净化：拆链接、去色/高亮、统一宋体、压空段、补段距（不覆盖原件）")
    dirty = dirty_document(os.path.join(work, "dirty.docx"))
    before = docx.Document(dirty)
    check(len(before.paragraphs) == 5, "原文档 5 段（含两个空段）", len(before.paragraphs))
    result = docx_clean({"path": dirty})
    stats = result["stats"]
    check(result["outputPath"] == os.path.abspath(os.path.join(work, "dirty_净化后.docx")),
          "默认另存「原名_净化后.docx」", result["outputPath"])
    check(len(docx.Document(dirty).paragraphs) == 5, "原件一个字没动（还是 5 段）")
    check("原文件已保留" in result["note"], "回执写明原件已保留", result.get("note"))
    check(stats["links"] == 1, "拆了 1 个超链接", stats)
    check(stats["emptyParagraphs"] == 1, "两个连续空段压成 1 个（只删 1 个）", stats)
    cleaned = docx.Document(result["outputPath"])
    check(len(cleaned.paragraphs) == 4, "剩 4 段（留下 1 个空段当呼吸）", [p.text for p in cleaned.paragraphs])
    text = "\n".join(paragraph.text for paragraph in cleaned.paragraphs)
    check("点这里看课件" in text, "超链接文字还在", text)
    check(not list(cleaned.element.body.iter(qn("w:hyperlink"))), "没有 w:hyperlink 了")
    runs = [run for paragraph in cleaned.paragraphs for run in paragraph.runs]
    check(all(run.font.color.rgb is None or str(run.font.color.rgb) == "000000" for run in runs), "颜色去掉了",
          [str(run.font.color.rgb) for run in runs])
    check(all(run.font.highlight_color is None for run in runs), "高亮去掉了")
    check(all(run.font.name in (None, "宋体") for run in runs), "字体统一成宋体",
          [run.font.name for run in runs])
    fonts = [element for element in cleaned.element.body.iter(qn("w:rFonts"))]
    check(fonts and all(element.get(qn("w:eastAsia")) == "宋体" for element in fonts), "eastAsia 也写了宋体",
          [element.get(qn("w:eastAsia")) for element in fonts])
    spacings = [element for element in cleaned.element.body.iter(qn("w:spacing"))]
    befores = [element.get(qn("w:before")) for element in spacings]
    afters = [element.get(qn("w:after")) for element in spacings]
    check("240" in befores, "原有段前 12 磅被保留（不再一律清零）", befores)
    check("360" in afters, "原有段后 18 磅被保留", afters)
    check("120" in afters, "没写段距的正文补了 6 磅段后", afters)

    print("[2] 净化到新文件 + 指定字体字号")
    copy = os.path.join(work, "clean-copy.docx")
    result = docx_clean({"path": dirty_document(os.path.join(work, "dirty2.docx")), "outputPath": copy,
                         "options": {"font": "楷体", "sizePt": 12}})
    check(os.path.exists(copy), "另存成功", copy)
    fonts = [element for element in docx.Document(copy).element.body.iter(qn("w:rFonts"))]
    check(fonts and all(element.get(qn("w:eastAsia")) == "楷体" for element in fonts), "字体是楷体")
    sizes = [element.get(qn("w:val")) for element in docx.Document(copy).element.body.iter(qn("w:sz"))]
    check(sizes and all(size == "24" for size in sizes), "12pt → w:sz 24", set(sizes))

    print("[3] 合并：两份带图片的文档 + 指定 8 开页面")
    first = os.path.join(work, "a.docx")
    document = docx.Document()
    document.add_paragraph("第一部分　知识点")
    document.add_paragraph("tall 高的")
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (24, 24), (200, 80, 80)).save(buffer, format="PNG")
    buffer.seek(0)
    document.add_picture(buffer)
    document.save(first)

    second = os.path.join(work, "b.docx")
    document = docx.Document()
    document.add_paragraph("第二部分　练习")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "题号"
    table.cell(0, 1).text = "答案"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "A"
    document.save(second)

    merged = os.path.join(work, "merged.docx")
    result = docx_merge({"inputs": [first, second], "outputPath": merged,
                         "page": {"size": "8K", "orientation": "landscape", "columns": 2},
                         "title": "复习卷合并"})
    check(result["images"] == 1, "图片关系搬过来了", result["images"])
    check(result["inputs"] == ["a.docx", "b.docx"], "回报了合并顺序", result["inputs"])
    merged_doc = docx.Document(merged)
    text = "\n".join(paragraph.text for paragraph in merged_doc.paragraphs)
    check("第一部分" in text and "第二部分" in text, "两段内容都在", text)
    check(len(merged_doc.tables) == 1, "表格还在", len(merged_doc.tables))
    check(abs(merged_doc.sections[0].page_width.mm - 368) < 2, "页面改成 8 开横向",
          merged_doc.sections[0].page_width.mm)
    check(len(merged_doc.inline_shapes) == 1, "图片是内联图（可读）", len(merged_doc.inline_shapes))
    check(merged_doc.core_properties.title == "复习卷合并", "标题属性写进去")
    packages = zipfile.ZipFile(merged).namelist()
    check(any(name.startswith("word/media/") for name in packages), "媒体部件在包里",
          [name for name in packages if "media" in name])
    check(len(packages) == len(set(packages)), "包内无重名部件")

    print("[4] 合并后的文字能再提取出来")
    extracted = docx_extract({"path": merged})
    check("tall 高的" in extracted["text"], "正文提取", extracted["text"][:60])
    check(extracted["tables"] and extracted["tables"][0][0] == ["题号", "答案"], "表格提取", extracted["tables"])
    check("题号 | 答案" in extracted["text"], "表格也进了纯文本", extracted["text"][-40:])

    print("[5] 错误形状")
    expect_error(lambda: docx_clean({}), "path-required", "没给文件报 path-required")
    expect_error(lambda: docx_clean({"path": os.path.join(work, "nope.docx")}), "file-not-found", "文件不存在")
    bad = os.path.join(work, "old.doc")
    open(bad, "wb").close()
    expect_error(lambda: docx_clean({"path": bad}), "unsupported-source", ".doc 报 unsupported-source")
    expect_error(lambda: docx_merge({"inputs": []}), "inputs-required", "没给 inputs")
    auto = docx_merge({"inputs": [first, second]})
    check(auto["outputPath"].startswith(os.path.join(work, "教师助手")), "没给路径就存到 文档/教师助手/文件加工",
          auto["outputPath"])

    print("[6] 命令表能按名字调到")
    extracted = TABLE["docx:extract"]({"path": first})
    check("第一部分" in extracted["text"], "docx:extract 可调")

    print("[7] 合并带批注/超链接的卷子（S1-1：Word 不再报「文件可能已经损坏」）")
    commented = commented_sample(base_document(os.path.join(work, "comment-base.docx"), "一、单选题　第 1 题"),
                                 os.path.join(work, "with-comment.docx"))
    linked = hyperlink_sample(base_document(os.path.join(work, "link-base.docx"), "（来源：某题库）"),
                              os.path.join(work, "with-link.docx"))
    merged_parts = os.path.join(work, "merged-parts.docx")
    result = docx_merge({"inputs": [commented, linked], "outputPath": merged_parts})
    check(result["comments"] == 1, "回报清掉 1 处批注", result.get("comments"))
    codes = [warning["code"] for warning in result["warnings"]]
    check("comments-dropped" in codes, "合并时说明批注没跟过来", codes)
    reopened = docx.Document(merged_parts)
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    check("一、单选题" in text and "来源：某题库" in text, "输入的文字都在", text)
    check(not dangling_rids(merged_parts), "包内没有悬空 rId", dangling_rids(merged_parts))
    names = zipfile.ZipFile(merged_parts).namelist()
    check(not [name for name in names if "comments" in name], "没有半截 comments 部件",
          [name for name in names if "comments" in name])
    check(not list(reopened.element.body.iter(qn("w:commentReference"))), "正文里没有残留批注锚点")
    links = [relationship for relationship in reopened.part.rels.values()
             if relationship.reltype == R_NS + "/hyperlink"]
    check(len(links) == 1 and links[0].target_ref == "https://example.com/tiku",
          "超链接地址重建到合并稿里", [(rel.rId, rel.target_ref) for rel in links])

    print("[8] 合并时每份之间插分页（S2-11）")
    page_a = base_document(os.path.join(work, "page-a.docx"), "第一份　试卷")
    page_b = base_document(os.path.join(work, "page-b.docx"), "第二份　答题页")
    merged_pages = os.path.join(work, "merged-pages.docx")
    result = docx_merge({"inputs": [page_a, page_b], "outputPath": merged_pages})
    xml = zipfile.ZipFile(merged_pages).read("word/document.xml").decode("utf-8")
    breaks = xml.count('<w:br w:type="page"/>') + len(re.findall(r"<w:pageBreakBefore[\s/>]", xml))
    check(breaks == 1, "两份之间正好 1 个分页", breaks)
    check(result["pageBreaks"] == 1, "回报插了 1 个分页", result.get("pageBreaks"))
    check(not ppr_order_violations(merged_pages), "分页标记插在 pPr 的合法位置（Word 不会提示修复）",
          ppr_order_violations(merged_pages))
    page_c = base_document(os.path.join(work, "page-c.docx"), "第三份　答案")
    result = docx_merge({"inputs": [page_a, page_b, page_c],
                         "outputPath": os.path.join(work, "merged-pages3.docx")})
    check(result["pageBreaks"] == 2, "三份之间 2 个分页", result.get("pageBreaks"))

    print("[9] 合并保留自动编号与自定义样式（S2-10）")
    numbered = numbered_sample(base_document(os.path.join(work, "number-base.docx"), "一、选择题"),
                               os.path.join(work, "numbered.docx"))
    merged_numbered = os.path.join(work, "merged-numbered.docx")
    docx_merge({"inputs": [numbered, page_a], "outputPath": merged_numbered})
    numbered_doc = docx.Document(merged_numbered)
    body_nums = [element.get(qn("w:val"))
                 for element in numbered_doc.element.body.iter(qn("w:numId"))]
    defined = {element.get(qn("w:numId"))
               for element in numbered_doc.part.numbering_part.element.findall(qn("w:num"))}
    check(len(body_nums) == 1, "自动编号的 w:numPr 跟过来了", body_nums)
    check(body_nums and body_nums[0] != "77", "源 numId 换成了新 id（不覆盖默认编号）", body_nums)
    check(set(body_nums) <= defined, "题号引用的 numId 在 numbering.xml 里有定义",
          (body_nums, sorted(defined)))
    check(not dangling_rids(merged_numbered), "编号合并后仍无悬空 rId", dangling_rids(merged_numbered))
    styled = os.path.join(work, "styled-src.docx")
    document = docx.Document()
    document.styles.add_style("卷首标题", WD_STYLE_TYPE.PARAGRAPH)
    heading = document.add_paragraph("2026 学年第一学期期末复习卷")
    heading.style = document.styles["卷首标题"]
    document.add_paragraph("一、单选题")
    document.save(styled)
    merged_styled = os.path.join(work, "merged-styled.docx")
    result = docx_merge({"inputs": [styled, page_a], "outputPath": merged_styled})
    styled_doc = docx.Document(merged_styled)
    check("卷首标题" in [style.name for style in styled_doc.styles], "自定义样式跨文档搬过来了",
          result.get("stylesMerged"))
    check(styled_doc.paragraphs[0].style.name == "卷首标题", "第一段样式没退回正文",
          styled_doc.paragraphs[0].style.name)

    print("[10] 两份都带同名图片：合并不撞车（S2-11）")
    picture_docs = []
    for index, color in enumerate(((200, 80, 80), (80, 200, 80))):
        path = os.path.join(work, "picture-%d.docx" % index)
        document = docx.Document()
        document.add_paragraph("图 %d" % index)
        buffer = io.BytesIO()
        Image.new("RGB", (20, 20), color).save(buffer, format="PNG")
        buffer.seek(0)
        document.add_picture(buffer)
        document.save(path)
        picture_docs.append(path)
    merged_pictures = os.path.join(work, "merged-pictures.docx")
    docx_merge({"inputs": picture_docs, "outputPath": merged_pictures})
    names = zipfile.ZipFile(merged_pictures).namelist()
    check(len(names) == len(set(names)), "包内无重名条目", [name for name in names if names.count(name) > 1])
    media = [name for name in names if name.startswith("word/media/")]
    check(len(media) == 2 and len(set(media)) == 2, "两张图各自落成独立部件", media)
    check(len(docx.Document(merged_pictures).inline_shapes) == 2, "两张内联图都能读",
          len(docx.Document(merged_pictures).inline_shapes))

    print("[11] 净化清批注（S2-8：备课批注不随文件发出去）")
    clean_comment = os.path.join(work, "clean-comment.docx")
    result = docx_clean({"path": commented, "outputPath": clean_comment})
    codes = [warning["code"] for warning in result["warnings"]]
    check("comments-removed" in codes, "warnings 有 comments-removed", codes)
    check(result["stats"]["comments"] == 1, "stats.comments=1", result["stats"])
    names = zipfile.ZipFile(clean_comment).namelist()
    check("word/comments.xml" not in names, "净化后 comments.xml 不存在",
          [name for name in names if "comments" in name])
    check(not list(docx.Document(clean_comment).element.body.iter(qn("w:commentReference"))),
          "净化后没有残留批注锚点")
    check(not dangling_rids(clean_comment), "净化产物无悬空 rId", dangling_rids(clean_comment))

    print("[12] 净化保留下划线（S2-9：划线题/填空横线是语义，不是花哨格式）")
    underlined = os.path.join(work, "underline.docx")
    document = docx.Document()
    run = document.add_paragraph().add_run("划线部分：cat")
    run.underline = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    document.save(underlined)
    underlined_out = os.path.join(work, "underline-default.docx")
    result = docx_clean({"path": underlined, "outputPath": underlined_out})
    check("underline-removed" not in [warning["code"] for warning in result["warnings"]],
          "默认净化不删下划线", result["warnings"])
    check(docx.Document(underlined_out).paragraphs[0].runs[0].underline, "下划线真的还在")
    check(result["stats"]["colors"] >= 1, "颜色照常去掉", result["stats"])
    check(not ppr_order_violations(underlined_out), "段间距插在 pPr 的合法位置",
          ppr_order_violations(underlined_out))
    forced = os.path.join(work, "underline-forced.docx")
    result = docx_clean({"path": underlined, "outputPath": forced, "options": {"keepUnderline": False}})
    message = " ".join(warning["message"] for warning in result["warnings"]
                       if warning["code"] == "underline-removed")
    check("underline-removed" in [warning["code"] for warning in result["warnings"]],
          "明确要求删时有 underline-removed", result["warnings"])
    check("下划线" in message and ("副本" in message or "keepUnderline" in message),
          "告警说清后果并给下一步", message)
    check(not docx.Document(forced).paragraphs[0].runs[0].underline, "keepUnderline=false 真的删掉了")

    print("[13] .doc / 加密 docx / 下载不完整 → 中文提示（S2-12）")

    def expect_friendly(fn, label):
        try:
            fn()
        except EngineError as error:
            check(error.code in ("unsupported-file", "unsupported-source"),
                  label + " 报 unsupported-file/source", error.code)
            check(re.search(r"[\u4e00-\u9fff]", error.message) is not None, label + " 提示是中文",
                  error.message)
            check(bool(error.hint), label + " 带一句下一步", error.hint)
            check(not re.search(r"[A-Za-z]{5,}Error|Traceback|PackageNotFound|BadZipFile", error.message),
                  label + " 没有把 Python 异常名抛出来", error.message)
            return
        check(False, label, "没有抛错")

    legacy = os.path.join(work, "legacy.doc")
    open(legacy, "wb").close()
    encrypted = os.path.join(work, "encrypted.docx")
    with open(encrypted, "wb") as stream:
        stream.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 256)
    intact = os.path.join(work, "intact.docx")
    with zipfile.ZipFile(intact, "w") as package:
        package.writestr("word/document.xml", "<w:document/>")
    with open(intact, "rb") as stream:
        blob = stream.read()
    truncated = os.path.join(work, "truncated.docx")
    with open(truncated, "wb") as stream:
        stream.write(blob[: len(blob) // 2])
    expect_friendly(lambda: docx_extract({"path": legacy}), "老 .doc 提文字")
    expect_friendly(lambda: docx_clean({"path": encrypted}), "加密 docx 净化")
    expect_friendly(lambda: docx_merge({"inputs": [encrypted], "outputPath": os.path.join(work, "x.docx")}),
                    "加密 docx 合并")
    expect_friendly(lambda: docx_extract({"path": truncated}), "下载不完整的 docx")

    print("[14] 净化默认不覆盖原件 / 原地写要 overwrite:true 且留备份（P0-5）")
    original = dirty_document(os.path.join(work, "inplace.docx"))
    result = docx_clean({"path": original})
    check(result["overwritten"] is False, "默认另存：overwritten=false", result["overwritten"])
    check(result.get("backup") is None, "没原地写就没有备份文件", result.get("backup"))
    check(result["changed"] is True, "changed 说的是内容真的变了", result["changed"])
    check(len(docx.Document(original).paragraphs) == 5, "默认净化后原件还是 5 段")
    expect_error(lambda: docx_clean({"path": original, "outputPath": original}),
                 "overwrite-required", "outputPath 指原文件又没给 overwrite → 拦下")
    check(len(docx.Document(original).paragraphs) == 5, "拦下之后原件也没动")
    result = docx_clean({"path": original, "overwrite": True})
    check(result["overwritten"] is True, "overwrite:true 才原地写", result["overwritten"])
    check(result["outputPath"] == os.path.abspath(original), "原地写的输出就是原路径", result["outputPath"])
    backup = result.get("backup")
    check(bool(backup) and os.path.exists(backup), "原地写留了备份", backup)
    check(os.path.basename(backup).startswith("inplace_原件备份"), "备份名是「原名_原件备份.docx」", backup)
    check(len(docx.Document(backup).paragraphs) == 5, "备份里是净化前的原件（5 段）")
    codes = [warning["code"] for warning in result["warnings"]]
    check("original-backed-up" in codes, "warning 说明备份在哪", codes)
    check("已原地净化" in result["note"], "note 说清原地写了、备份在哪", result.get("note"))
    check(len(docx.Document(original).paragraphs) == 4, "原地写后原文件变成净化版（4 段）")
    check(not ppr_order_violations(original), "原地写的段距插在 pPr 合法位置",
          ppr_order_violations(original))
    again = docx_clean({"path": original, "overwrite": True})
    check(again.get("backup") != backup, "第二次原地写不覆盖上一份备份", (backup, again.get("backup")))
    check(again["changed"] is False, "已经干净的文件再净化 changed=false（不再骗人）", again["stats"])

    print("[14a] 段距：keepSpacing 默认保留段前段后，只把行距拉到 1.5 倍（P1-9）")
    spaced = os.path.join(work, "spaced.docx")
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run("题目一行挨一行")
    document.save(spaced)
    result = docx_clean({"path": spaced})
    kept = docx.Document(result["outputPath"]).paragraphs[0].paragraph_format
    check(kept.space_before == Pt(6), "段前 6 磅原样保留", kept.space_before)
    check(kept.space_after == Pt(4), "段后 4 磅原样保留", kept.space_after)
    check(abs(float(kept.line_spacing) - 1.5) < 0.01, "单倍行距统一成 1.5 倍", kept.line_spacing)
    tight = docx_clean({"path": spaced, "outputPath": os.path.join(work, "spaced-tight.docx"),
                        "options": {"keepSpacing": False}})
    loose = docx.Document(tight["outputPath"]).paragraphs[0].paragraph_format
    check(not loose.space_before and not loose.space_after, "keepSpacing=false 才清零段前段后",
          (loose.space_before, loose.space_after))

    print("[14b] 空段：默认只压连续空段，stripEmpty 才全删（P1-9）")
    blanks = os.path.join(work, "blanks.docx")
    document = docx.Document()
    document.add_paragraph("一、单选题")
    for _index in range(4):
        document.add_paragraph("")
    document.add_paragraph("二、阅读")
    document.save(blanks)
    result = docx_clean({"path": blanks})
    texts = [paragraph.text for paragraph in docx.Document(result["outputPath"]).paragraphs]
    check(result["stats"]["emptyParagraphs"] == 3, "4 个连续空段压成 1 个（删 3 个）", result["stats"])
    check(len(texts) == 3 and texts[1] == "", "留下 2 段正文 + 1 个空段", texts)
    stripped = docx_clean({"path": blanks, "outputPath": os.path.join(work, "blanks-stripped.docx"),
                           "options": {"stripEmpty": True}})
    texts = [paragraph.text for paragraph in docx.Document(stripped["outputPath"]).paragraphs]
    check(texts == ["一、单选题", "二、阅读"], "stripEmpty=true 把空段全删了", texts)

    print("[14c] 页眉页脚：删网站广告、留学校抬头（P1-9）")
    ad_doc = os.path.join(work, "ad-header.docx")
    document = docx.Document()
    document.add_paragraph("一、单选题　第 1 题")
    header = document.sections[0].header
    header.paragraphs[0].text = "www.xxjyw.com 免费下载 请勿商用"
    header.add_paragraph("XX 中学 七年级英语组　第 1 页")
    document.sections[0].footer.paragraphs[0].text = "下载地址：https://tiku.example.com"
    document.save(ad_doc)
    result = docx_clean({"path": ad_doc})
    check(result["stats"]["ads"] == 2, "页眉 1 行 + 页脚 1 行广告都清掉", result["stats"])
    codes = [warning["code"] for warning in result["warnings"]]
    check("header-ad-removed" in codes, "warnings 有 header-ad-removed", codes)
    cleaned = docx.Document(result["outputPath"])
    header_text = "\n".join(paragraph.text for paragraph in cleaned.sections[0].header.paragraphs)
    footer_text = "\n".join(paragraph.text for paragraph in cleaned.sections[0].footer.paragraphs)
    check("www.xxjyw.com" not in header_text, "页眉里的网址没了", header_text)
    check("免费下载" not in header_text, "「免费下载 请勿商用」那行删掉了", header_text)
    check("XX 中学" in header_text, "学校抬头留着", header_text)
    check("下载地址" not in footer_text, "页脚里的下载地址没了", footer_text)
    check(len(cleaned.sections[0].header.paragraphs) >= 1, "页眉没被删空（Word 不会提示修复）")
    check(not dangling_rids(result["outputPath"]), "清完页眉后包内无悬空 rId",
          dangling_rids(result["outputPath"]))
    kept_ad = docx_clean({"path": ad_doc, "outputPath": os.path.join(work, "ad-kept.docx"),
                          "options": {"cleanHeader": False, "cleanFooter": False}})
    kept_text = "\n".join(paragraph.text
                          for paragraph in docx.Document(kept_ad["outputPath"]).sections[0].header.paragraphs)
    check("www.xxjyw.com" in kept_text, "cleanHeader=false 时不动页眉", kept_text)

    print("[14d] 超链接文字：连 Hyperlink 字符样式一起摘掉（P1-9）")
    link_style_doc = os.path.join(work, "link-style.docx")
    document = docx.Document()
    style = document.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    style.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    run = document.add_paragraph().add_run("（来源：某题库）")
    run.style = style
    document.save(link_style_doc)
    result = docx_clean({"path": link_style_doc})
    cleaned = docx.Document(result["outputPath"])
    check(cleaned.paragraphs[0].text == "（来源：某题库）", "文字还在", cleaned.paragraphs[0].text)
    check(cleaned.paragraphs[0].runs[0].font.color.rgb is None, "蓝字去掉了",
          cleaned.paragraphs[0].runs[0].font.color.rgb)
    check(not list(cleaned.element.body.iter(qn("w:rStyle"))), "Hyperlink 字符样式摘掉了",
          [element.get(qn("w:val")) for element in cleaned.element.body.iter(qn("w:rStyle"))])

    print("[14e] 本来就干净的文件：changed=false，不乱改")
    plain = os.path.join(work, "plain.docx")
    document = docx.Document()
    document.add_paragraph("一、单选题　第 1 题")
    document.save(plain)
    result = docx_clean({"path": plain,
                         "options": {"font": "", "normalizeSpacing": False,
                                     "removeEmptyParagraphs": False}})
    check(result["changed"] is False, "没东西可改时 changed=false", result["stats"])
    check(len(docx.Document(plain).paragraphs) == 1, "这种情形原件也没动")

    print("[15] 合并默认沿用第一份的页眉页脚（学校抬头别丢）")
    head_doc = os.path.join(work, "merge-head.docx")
    document = docx.Document()
    document.sections[0].header.is_linked_to_previous = False
    document.sections[0].header.paragraphs[0].text = "XX 市第三中学　英语组"
    document.add_paragraph("第一份正文")
    document.save(head_doc)
    tail_doc = os.path.join(work, "merge-tail.docx")
    document = docx.Document()
    document.add_paragraph("第二份正文")
    document.save(tail_doc)
    merged_path = os.path.join(work, "merge-heads.docx")
    merged = docx_merge({"inputs": [head_doc, tail_doc], "outputPath": merged_path, "overwrite": True})
    check(merged["headersInherited"] >= 1, "回报沿用了页眉", merged["headersInherited"])
    merged_doc = docx.Document(merged_path)
    head_text = merged_doc.sections[0].header.paragraphs[0].text
    check("第三中学" in head_text, "合并稿带着第一份的学校抬头", head_text)
    check(any("第一份正文" in item.text for item in merged_doc.paragraphs), "第一份正文还在")
    check(any("第二份正文" in item.text for item in merged_doc.paragraphs), "第二份正文还在")
    codes = [item.get("code") for item in merged["warnings"]]
    check("headers-inherited" in codes, "给了一条「沿用第一份页眉」的说明", codes)
    no_heads_path = os.path.join(work, "merge-heads-off.docx")
    docx_merge({"inputs": [head_doc, tail_doc], "outputPath": no_heads_path, "overwrite": True,
                "mergeHeaders": False})
    off_text = docx.Document(no_heads_path).sections[0].header.paragraphs[0].text
    check("第三中学" not in off_text, "选了不带页眉就不带", off_text)
    check(not dangling_rids(merged_path), "带页眉的合并件没有悬空 rId")

    print("[16] 逐份抬头：mergeHeaders='each' 每份文件各占一节，各带各的页眉页脚（评审待办）")
    school_a = header_document(os.path.join(work, "head-a.docx"), "A 中学　七年级英语组　练习册",
                               header_text="A 中学　七年级英语组", footer_text="第 1 页")
    school_b = header_document(os.path.join(work, "head-b.docx"), "B 小学　六年级语文组　卷子",
                               header_text="B 小学　六年级语文组", footer_text="第 2 页")
    each_path = os.path.join(work, "merge-each.docx")
    each = docx_merge({"inputs": [school_a, school_b], "outputPath": each_path, "overwrite": True,
                       "mergeHeaders": "each"})
    check(each["headersMode"] == "each", "回报 headersMode=each", each.get("headersMode"))
    check(each["headerSections"] == 2, "两节都带出了自己的抬头", each.get("headerSections"))
    each_doc = docx.Document(each_path)
    check(len(each_doc.sections) == 2, "两份文件各占一节（分节而不是只分页）", len(each_doc.sections))
    first_head = section_header_text(each_doc.sections[0])
    second_head = section_header_text(each_doc.sections[1])
    check("A 中学" in first_head and "B 小学" not in first_head, "第一节是 A 中学的抬头", first_head)
    check("B 小学" in second_head and "A 中学" not in second_head, "第二节是 B 小学的抬头", second_head)
    body = "\n".join(paragraph.text for paragraph in each_doc.paragraphs)
    check("A 中学　七年级英语组　练习册" in body and "B 小学　六年级语文组　卷子" in body, "两段正文都在", body)
    check(each["pageBreaks"] == 1, "分节符就是分页，不再多插一张空白页", each.get("pageBreaks"))
    check(not ppr_order_violations(each_path), "分节后 pPr 顺序没乱", ppr_order_violations(each_path))
    check(all(not section.footer.is_linked_to_previous for section in each_doc.sections),
          "每节页脚都断开了「同上一节」")
    check("第 2 页" in section_footer_text(each_doc.sections[1]), "第二节页脚是 B 小学自己的",
          section_footer_text(each_doc.sections[1]))
    codes = [warning["code"] for warning in each["warnings"]]
    check("headers-per-section" in codes, "说明每份文件保留了自己的抬头", codes)
    check("headers-inherited" not in codes, "逐份模式下不再说「沿用第一份」", codes)
    check(not dangling_rids(each_path), "逐份抬头后包内没有悬空 rId", dangling_rids(each_path))

    print("[17] 默认口径一个字没变：不传还是沿用第一份，只合成一节")
    default_path = os.path.join(work, "merge-each-default.docx")
    default = docx_merge({"inputs": [school_a, school_b], "outputPath": default_path, "overwrite": True})
    check(default["headersMode"] == "first", "不传 mergeHeaders 还是 first", default.get("headersMode"))
    check(default["headerSections"] == 0, "默认模式不按节带抬头", default.get("headerSections"))
    default_doc = docx.Document(default_path)
    check(len(default_doc.sections) == 1, "默认还是一节（老版面不变）", len(default_doc.sections))
    default_head = section_header_text(default_doc.sections[0])
    check("A 中学" in default_head and "B 小学" not in default_head, "默认只有第一份的抬头", default_head)
    truly = docx_merge({"inputs": [school_a, school_b], "outputPath": os.path.join(work, "merge-each-true.docx"),
                        "overwrite": True, "mergeHeaders": True})
    check(truly["headersMode"] == "first", "mergeHeaders=true 显式写也还是 first", truly.get("headersMode"))

    print("[18] mergeHeaderMode='each' 这个别名也认")
    alias_path = os.path.join(work, "merge-each-alias.docx")
    alias = docx_merge({"inputs": [school_a, school_b], "outputPath": alias_path, "overwrite": True,
                        "mergeHeaderMode": "each"})
    check(alias["headersMode"] == "each", "mergeHeaderMode='each' 解析成 each", alias.get("headersMode"))
    check(len(docx.Document(alias_path).sections) == 2, "别名走的是同一套分节逻辑",
          len(docx.Document(alias_path).sections))

    print("[19] 缺页眉的源文件：那一节必须是空白，不能继承上一份的校名")
    bare = base_document(os.path.join(work, "head-bare.docx"), "第三份正文（本来就没抬头）")
    blank_path = os.path.join(work, "merge-each-blank.docx")
    blank = docx_merge({"inputs": [school_a, bare], "outputPath": blank_path, "overwrite": True,
                        "mergeHeaders": "each"})
    check(blank["headerSections"] == 1, "只有真带抬头的那节算进 headerSections", blank.get("headerSections"))
    blank_doc = docx.Document(blank_path)
    check(len(blank_doc.sections) == 2, "没抬头的文件也占一节", len(blank_doc.sections))
    check(section_header_text(blank_doc.sections[1]) == "",
          "第二节页眉是空白（没印上一份的校名）", section_header_text(blank_doc.sections[1]))

    print("[20] 图片抬头：图片不带，那一节留空，不能留悬空 rId")
    from PIL import Image
    picture = io.BytesIO()
    Image.new("RGB", (24, 24), (80, 120, 200)).save(picture, format="PNG")
    picture.seek(0)
    pic_doc_path = header_document(os.path.join(work, "head-pic.docx"), "带校徽抬头的那份",
                                   header_image=picture)
    pic_path = os.path.join(work, "merge-each-pic.docx")
    pic = docx_merge({"inputs": [school_a, pic_doc_path], "outputPath": pic_path, "overwrite": True,
                      "mergeHeaders": "each"})
    check(not dangling_rids(pic_path), "图片抬头没带出悬空 rId", dangling_rids(pic_path))
    pic_merged = docx.Document(pic_path)
    check(not list(pic_merged.sections[1].header._element.iter(qn("w:drawing"))),
          "图片没搬过来（那一节页眉留空）")
    check(pic["headerSections"] == 1, "图片抬头那节按空白算", pic.get("headerSections"))
    check("带校徽抬头的那份" in "\n".join(paragraph.text for paragraph in pic_merged.paragraphs),
          "这份的正文照样在")

    print("[21] 复杂页码域：搬成 w:fldSimple，不留散架的 fldChar")
    field_doc = os.path.join(work, "head-field.docx")
    document = docx.Document()
    document.sections[0].header.is_linked_to_previous = False
    document.sections[0].header.paragraphs[0].text = "C 学校　英语组"
    footer = document.sections[0].footer
    footer.is_linked_to_previous = False
    add_complex_page_field(footer.paragraphs[0])
    document.add_paragraph("带复杂页码域的那份")
    document.save(field_doc)
    field_path = os.path.join(work, "merge-each-field.docx")
    docx_merge({"inputs": [school_a, field_doc], "outputPath": field_path, "overwrite": True,
                "mergeHeaders": "each"})
    field_footer = docx.Document(field_path).sections[1].footer._element
    simples = list(field_footer.iter(qn("w:fldSimple")))
    check(any("PAGE" in (element.get(qn("w:instr")) or "") for element in simples),
          "页码域变成 w:fldSimple 且带 PAGE", [element.get(qn("w:instr")) for element in simples])
    check(not list(field_footer.iter(qn("w:fldChar"))), "复杂域的 fldChar 记号没留下")

    print()
    if FAILS:
        print("失败 " + str(len(FAILS)) + " 项：")
        for label in FAILS:
            print(" - " + label)
        sys.exit(1)
    print("全部通过 ✅")


if __name__ == "__main__":
    main()