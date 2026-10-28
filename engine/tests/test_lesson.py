#!/usr/bin/env python
"""课件：跑 python test_lesson.py。样本按老师真实粘贴的样子写。"""

import os
import sys
import tempfile
import zipfile

from pptx import Presentation

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def template_with_sections(path):
    """造一个带示例页 + 分节 + 自定义放映的学校模板（评审 S2-7 的现场）。"""
    from lxml import etree

    presentation = Presentation()
    for title in ("单击输入标题", "示例页二", "示例页三"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "示例"
    presentation.save(path)
    with zipfile.ZipFile(path) as archive:
        items = {name: archive.read(name) for name in archive.namelist()}
    tree = etree.fromstring(items["ppt/presentation.xml"])
    id_list = tree.find("{%s}sldIdLst" % NS_P)
    ids = [node.get("id") for node in id_list]
    rel_ids = [node.get("{%s}id" % NS_R) for node in id_list]
    ext_lst = etree.SubElement(tree, "{%s}extLst" % NS_P)
    ext = etree.SubElement(ext_lst, "{%s}ext" % NS_P)
    ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    section_lst = etree.SubElement(ext, "{%s}sectionLst" % NS_P14)
    for name, chunk in (("模板第一节", ids[:2]), ("模板第二节", ids[2:])):
        section = etree.SubElement(section_lst, "{%s}section" % NS_P14)
        section.set("name", name)
        section.set("id", "{6B1F1B2C-2222-4A2A-9E33-00000000000%d}" % len(section_lst))
        sld_id_lst = etree.SubElement(section, "{%s}sldIdLst" % NS_P14)
        for slide_id in chunk:
            node = etree.SubElement(sld_id_lst, "{%s}sldId" % NS_P14)
            node.set("id", slide_id)
    cust_show_lst = etree.Element("{%s}custShowLst" % NS_P)
    cust_show = etree.SubElement(cust_show_lst, "{%s}custShow" % NS_P)
    cust_show.set("name", "只看重点")
    cust_show.set("id", "1")
    sld_lst = etree.SubElement(cust_show, "{%s}sldLst" % NS_P)
    for rel_id in rel_ids:
        node = etree.SubElement(sld_lst, "{%s}sld" % NS_P)
        node.set("{%s}id" % NS_R, rel_id)
    ext_lst.addprevious(cust_show_lst)
    items["ppt/presentation.xml"] = etree.tostring(
        tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in items.items():
            archive.writestr(name, blob)


def slide_refs(path):
    """读成品 pptx 的分节/自定义放映引用，用来查有没有指向已删页的悬空 sldId。"""
    from lxml import etree

    with zipfile.ZipFile(path) as archive:
        tree = etree.fromstring(archive.read("ppt/presentation.xml"))
        names = archive.namelist()
    ids = {node.get("id") for node in tree.findall("{%s}sldIdLst/{%s}sldId" % (NS_P, NS_P))}
    sections = [[node.get("id") for node in section.iter("{%s}sldId" % NS_P14)]
                for section in tree.iter("{%s}section" % NS_P14)]
    shows = [[node.get("{%s}id" % NS_R) for node in show.iter("{%s}sld" % NS_P)]
             for show in tree.iter("{%s}custShow" % NS_P)]
    return {
        "ids": ids,
        "sections": sections,
        "shows": shows,
        "has_section_lst": any(True for _ in tree.iter("{%s}sectionLst" % NS_P14)),
        "has_custom_show_lst": tree.find("{%s}custShowLst" % NS_P) is not None,
        "names": names,
    }

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

from teacher_kit.commands import TABLE  # noqa: E402
from teacher_kit.errors import EngineError  # noqa: E402
from teacher_kit.lesson import build_deck, plan_deck  # noqa: E402
from teacher_kit.pptx_build import build_pptx  # noqa: E402

FAILS = []


def check(condition, label, detail=""):
    if condition:
        print("  ✓ " + label)
    else:
        FAILS.append(label)
        print("  ✗ " + label + (" — " + str(detail) if detail else ""))


def expect_error(fn, code, label):
    try:
        fn()
    except EngineError as error:
        check(error.code == code, label, "拿到的是 " + error.code)
        return
    check(False, label, "没有抛错")


def title_only_pages(planned):
    """除封面外，只有标题、一条正文都没有的页——投影上老师只能站在原地翻页。"""
    return [slide.get("title") for slide in planned["slides"][1:]
            if not (slide.get("bullets") or [])]


BRACKET = """【学习目标】
1. 掌握本单元 5 个核心词汇的读音和拼写。
2. 能用 What does he look like? 询问外貌。
【新课导入】
- 看图片，说出你看到的人。
讲：这里放两张明星照片，学生先说，再带读词汇。
【词汇操练】
- tall / short / thin / heavy
- 用这四个词描述同桌。
【作业】
完成练习册 3-1，抄写单词三遍。
"""

MD = """# 七上 Unit 3 第一课时
## 学习目标
- 认识 6 个新单词
- 会读会写 three sentences
## 重点
- Where is it? It's on the desk.
"""

CN = """一、复习导入
上节课我们学了家庭成员。
今天继续学外貌描述。
二、新授
tall 高的
short 矮的
三、小结
今天我们学了四个形容词。
四、作业
抄写单词。
"""

# 老师从备课稿直接粘过来的样子：小标题光秃秃一行，既没有【】也没有 # 和编号（P1-8）。
# 以前这种稿子整篇认不出小标题，6 页全叫课题名，还一路「（续）」。
BARE_DECK = """Unit 3 What does he look like? Section A 1a—2c
学习目标
• 掌握 tall / short / thin / heavy 四个形容词的读音和拼写
• 能用 What does he look like? 询问并描述外貌
New words
• tall / short / thin / heavy / medium height
• 认读并抄写一遍
Practice
• 两人一组，互相描述同桌的外貌
• 用 He is tall and thin. 造句
Summary
• 今天我们学了四个描述外貌的形容词
Homework
• 抄写单词三遍，完成练习册 3-1
"""


def main():
    work = tempfile.mkdtemp(prefix="lesson-")
    os.environ["TEACHER_KIT_HOME"] = os.path.join(work, "教师助手")

    print("[1] 【】小标题 + 编号 + 项目符号 + 讲稿")
    planned = plan_deck({"text": BRACKET, "title": "What does he look like？",
                         "options": {"grade": "七年级", "subject": "英语", "teacher": "王老师"}})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles[0] == "What does he look like？", "封面是第一页", titles)
    check("学习目标" in titles and "词汇操练" in titles and "作业" in titles, "四个小标题都成页", titles)
    check(planned["slides"][0].get("subtitle") == "七年级 · 英语 · 王老师", "封面副标题由年级学科老师拼",
          planned["slides"][0].get("subtitle"))
    check(not planned["warnings"], "没有多余告警", planned["warnings"])
    goal_bullets = planned["slides"][1]["bullets"]
    check(len(goal_bullets) == 2, "目标页两条", goal_bullets)
    check(goal_bullets[0]["text"].startswith("掌握本单元"), "编号被去掉", goal_bullets[0]["text"])
    notes = [slide.get("notes") for slide in planned["slides"] if slide.get("notes")]
    check(any(note and "明星照片" in note for note in notes), "讲：写进了备注", notes)

    print("[2] markdown 标题")
    planned = plan_deck({"text": MD, "title": "Unit 3"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check("七上 Unit 3 第一课时" in titles, "# 标题成页", titles)
    check("学习目标" in titles and "重点" in titles, "## 标题成页", titles)
    target = [slide for slide in planned["slides"] if slide.get("title") == "学习目标"][0]
    check(len(target["bullets"]) == 2, "两个要点", target["bullets"])

    print("[3] 一、二、三、 中文编号；带句号的正文不能被当标题")
    planned = plan_deck({"text": CN, "title": "外貌"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check("复习导入" in titles and "新授" in titles and "小结" in titles and "作业" in titles, "四个中文标题",
          titles)
    review = [slide for slide in planned["slides"] if slide.get("title") == "复习导入"][0]
    texts = [bullet["text"] for bullet in review["bullets"]]
    check(len(texts) == 2, "标题行下的两句话都在", texts)
    check(any(text.endswith("。") for text in texts), "整句话原样保留", texts)

    print("[4] 一页最多 6 条，超了拆续页")
    long_text = "【知识点】\n" + "\n".join("- 要点 " + str(index) for index in range(1, 15))
    planned = plan_deck({"text": long_text, "title": "多要点"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles.count("知识点") + titles.count("知识点（续）") == 3, "14 条拆成 3 页", titles)
    check(any(w["code"] == "split-slide" for w in planned["warnings"]), "拆页给了提示", planned["warnings"])

    print("[5] 一个标题都没有 → 自动分页并提示加小标题")
    plain = "\n".join("第 " + str(index) + " 句话的内容。" for index in range(1, 11))
    planned = plan_deck({"text": plain, "title": "没标题的课"})
    check(len(planned["slides"]) >= 3, "按 6 条一页切开", len(planned["slides"]))
    check(any(w["code"] == "no-heading" for w in planned["warnings"]), "提示加小标题", planned["warnings"])

    print("[6] 出 pptx：16:9、中文落在微软雅黑、备注写进去")
    output = os.path.join(work, "deck.pptx")
    result = build_deck({"text": BRACKET, "outputPath": output, "title": "What does he look like？",
                         "options": {"grade": "七年级", "subject": "英语", "teacher": "王老师",
                                     "notes": "第一课时，40 分钟"}})
    check(os.path.exists(output), "文件落盘", output)
    presentation = Presentation(output)
    check(len(presentation.slides) == result["slides"], "页数对上", (len(presentation.slides), result["slides"]))
    check(abs(presentation.slide_width.mm - 338.7) < 1, "16:9 宽 338.7mm", presentation.slide_width.mm)
    cover = presentation.slides[0]
    check(cover.shapes.title.text_frame.text == "What does he look like？", "封面标题",
          cover.shapes.title.text_frame.text)
    check(cover.notes_slide.notes_text_frame.text == "第一课时，40 分钟", "封面备注")
    speak_slide = [slide for slide in presentation.slides
                   if slide.has_notes_slide and "明星" in slide.notes_slide.notes_text_frame.text]
    check(len(speak_slide) == 1, "讲稿落到对应页", len(speak_slide))
    fonts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.append(run.font.name)
    check(fonts and all(name == "微软雅黑" for name in fonts), "所有 run 都是微软雅黑", set(fonts))
    check(any(shape.has_text_frame and shape.text_frame.text for shape in presentation.slides[1].shapes),
          "第二页有正文")

    print("[7] 用学校模板：清掉示例页、保留模板尺寸与母版")
    template = os.path.join(work, "school-template.pptx")
    build_pptx({"outputPath": template, "slideSize": "4:3",
                "slides": [{"title": "单击输入标题", "bullets": ["示例"]}, {"title": "示例页二", "bullets": ["示例"]}]})
    result = build_deck({"text": MD, "outputPath": os.path.join(work, "deck-template.pptx"),
                         "title": "Unit 3", "template": template})
    check(result["template"] == os.path.abspath(template), "回报用了哪个模板", result["template"])
    baked = Presentation(os.path.join(work, "deck-template.pptx"))
    check(len(baked.slides) == result["slides"], "示例页没留下", (len(baked.slides), result["slides"]))
    check(abs(baked.slide_width.mm - 254.0) < 1, "沿用模板 4:3 宽 254mm", baked.slide_width.mm)
    texts = []
    for slide in baked.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    check(not any("单击输入标题" in text for text in texts), "模板示例文字没了", texts[:5])
    names = zipfile.ZipFile(os.path.join(work, "deck-template.pptx")).namelist()
    check(len(names) == len(set(names)), "包内没有重名部件（有重名 WPS 会报损坏）",
          [name for name in names if names.count(name) > 1][:5])

    print("[8] 错误形状")
    expect_error(lambda: plan_deck({"text": "  "}), "text-required", "空内容报 text-required")
    auto = build_deck({"text": MD, "title": "自动存"})
    check(auto["outputPath"].startswith(os.path.join(work, "教师助手")), "没给路径就存到 文档/教师助手/课件",
          auto["outputPath"])
    check(auto["outputPath"].endswith(".pptx") and os.path.exists(auto["outputPath"]), "真的生成了课件")
    expect_error(lambda: build_deck({"text": MD, "outputPath": os.path.join(work, "x.pptx"),
                                     "template": os.path.join(work, "nope.pptx")}),
                 "file-not-found", "模板不存在报 file-not-found")

    print("[9] 命令表能按名字调到")
    outline = TABLE["lesson:outline"]({"text": CN, "title": "外貌"})
    check(len(outline["outline"]) >= 5, "lesson:outline 返回大纲", outline["outline"])

    print("[10] 数字编号大纲：连续编号是有序列表，不是 N 个空小节（S1-2）")
    numbered = """Unit 3 Where did you go?
1. 学习 go 的过去式 went
2. 句型：Where did you go? I went to Sanya.
3. 单词：beach, hotel, delicious
4. 注意 e.g. 和 i.e. 的用法
5. 作业：抄写 P.45 的 10 个单词
6. 下节课听写
7. 预习 Section B"""
    planned = plan_deck({"text": numbered, "title": "Unit 3 Where did you go?"})
    content = planned["slides"][1:]
    check(len(planned["slides"]) <= 4, "7 条编号没有变成 7 页", len(planned["slides"]))
    check(all(slide.get("bullets") for slide in content), "没有只有标题的空页",
          [(slide.get("title"), len(slide.get("bullets") or [])) for slide in content])
    numbered_texts = [bullet["text"] for slide in content for bullet in slide.get("bullets") or []]
    check(len(numbered_texts) == 7, "7 条要点一条不少", numbered_texts)
    check(not any(text.startswith(("1.", "2.", "3.")) for text in numbered_texts),
          "投影上不顶着编号", numbered_texts[:3])

    print("[11] 常规英语课大纲：三节各自成页、讲稿进备注、长段落切成短要点")
    routine = """1. 学习目标
- 能听懂并会读本课 5 个新单词。
- 能用 What does he look like? 询问外貌。
讲：先放两张照片，学生先说，再带读。
2. 重点词汇
- tall / short / thin / heavy
讲：用这四个词描述同桌。
3. 句型操练
- What does he look like? He is tall and thin.
- 两人一组，互相描述外貌。
讲：抽查两组到前面表演。"""
    planned = plan_deck({"text": routine, "title": "What does he look like?"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles[1:4] == ["学习目标", "重点词汇", "句型操练"], "三节各自成页", titles)
    check(all(slide.get("bullets") for slide in planned["slides"][1:]),
          "常规大纲没有只有标题的空页",
          [(slide.get("title"), len(slide.get("bullets") or [])) for slide in planned["slides"][1:]])
    check(not [w for w in planned["warnings"] if w["code"] == "empty-slide"],
          "常规大纲不报 empty-slide", planned["warnings"])
    notes = [slide.get("notes") for slide in planned["slides"] if slide.get("notes")]
    check(len(notes) == 3 and any("照片" in note for note in notes), "讲稿进备注且不混进正文", notes)

    long_para = "光合作用是绿色植物通过叶绿体利用光能把二氧化碳和水转化成储存能量的有机物并释放氧气的过程" * 12
    planned = plan_deck({"text": "光合作用\n" + long_para + "\n叶绿体", "title": "光合作用"})
    long_texts = [bullet["text"] for slide in planned["slides"] for bullet in slide.get("bullets") or []]
    check(bool(long_texts) and max(len(text) for text in long_texts) <= 60,
          "整段 500+ 字切成 ≤60 字的要点，一条不塞一页",
          max(len(text) for text in long_texts) if long_texts else None)
    check("叶绿体" in long_texts and len(long_texts) >= 9, "长段落内容没丢", len(long_texts))

    print("[12] 真的只有标题的页要警告，不再静默")
    planned = plan_deck({"text": "【学习目标】\n【重点词汇】\n- tall / short", "title": "词汇课"})
    check(any(w["code"] == "empty-slide" for w in planned["warnings"]),
          "空页给 empty-slide 告警", planned["warnings"])

    print("[13] 模板清页：分节/自定义放映里不留悬空 sldId（S2-7）")
    sectioned = os.path.join(work, "sectioned-template.pptx")
    template_with_sections(sectioned)
    before = slide_refs(sectioned)
    check(before["has_section_lst"] and before["has_custom_show_lst"], "模板本身带分节与自定义放映", before)
    sectioned_deck = os.path.join(work, "sectioned-deck.pptx")
    result = build_deck({"text": "一、导入\n看图片", "title": "两页课",
                         "template": sectioned, "outputPath": sectioned_deck})
    check(result.get("removedTemplateSlides") == 3, "清掉 3 张示例页并回传",
          result.get("removedTemplateSlides"))
    check(result["slides"] == 2, "模板 3 页 → 出 2 页", result["slides"])
    after = slide_refs(sectioned_deck)
    dangling = [slide_id for section in after["sections"] for slide_id in section
                if slide_id not in after["ids"]]
    check(not dangling, "分节里没有指向已删页的 sldId", dangling)
    check(not after["has_section_lst"] or not after["sections"],
          "模板旧节名没有落到新页上", after["sections"])
    check(not after["has_custom_show_lst"] or not after["shows"],
          "自定义放映里没有悬空页", after["shows"])
    check(len(after["names"]) == len(set(after["names"])), "包内没有重名部件",
          [name for name in set(after["names"]) if after["names"].count(name) > 1][:5])
    kept = os.path.join(work, "sectioned-keep.pptx")
    build_pptx({"outputPath": kept, "template": sectioned, "clearTemplateSlides": False,
                "slides": [{"title": "新页", "bullets": ["a"]}]})
    check(slide_refs(kept)["sections"] == before["sections"], "保留示例页时分节原样不动",
          slide_refs(kept)["sections"])

    # [14] 光秃秃的小标题也得当页标题（P1-8）：学习目标 / New words / Practice
    print("[14] 裸小标题：学习目标 / New words / Practice")
    planned = plan_deck({"text": BARE_DECK, "title": "Unit 3 What does he look like?"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(planned["slides"][0]["layout"] == "title", "第一页还是封面",
          planned["slides"][0].get("layout"))
    for heading in ("学习目标", "New words", "Practice", "Summary", "Homework"):
        check(heading in titles, "「" + heading + "」自己成一页", titles)
    check(not any(str(title).endswith("（续）") for title in titles),
          "没有把每页都堆成「课题名（续）」", titles)
    projected = [bullet["text"] for slide in planned["slides"]
                 for bullet in (slide.get("bullets") or [])]
    check("学习目标" not in projected, "「学习目标」没被当成一条要点留在页里", projected[:3])
    check(not any(str(text).startswith("•") for text in projected),
          "投影文本里没有「• 学习目标」这种项目符号残留", projected[:3])
    goal = [slide for slide in planned["slides"] if slide.get("title") == "学习目标"][0]
    check(str(goal["bullets"][0]["text"]).startswith("掌握 tall"),
          "目标页第一条就是目标正文", [bullet["text"] for bullet in goal["bullets"]][:2])
    words = [slide for slide in planned["slides"] if slide.get("title") == "New words"][0]
    texts = [bullet["text"] for bullet in words["bullets"]]
    check(len(texts) == 2 and "medium height" in texts[0], "生词页两条要点一条不少", texts)
    practice = [slide for slide in planned["slides"] if slide.get("title") == "Practice"][0]
    check(len(practice["bullets"]) == 2, "操练页两条要点都在", practice["bullets"])
    check(not [warning for warning in planned["warnings"]
               if warning["code"] in ("no-heading", "empty-slide")],
          "认出小标题后不再报 no-heading / empty-slide", planned["warnings"])
    check(len(planned["outline"]) == len(planned["slides"]), "目录和页数对得上",
          len(planned["outline"]))

    # 同一小标题写重了：并成一页，不要各出一页各挂个「（续）」
    planned = plan_deck({"text": "学习目标\n- 掌握四个形容词\n学习目标\n- 能用 What does he look like? 问外貌\n",
                         "title": "重复小标题"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles == ["重复小标题", "学习目标"], "同名小标题写两遍只出一页", titles)
    merged = [bullet["text"] for bullet in planned["slides"][1]["bullets"]]
    check(len(merged) == 2, "两段要点并到同一页", merged)

    # 「（续）」只在要点真的超量、需要分页时才用
    planned = plan_deck({"text": "操练\n" + "\n".join("- 要点 " + str(index) for index in range(1, 15)),
                         "title": "多要点"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles[1] == "操练" and all(title == "操练（续）" for title in titles[2:]),
          "要点超量才拆「操练（续）」", titles)

    print("[15] 拆页要均衡：7 条别拆成「6 条 + 1 条孤零零」（persona P2）")
    seven = plan_deck({"text": "生词\ntall\nshort\nthin\nheavy\nmedium height\ncurly hair\nstraight hair",
                       "title": "Unit 3"})
    word_slides = [slide for slide in seven["slides"] if str(slide.get("title") or "").startswith("生词")]
    counts = [len(slide.get("bullets") or []) for slide in word_slides]
    check(len(counts) == 2 and min(counts) >= 2, "7 条拆成两页且没有单条孤页", counts)
    check(sum(counts) == 7, "7 条一条不少", sum(counts))
    check(counts[0] > counts[1] or counts[0] == counts[1], "前页不少于后页（读起来顺）", counts)

    print("[16] 活动名不是小标题：Pair work 留在操练页，不再多一页空「操练」（persona P1-4）")
    # 老师真实稿子（Unit3-课件稿.txt）的操练段：「- Pair work：描述你的同桌」曾被当成一节标题，
    # 于是「操练」单独成一页、正文全被搬到下一页「Pair work」。
    drill_deck = """【操练】
- Pair work：描述你的同桌
- 听力 1b：圈出听到的单词
- Group work：四人一组编对话
- Pairwork：给同桌描述家人的照片
【小结】
- is + 形容词
- has + 名词"""
    planned = plan_deck({"text": drill_deck, "title": "Unit 3 What does he look like?"})
    titles = [slide.get("title") for slide in planned["slides"]]
    folded = [str(title).casefold() for title in titles]
    check(not any(word in title for title in folded
                  for word in ("pair work", "pairwork", "group work")),
          "Pair work / Group work / Pairwork 都没有单独成页", titles)
    check(titles == ["Unit 3 What does he look like?", "操练", "小结"],
          "4 段操练内容都留在「操练」这一页里", titles)
    drill_texts = [bullet["text"] for bullet in planned["slides"][1]["bullets"]]
    check(len(drill_texts) == 4, "操练页 4 条要点一条不少", drill_texts)
    check(any("pair work" in text.casefold() for text in drill_texts),
          "操练页正文里有「pair work」", drill_texts)
    check("Pair work：描述你的同桌" in drill_texts, "这条要点的原文一个字没改", drill_texts)
    check(title_only_pages(planned) == [], "没有任何一页只有标题没内容", title_only_pages(planned))
    check(not [warning for warning in planned["warnings"] if warning["code"] == "empty-slide"],
          "不再报 empty-slide 空页告警", planned["warnings"])
    body_total = sum(len(slide.get("bullets") or []) for slide in planned["slides"][1:])
    check(body_total == 6, "合计正文行数不变（操练 4 + 小结 2）", body_total)

    # 光秃秃的活动名独立一行、下面跟要点：也只算上一节的活动要点，不升页
    planned = plan_deck({"text": "热身\n- 看照片猜人\nPair work\n- 两人互相描述\nGuessing game\n- 猜猜是谁",
                         "title": "Unit 3"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles == ["Unit 3", "热身"], "Pair work / Guessing game 独立一行也不升页", titles)
    texts = [bullet["text"] for bullet in planned["slides"][1]["bullets"]]
    check(texts == ["看照片猜人", "Pair work", "两人互相描述", "Guessing game", "猜猜是谁"],
          "活动名按原顺序留在热身页当要点", texts)

    # 真的小标题（导入 / 操练 独立一行、下面有正文）照样各成一页
    planned = plan_deck({"text": "导入\n看图片，说出你看到的人。\n操练\n- 两人一组对话\n- 抽查两组表演",
                         "title": "Unit 3"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles == ["Unit 3", "导入", "操练"], "「导入」「操练」独立一行仍然各自成页", titles)
    counts = [len(slide.get("bullets") or []) for slide in planned["slides"][1:]]
    check(counts == [1, 2], "这两页的正文都在", counts)
    check(title_only_pages(planned) == [], "这两页都不是空标题页", title_only_pages(planned))

    # 某一页最后只剩标题（不管什么原因）：并回上一页，不许单独留页
    planned = plan_deck({"text": "【导入】\n看图片，说出你看到的人。\n【小结】", "title": "Unit 3"})
    titles = [slide.get("title") for slide in planned["slides"]]
    check(titles == ["Unit 3", "导入"], "只有标题的「小结」没有单独成页", titles)
    texts = [bullet["text"] for bullet in planned["slides"][1]["bullets"]]
    check(texts == ["看图片，说出你看到的人。", "小结"],
          "并回上一页后标题以要点形式留着，一条不丢", texts)
    check(not [warning for warning in planned["warnings"] if warning["code"] == "empty-slide"],
          "并页之后不报空页", planned["warnings"])

    print()
    if FAILS:
        print("失败 " + str(len(FAILS)) + " 项：")
        for label in FAILS:
            print(" - " + label)
        sys.exit(1)
    print("全部通过 ✅")


if __name__ == "__main__":
    main()