#!/usr/bin/env python
"""教案引擎的验收脚本：跑 "python test_plan.py"。

样本都按老师真实粘贴的样子写：有的用「一、教学目标」，有的用「【教学目标】」，
有的是从微信/公众号里复制出来的、前后带着一堆没头没尾的行。
断言一律打在成品 docx 上（表格数量、列数、单元格文字、页眉页脚），
因为老师能不能交检，看的是那份文件，不是函数返回的 dict。
"""

import os
import sys
import tempfile

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

from teacher_kit.errors import EngineError  # noqa: E402
from teacher_kit.plan import build_plan, parse_plan, plan_fields  # noqa: E402

FAILS = []


def check(condition, label, detail=""):
    if condition:
        print("  ✓ " + label)
    else:
        print("  ✗ " + label + ("　→ " + str(detail) if detail else ""))
        FAILS.append(label)


def expect_error(fn, code, label):
    try:
        fn()
    except EngineError as error:
        check(error.code == code, label + "（code=" + code + "）", "实际 code=" + error.code)
        return error
    except Exception as error:  # noqa: BLE001
        check(False, label + "（code=" + code + "）", "抛了别的异常：" + repr(error))
        return None
    check(False, label + "（code=" + code + "）", "没有抛异常")
    return None


# ── 样本 ──────────────────────────────────────────────────────────────────
# 最标准的一份：人教版七年级英语，八块齐全，教学过程用「（一）导入」分环节
SAMPLE = """课题：What does he look like?
课型：新授课
课时：1 课时
班级：七年级（3）班
授课教师：王老师
学情分析：学生已掌握 tall、short 等基础词，但对 thin、heavy 不熟。
一、教学目标
1. 知识目标：掌握 tall、short、thin、heavy 四个形容词的读音与拼写。
2. 能力目标：能用 What does he look like? 询问并描述外貌。
3. 情感目标：学会欣赏他人，不以外貌取人。
二、教学重点
掌握四个形容词的读音与拼写。
三、教学难点
在真实情境中正确使用四个形容词描述他人外貌。
四、教学准备
单词卡片、人物图片、多媒体课件。
五、教学过程
（一）导入（5 分钟）
教师：出示两张明星照片，提问 Who is he?
学生：自由回答。
（二）新授
教师：带读 tall / short / thin / heavy。
学生：跟读，开火车读。
（三）操练
两人一组，互相描述同桌的外貌。
（四）小结
教师：今天我们学了四个描述外貌的形容词。
六、板书设计
tall   short   thin   heavy
What does he look like? He is tall.
七、作业布置
抄写单词三遍，完成练习册 3-1。
八、教学反思
学生参与度高，但 thin 的读音还需纠正。
"""

# 方括号、中英文冒号、环节一/Step 混着写 —— 从别处复制过来最常见的脏样子
BRACKET = """【课题】Unit 3 Where did you go?
【课型】: 复习课
[课时]：2 课时
[班级]: 九年级（1）班
授课教师: 李老师
一、教学目标
1. 复习 go/went 的用法。
2. 能用过去式谈论周末。
环节一：导入
看图片，说出去过的地方。
Step 2: Practice
两人一组对话。
环节三：小结
回顾本节课的重点。
"""

# 开头几行是复制时带进来的废话，不属于任何一块
NOISY = """今天这节课是复习课。

课题：My Weekend

一、教学目标
能用过去式谈论周末。
"""

DUP = """课题：重复测试
一、教学目标
1. 第一条目标。
二、教学重点
重点内容。
三、教学目标
2. 第二条目标（老师又补了一句）。
四、教学反思
"""

ROUTING = """一、教学过程
（一）热身（3 分钟）
教师：播放歌曲。
学生：跟唱并做动作。
（二）新授
教师：讲解 go 的过去式。
学生：做笔记。
时间：10 分钟
"""


# 教研组最常见的写法：环节名 + 括号时长，前面既没有「一、」也没有冒号（P1-7）。
# 以前这种写法一个环节都认不出，整段过程挤成一格，还照样出成品。
BARE_SECTIONS = """课题：What does he look like?
课型：新授课
一、教学目标
1. 能用 What does he look like? 描述外貌。
五、教学过程
导入（5 分钟）
教师：出示两张明星照片。
学生：自由回答。
新授(10min)
教师：带读 tall / short / thin / heavy。
学生：跟读。
操练 12 分钟
两人一组，互相描述同桌的外貌。
小结（3 分钟）
教师：总结今天学的四个形容词。
"""


# 一个「学生活动」都没写的教案：语法上完全正常（老师常常只写教师这边）
BLANK_STUDENT = """课题：My Family
课型：新授课
课时：1 课时
一、教学目标
能用 This is my... 介绍家人。
二、教学过程
（一）导入
教师：出示全家福照片。
（二）新授
教师：带读 family、father、mother。
（三）小结
教师：总结本课单词。
"""


def load(path):
    return docx.Document(path)


def find_table(document, first_cell):
    """按第一格文字找表：教案里每张表的第一格就是字段名，天然唯一。"""
    for table in document.tables:
        if table.cell(0, 0).text.strip().startswith(first_cell):
            return table
    return None


def body_paragraphs(document):
    return [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]


def footer_text(document):
    return "\n".join(paragraph.text for paragraph in document.sections[0].footer.paragraphs)


def first_footer_field(document):
    paragraph = document.sections[0].footer.paragraphs[0]
    return paragraph._p.find(qn("w:fldSimple"))


def main():
    work = tempfile.mkdtemp(prefix="teacher-kit-plan-")
    os.environ["TEACHER_KIT_HOME"] = os.path.join(work, "教师助手")
    print("工作目录：" + work)

    def out(name):
        return os.path.join(work, name + ".docx")

    # [1] 标准教案：八块（加课题信息共 14 个字段）全命中
    print("[1] 标准教案字段全命中")
    parsed = parse_plan(SAMPLE)
    fields = parsed["fields"]
    for key, label in (("title", "课题"), ("lessonType", "课型"), ("periods", "课时"),
                       ("className", "班级"), ("teacher", "授课教师"),
                       ("studentAnalysis", "学情分析"), ("objectives", "教学目标"),
                       ("keyPoints", "教学重点"), ("difficultPoints", "教学难点"),
                       ("preparation", "教学准备"), ("procedure", "教学过程"),
                       ("boardDesign", "板书设计"), ("homework", "作业布置"),
                       ("reflection", "教学反思")):
        check(key in fields, "认出了 " + label + "（" + key + "）", sorted(fields))
    check(fields.get("title") == "What does he look like?", "课题取值正确", fields.get("title"))
    check(fields.get("teacher") == "王老师", "授课教师取值正确", fields.get("teacher"))
    # 目标保留 1./2./3. 分层原文，不能被压成一行
    check("1. 知识目标" in fields.get("objectives", ""), "教学目标保留 1./2./3. 编号")
    check("情感目标" in fields.get("objectives", ""), "教学目标保留情感层")
    check(len(fields["objectives"].split("\n")) == 3, "教学目标仍是三行",
          fields.get("objectives"))
    check("教师：" not in fields["teacher"], "「教师：」没被当成授课教师字段",
          fields.get("teacher"))
    check([section["title"] for section in parsed["sections"]] == ["导入", "新授", "操练", "小结"],
          "教学过程分出 4 个环节", [section["title"] for section in parsed["sections"]])
    check(parsed["warnings"] == [], "标准教案没有提示", parsed["warnings"])

    # [2] 【】/[] 标记 + 中英文冒号 + 环节一/Step 变体
    print("[2] 方括号标记与冒号变体")
    parsed = parse_plan(BRACKET)
    fields = parsed["fields"]
    check(fields.get("title") == "Unit 3 Where did you go?", "【课题】认出", fields.get("title"))
    check(fields.get("lessonType") == "复习课", "【课型】: 认出（带英文冒号+空格）",
          fields.get("lessonType"))
    check(fields.get("periods") == "2 课时", "[课时]：认出（中文冒号）", fields.get("periods"))
    check(fields.get("className") == "九年级（1）班", "[班级]: 认出", fields.get("className"))
    check(fields.get("teacher") == "李老师", "授课教师: 认出", fields.get("teacher"))
    check("2. 能用过去式谈论周末。" in fields.get("objectives", ""), "两行目标都在",
          fields.get("objectives"))
    check([section["title"] for section in parsed["sections"]] == ["导入", "Practice", "小结"],
          "环节一 / Step 2 / 环节三 都认成环节", [section["title"] for section in parsed["sections"]])
    # 这一份没写「教学过程」四个字，直接从环节往下写：要能自动归位并提示一句
    check([warning["code"] for warning in parsed["warnings"]] == ["procedure-inferred"],
          "没写「教学过程」标题时按环节行自动归位", parsed["warnings"])
    check(parsed["warnings"] and "自动归到教学过程" in parsed["warnings"][0]["message"],
          "自动归位给的是中文提示", parsed["warnings"][:1])
    check(parsed["sections"][0]["lines"] == ["看图片，说出去过的地方。"],
          "环节正文跟着环节走", parsed["sections"][0]["lines"])

    # [3] 认不出的行进教学过程 + 空行不报
    print("[3] 认不出的行与空行")
    parsed = parse_plan(NOISY)
    warnings = parsed["warnings"]
    check(len(warnings) == 1, "只报 1 条（空行不算）", warnings)
    check(warnings[0]["code"] == "unknown-line", "提示类型是 unknown-line", warnings)
    check(warnings[0]["message"] == "第 1 行没认出属于哪一块，已放进教学过程",
          "提示文案是给老师看的中文", warnings[0]["message"] if warnings else "")
    check(parsed["fields"].get("title") == "My Weekend", "后面的字段照样认出来",
          parsed["fields"].get("title"))
    check("今天这节课是复习课。" in parsed["fields"].get("procedure", ""),
          "认不出的行没丢，进了教学过程", parsed["fields"].get("procedure"))
    check("能用过去式谈论周末。" in parsed["fields"].get("objectives", ""),
          "教学目标没被抢走", parsed["fields"].get("objectives"))

    # [4] form 版：表格数量 = 字段数 + 顶部表 + 板书框；教学过程 4 列
    print("[4] form 版结构")
    result = build_plan({"text": SAMPLE, "outputPath": out("form")})
    document = load(result["outputPath"])
    check(result["layout"] == "form", "默认版式是 form", result["layout"])
    check(result["fieldCount"] == 8, "正文排了 8 块字段（课时挪进抬头表了）", result["fieldCount"])
    check(len(document.tables) == result["fieldCount"] + 1 + 1,
          "表格数 = 字段数 + 顶部信息表 + 板书框", len(document.tables))
    header = document.tables[0]
    check(len(header.rows) == 3 and len(header.columns) == 4, "顶部信息表是 4 列 3 行",
          (len(header.rows), len(header.columns)))
    check(header.cell(0, 0).text == "课题", "信息表第一行是课题", header.cell(0, 0).text)
    check(header.cell(0, 1).text == "What does he look like?", "课题值进格子",
          header.cell(0, 1).text)
    span = header.cell(0, 1)._tc.tcPr.find(qn("w:gridSpan"))
    check(span is not None and span.get(qn("w:val")) == "3",
          "课题值横跨三列（长课题不折行）", span is not None and span.get(qn("w:val")))
    check(header.cell(1, 0).text == "课型" and header.cell(1, 2).text == "课时",
          "第二行是 课型|课时", (header.cell(1, 0).text, header.cell(1, 2).text))
    check(header.cell(1, 1).text == "新授课" and header.cell(1, 3).text == "1 课时",
          "课时和它的值在同一行（老师反馈过它「掉到下一行」）",
          (header.cell(1, 1).text, header.cell(1, 3).text))
    # 掉行的真实原因：值格子窄、12pt 正文尺寸放不下 → 折成两行。这里按格宽缩了字号
    check("\n" not in header.cell(1, 3).text, "课时值没折成两行", repr(header.cell(1, 3).text))
    size = header.cell(1, 3).paragraphs[0].runs[0].font.size
    check(size is not None and size.pt <= 11, "课时值字号收在格子宽度里", size and size.pt)
    check(header.cell(2, 0).text == "授课教师" and header.cell(2, 2).text == "班级",
          "第三行是 授课教师|班级", (header.cell(2, 0).text, header.cell(2, 2).text))
    check(header.cell(2, 1).text == "王老师" and header.cell(2, 3).text == "七年级（3）班",
          "教师与班级进格子", (header.cell(2, 1).text, header.cell(2, 3).text))
    check(find_table(document, "课时") is None,
          "课时不再另排成一块字段表（那就是「值掉到下一行」的观感来源）",
          [table.cell(0, 0).text for table in document.tables])
    labels = header.cell(0, 0).paragraphs[0].runs[0]
    check(labels.bold is True, "信息表字段名加粗", labels.bold)

    objectives = find_table(document, "教学目标")
    check(objectives is not None and len(objectives.columns) == 1 and len(objectives.rows) == 2,
          "教学目标是一块 1 列 2 行的表", objectives and (len(objectives.rows), len(objectives.columns)))
    check("知识目标" in objectives.cell(1, 0).text, "内容进第二行", objectives.cell(1, 0).text[:24])
    check(objectives.cell(0, 0).paragraphs[0].runs[0].bold is True, "字段名一行加粗")

    procedure = find_table(document, "环节")
    check(procedure is not None, "有教学过程 4 列表", procedure)
    check(len(procedure.columns) == 4, "教学过程表是 4 列", len(procedure.columns))
    check([procedure.cell(0, index).text for index in range(4)] ==
          ["环节", "教师活动", "学生活动", "时间"], "表头是 环节|教师活动|学生活动|时间",
          [procedure.cell(0, index).text for index in range(4)])
    check(len(procedure.rows) == 5, "4 个环节 + 表头 5 行", len(procedure.rows))
    check(procedure.cell(1, 0).text == "导入", "第一个环节是导入", procedure.cell(1, 0).text)
    check("明星照片" in procedure.cell(1, 1).text, "教师活动进第二格", procedure.cell(1, 1).text)
    check(procedure.cell(1, 2).text == "自由回答。", "学生活动进第三格", procedure.cell(1, 2).text)
    check(procedure.cell(1, 3).text == "5 分钟", "时间进第四格", procedure.cell(1, 3).text)
    # 学生活动留白是故意的，但必须让老师看得出「是给我填的」，而不是以为没生成
    check(result["table"] == {"rows": 4,
                              "teacher": {"filled": 4, "blank": 0},
                              "student": {"filled": 2, "blank": 2},
                              "time": {"filled": 1, "blank": 3}},
          "返回值里有教学过程表的填写统计", result["table"])
    notes = [paragraph for paragraph in body_paragraphs(document)
             if paragraph.text.startswith("注：")]
    check(len(notes) == 1 and "留给你" in notes[0].text,
          "表下一行小字说明空白是留给你填的", [item.text[:16] for item in notes])
    note_run = notes[0].runs[0] if notes else None
    note_rpr = note_run._r.rPr if note_run is not None else None
    note_color = note_rpr.find(qn("w:color")) if note_rpr is not None else None
    check(note_color is not None and (note_color.get(qn("w:val")) or "").upper() == "808080",
          "说明小字是灰的（打印出来是「注」，不抢正文）",
          note_color is not None and note_color.get(qn("w:val")))
    check(note_run is not None and note_run.font.size is not None
          and round(note_run.font.size.pt) == 9, "说明小字是 9pt",
          note_run and note_run.font.size)

    board = find_table(document, "板书设计")
    check(board is not None and len(board.rows) == 1 and len(board.columns) == 1,
          "板书是 1×1 带框格子", board and (len(board.rows), len(board.columns)))
    check("tall" in board.cell(0, 0).text, "板书内容进格子", board.cell(0, 0).text[:20])
    check(board.cell(0, 0).paragraphs[0].runs[0].bold is True, "板书标签加粗")

    section = document.sections[0]
    check(round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297,
          "A4 竖版", (round(section.page_width.mm), round(section.page_height.mm)))
    text = footer_text(document)
    check("第" in text and "页" in text, "默认页脚有页码", text)
    field = first_footer_field(document)
    check(field is not None and "PAGE" in (field.get(qn("w:instr")) or ""),
          "页脚页码是真域（不是死字）", field is not None and field.get(qn("w:instr")))
    check("{{" not in text, "页脚没有留下没替换的占位符", text)
    check("出卷老师" not in text and "授课教师" in text, "页脚是教案版（不是卷子版的「出卷老师」）",
          text)
    page_fields = [element for element in section.footer.paragraphs[0]._p
                   if element.tag == qn("w:fldSimple") and "PAGE" in (element.get(qn("w:instr")) or "")
                   and "NUMPAGES" not in (element.get(qn("w:instr")) or "")]
    check(len(page_fields) == 1, "页脚只有一个页码域（没有上一版页脚的残留）", len(page_fields))
    check(body_paragraphs(document)[0].text == "《What does he look like?》教案",
          "首行大标题套书名号", body_paragraphs(document)[0].text)
    check(result["bytes"] > 4000, "文件真的写出来了", result["bytes"])

    # [5] flow 版：没有字段表格，字段名以加粗段落出现，板书仍是 1×1
    print("[5] flow 版结构")
    result = build_plan({"text": SAMPLE, "layout": "flow", "outputPath": out("flow")})
    document = load(result["outputPath"])
    check(result["layout"] == "flow", "版式记成 flow", result["layout"])
    check(result["fieldCount"] == 8, "字段块数量和 form 版一致", result["fieldCount"])
    check(len(document.tables) == 2, "只剩顶部信息表 + 板书框", len(document.tables))
    check(len(document.tables[0].rows) == 3, "flow 版抬头表也是 3 行",
          len(document.tables[0].rows))
    check(result["table"]["student"] == {"filled": 0, "blank": 0},
          "整段版没有活动格子，计数全 0（别让面板以为有留白）", result["table"]["student"])
    check(not [paragraph for paragraph in body_paragraphs(document)
               if paragraph.text.startswith("注：")], "整段版不印「学生活动」说明")
    one_by_two = [table for table in document.tables
                  if len(table.columns) == 1 and len(table.rows) == 2]
    check(one_by_two == [], "没有「字段名一行/内容一行」的字段表", len(one_by_two))
    bold_labels = [paragraph.text for paragraph in body_paragraphs(document)
                   if paragraph.runs and all(run.bold for run in paragraph.runs)]
    check("教学目标" in bold_labels, "字段名是加粗段落", bold_labels[:6])
    check("教学准备" in bold_labels, "后面的字段也是加粗段落", bold_labels[:8])
    check(body_paragraphs(document)[0].text.endswith("教案"), "标题仍在最前面",
          body_paragraphs(document)[0].text)
    board = find_table(document, "板书设计")
    check(board is not None and len(board.rows) == 1 and len(board.columns) == 1,
          "flow 版板书仍是 1×1 格子", board)

    # [6] 边界：空 text / 只有标题 / 5000 字长文 / 坏版式 / 覆盖
    print("[6] 边界与错误")
    expect_error(lambda: build_plan({"text": ""}), "text-required", "空 text 报 text-required")
    expect_error(lambda: build_plan({"text": "   \n\t  "}), "text-required", "全空白也报错")
    expect_error(lambda: parse_plan(None), "text-not-string", "parse_plan(None) 报错")
    expect_error(lambda: build_plan({"text": SAMPLE, "layout": "卡片"}),
                 "unknown-layout", "不认识的版式报错")

    result = build_plan({"text": "课题：春", "outputPath": out("only-title")})
    document = load(result["outputPath"])
    check(result["title"] == "春", "只有课题时标题取课题", result["title"])
    check(result["fieldCount"] == 0, "没有正文字段", result["fieldCount"])
    check(len(document.tables) == 2, "仍然给出信息表 + 板书框", len(document.tables))
    codes = [warning["code"] for warning in result["warnings"]]
    check("no-procedure" in codes and "empty-plan" in codes,
          "提示缺教学过程、缺正文字段", codes)

    long_text = ("课题：长文测试\n一、教学目标\n" + "理解光合作用的过程与意义。" * 260
                 + "\n五、教学过程\n（一）导入\n" + "教师：讲解光合作用。" * 240)
    check(len(long_text) > 5000, "长文样本超过 5000 字", len(long_text))
    result = build_plan({"text": long_text, "outputPath": out("long")})
    document = load(result["outputPath"])
    check(os.path.getsize(result["outputPath"]) > 5000, "长文成品有内容", result["bytes"])
    check(result["sections"] == ["导入"], "长文里的环节仍认得出", result["sections"])
    check(len(document.tables) >= 4, "长文表格没崩", len(document.tables))

    target = out("overwrite")
    build_plan({"text": SAMPLE, "outputPath": target})
    expect_error(lambda: build_plan({"text": SAMPLE, "outputPath": target}),
                 "output-exists", "同名文件不覆盖")
    again = build_plan({"text": SAMPLE, "outputPath": target, "overwrite": True})
    check(os.path.isfile(again["outputPath"]), "overwrite=True 才覆盖", again["outputPath"])

    # [7] baseTemplate：页眉页脚保留、模板正文清掉
    print("[7] 套模板")
    template = os.path.join(work, "学校模板.docx")
    source = docx.Document()
    source.sections[0].header.paragraphs[0].text = "XX中学电子教案（2025—2026 学年）"
    source.sections[0].footer.paragraphs[0].text = "模板页脚：教务处存档"
    source.add_paragraph("模板里的示例正文，生成时应被清掉")
    source.save(template)
    result = build_plan({"text": SAMPLE, "baseTemplate": template, "outputPath": out("tpl")})
    document = load(result["outputPath"])
    check(document.sections[0].header.paragraphs[0].text == "XX中学电子教案（2025—2026 学年）",
          "页眉保留", document.sections[0].header.paragraphs[0].text)
    check("模板页脚" in footer_text(document), "模板页脚保留（不写默认页脚）",
          footer_text(document))
    check(not any("模板里的示例正文" in paragraph.text for paragraph in document.paragraphs),
          "模板正文被清掉")
    check(len(document.tables) == result["fieldCount"] + 2, "表格还是按字段排",
          len(document.tables))
    expect_error(lambda: build_plan({"text": SAMPLE,
                                     "baseTemplate": os.path.join(work, "没这个文件.docx")}),
                 "file-not-found", "模板不存在报 file-not-found")
    broken = os.path.join(work, "坏模板.docx")
    with open(broken, "w", encoding="utf-8") as handle:
        handle.write("这不是一个 docx")
    expect_error(lambda: build_plan({"text": SAMPLE, "baseTemplate": broken}),
                 "template-broken", "坏模板报 template-broken")

    # [8] 标题缺省与返回结构
    print("[8] 标题缺省与返回字段")
    result = build_plan({"text": SAMPLE, "outputPath": out("no-title")})
    check(result["title"] == "What does he look like?", "没传 title 时用课题", result["title"])
    check(set(result) == {"outputPath", "title", "layout", "fieldCount", "sections", "warnings",
                          "bytes", "table"}, "返回字段齐全且不多余", sorted(result))
    check(set(result["table"]) == {"rows", "teacher", "student", "time"},
          "table 统计含三列填写情况", sorted(result["table"]))
    check(all(set(result["table"][key]) == {"filled", "blank"}
              for key in ("teacher", "student", "time")), "每列都有 filled/blank", result["table"])
    check(result["sections"] == ["导入", "新授", "操练", "小结"], "返回环节清单",
          result["sections"])
    result = build_plan({"text": "一、教学目标\n会用一般过去时。"})
    check(result["title"] == "教案", "连课题都没有就落到「教案」", result["title"])
    check(result["outputPath"].startswith(os.path.join(work, "教师助手", "教案")),
          "默认输出到 教师助手/教案/", result["outputPath"])
    check(os.path.isfile(result["outputPath"]), "默认路径真的落了盘", result["outputPath"])
    document = load(result["outputPath"])
    check(body_paragraphs(document)[0].text == "教案", "没课题时标题不套书名号",
          body_paragraphs(document)[0].text)
    result = build_plan({"text": SAMPLE, "subject": "英语", "outputPath": out("subject")})
    document = load(result["outputPath"])
    check(body_paragraphs(document)[0].text.startswith("英语 "), "科目加在标题前面",
          body_paragraphs(document)[0].text)

    # [9] 教学过程四格路由 + 认不出环节时的合并
    print("[9] 教学过程的四格")
    result = build_plan({"text": ROUTING, "outputPath": out("routing")})
    document = load(result["outputPath"])
    procedure = find_table(document, "环节")
    check(len(procedure.rows) == 3, "2 个环节 + 表头", len(procedure.rows))
    check(procedure.cell(1, 0).text == "热身" and procedure.cell(2, 0).text == "新授",
          "环节名去掉编号与时间", (procedure.cell(1, 0).text, procedure.cell(2, 0).text))
    check(procedure.cell(1, 3).text == "3 分钟", "环节标题里的时间进时间列",
          procedure.cell(1, 3).text)
    check(procedure.cell(2, 3).text == "10 分钟", "「时间：」行进时间列",
          procedure.cell(2, 3).text)
    check(procedure.cell(1, 1).text == "播放歌曲。", "教师：前缀被去掉", procedure.cell(1, 1).text)
    check(procedure.cell(2, 2).text == "做笔记。", "学生：进学生活动列", procedure.cell(2, 2).text)

    result = build_plan({"text": "一、教学过程\n先复习单词，再讲句型，最后做练习。",
                         "outputPath": out("merged")})
    document = load(result["outputPath"])
    check("no-sections" in [warning["code"] for warning in result["warnings"]],
          "认不出环节时给提示", result["warnings"])
    procedure = find_table(document, "环节")
    check(len(procedure.rows) == 2, "整段排成一行", len(procedure.rows))
    check("先复习单词" in procedure.cell(1, 0).text, "整段话在第一格",
          procedure.cell(1, 0).text[:20])
    span = procedure.cell(1, 0)._tc.tcPr.find(qn("w:gridSpan"))
    check(span is not None and span.get(qn("w:val")) == "4",
          "第一格横跨 4 列，内容不被挤成竖排", span is not None and span.get(qn("w:val")))
    check(len(procedure.columns) == 4, "合并后仍是 4 列网格", len(procedure.columns))
    # 合并成一行时后三列本来就是留给老师填的：表头要点明「老师自填」，
    # 计数也要如实报 1 行留白（以前报全 0，面板与回执看着像没生成）
    check(procedure.cell(0, 2).text == "学生活动（老师自填）",
          "整段排一行时表头也标「老师自填」（后三列是留白）", procedure.cell(0, 2).text)
    check(result["table"]["student"] == {"filled": 0, "blank": 1},
          "整段排一行时活动计数报 1 行留白", result["table"]["student"])
    check(result["table"]["teacher"] == {"filled": 1, "blank": 0},
          "整段排一行时教师活动算填了", result["table"]["teacher"])

    # [10] 重复字段合并、空字段提示、面板字段表
    print("[10] 重复字段与面板字段表")
    parsed = parse_plan(DUP)
    codes = [warning["code"] for warning in parsed["warnings"]]
    check("duplicate-field" in codes, "重复的块给提示", codes)
    check(parsed["fields"]["objectives"].count("目标") >= 2, "两段目标合并成一块",
          parsed["fields"]["objectives"])
    check("第二条目标" in parsed["fields"]["objectives"], "后出现的内容没丢",
          parsed["fields"]["objectives"])
    check(parsed["fields"].get("reflection") == "", "只有标题没内容的字段记成空串",
          parsed["fields"].get("reflection"))
    result = build_plan({"text": DUP, "outputPath": out("dup")})
    check("empty-field" in [warning["code"] for warning in result["warnings"]],
          "空字段提醒老师补一句", result["warnings"])
    check(result["fieldCount"] == 2, "空字段不印表格", result["fieldCount"])

    items = plan_fields()
    check(len(items) == 14, "面板字段表 14 个字段", len(items))
    check(all(set(item) == {"key", "label", "aliases"} for item in items),
          "每个字段都是 key/label/aliases 三样", items[0])
    by_key = {item["key"]: item for item in items}
    check(by_key["objectives"]["label"] == "教学目标", "教学目标的中文名",
          by_key["objectives"]["label"])
    check("学习目标" in by_key["objectives"]["aliases"], "教学目标认「学习目标」别名",
          by_key["objectives"]["aliases"])
    check("重点" in by_key["keyPoints"]["aliases"], "教学重点认「重点」别名",
          by_key["keyPoints"]["aliases"])
    check(items[0]["key"] == "title", "第一个字段是课题", items[0]["key"])

    # [11] 学生活动整列空白：不瞎编，但必须一眼看出「留给你填」
    print("[11] 学生活动整列空白")
    result = build_plan({"text": BLANK_STUDENT, "outputPath": out("blank-student")})
    document = load(result["outputPath"])
    procedure = find_table(document, "环节")
    check(procedure.cell(0, 2).text == "学生活动（老师自填）",
          "整列空白时表头写明「老师自填」", procedure.cell(0, 2).text)
    check(all(not procedure.cell(index, 2).text.strip() for index in (1, 2, 3)),
          "格子仍然空着（不瞎编学生的反应）",
          [procedure.cell(index, 2).text for index in (1, 2, 3)])
    codes = [warning["code"] for warning in result["warnings"]]
    check("student-column-blank" in codes, "整列空白给一条 student-column-blank", codes)
    check(result["table"]["student"] == {"filled": 0, "blank": 3},
          "返回值说清几格留白", result["table"]["student"])
    check(any(paragraph.text.startswith("注：") for paragraph in body_paragraphs(document)),
          "表下有一行说明小字")
    check(all("自由回答" not in procedure.cell(index, 2).text for index in (1, 2, 3)),
          "没有塞进任何编造的学生活动")

    # 只空个别格子时不能报「整列空白」——条条报警，老师连真问题都不看了
    result = build_plan({"text": SAMPLE, "outputPath": out("partial-student")})
    check("student-column-blank" not in [warning["code"] for warning in result["warnings"]],
          "只空几格时不发「整列空白」警告")
    document = load(result["outputPath"])
    check(find_table(document, "环节").cell(0, 2).text == "学生活动",
          "只空几格时表头保持原样", find_table(document, "环节").cell(0, 2).text)

    # [12] 长课题：抬头表任何一格都不许折行（老师第一眼看的就是这块）
    print("[12] 长课题也不折行")
    long_title = "人教版七年级下册 Unit 3 Section A 2d 听说课教学设计"
    result = build_plan({"text": "课题：" + long_title + "\n课型：新授课\n课时：1 课时\n"
                                 "班级：七年级（3）班\n一、教学目标\n能听懂并复述对话。",
                         "outputPath": out("long-title")})
    document = load(result["outputPath"])
    header = document.tables[0]
    check(header.cell(0, 1).text == long_title, "长课题整串进格子", header.cell(0, 1).text)
    check(header.cell(1, 3).text == "1 课时", "长课题时课时仍在同一行", header.cell(1, 3).text)
    folded = [cell.text for cell in (header.cell(0, 1), header.cell(1, 3), header.cell(2, 1),
                                     header.cell(2, 3)) if "\n" in cell.text]
    check(folded == [], "抬头表格子里没有换行", folded)
    size = header.cell(0, 1).paragraphs[0].runs[0].font.size
    check(size is not None and 8 <= size.pt <= 10.5, "长课题按格宽缩了字号", size and size.pt)

    # [13] 老师最自然的写法：「导入（5 分钟）」这种裸环节名（P1-7）
    print("[13] 裸环节名与括号时长")
    parsed = parse_plan(BARE_SECTIONS)
    check([section["title"] for section in parsed["sections"]] == ["导入", "新授", "操练", "小结"],
          "「导入（5 分钟）/新授(10min)/操练 12 分钟/小结（3 分钟）」四条全中",
          [section["title"] for section in parsed["sections"]])
    check(parsed["warnings"] == [], "裸环节名不再刷提示", parsed["warnings"])
    check(parsed["sections"][0]["lines"] == ["教师：出示两张明星照片。", "学生：自由回答。"],
          "环节正文跟着环节走", parsed["sections"][0]["lines"])
    result = build_plan({"text": BARE_SECTIONS, "outputPath": out("bare-sections")})
    document = load(result["outputPath"])
    procedure = find_table(document, "环节")
    check(len(procedure.rows) == 5, "4 个环节 + 表头 5 行", len(procedure.rows))
    check([procedure.cell(index, 0).text for index in range(1, 5)] == ["导入", "新授", "操练", "小结"],
          "环节列是四个名字", [procedure.cell(index, 0).text for index in range(1, 5)])
    check([procedure.cell(index, 3).text for index in range(1, 5)]
          == ["5 分钟", "10min", "12 分钟", "3 分钟"],
          "四种时长写法都解析进「时间」列（括号 / 空格 / min）",
          [procedure.cell(index, 3).text for index in range(1, 5)])
    check(result["table"]["time"] == {"filled": 4, "blank": 0}, "时间列 4 格都填上",
          result["table"]["time"])
    check(result["sections"] == ["导入", "新授", "操练", "小结"], "返回值里的环节清单",
          result["sections"])
    check("no-sections" not in [warning["code"] for warning in result["warnings"]],
          "认出环节后不再发 no-sections", result["warnings"])
    check("明星照片" in procedure.cell(1, 1).text and procedure.cell(1, 2).text == "自由回答。",
          "教师活动 / 学生活动照旧各进各的格子",
          (procedure.cell(1, 1).text, procedure.cell(1, 2).text))

    # 反例：正文句子不能被切成环节（宁可少猜，也不要把表切碎）
    parsed = parse_plan("一、教学过程\n导入\n教师：出示图片。\n复习单词三遍\n学生：齐读。\n")
    check([section["title"] for section in parsed["sections"]] == ["导入"],
          "「复习单词三遍」这种正文没被当成新环节",
          [section["title"] for section in parsed["sections"]])

    # 整段还是认不出时，提示要告诉老师正确写法
    result = build_plan({"text": "一、教学过程\n先复习单词，再讲句型，最后做练习。",
                         "outputPath": out("no-sections-hint")})
    messages = [warning["message"] for warning in result["warnings"]
                if warning["code"] == "no-sections"]
    check(bool(messages) and "导入（5 分钟）" in messages[0] and "（一）导入" in messages[0],
          "认不出环节时提示给出正确写法示例", messages[:1])

    # [17] 超长环节：一行比一页还高时，Word/LO 只能从中间切断，续页第一列空白、
    # 表头也不再现（0.18 审查 L-01）。现在按字数拆成「（续）」行 + 表头行标跨页重复。
    print("[17] 超长环节拆成「（续）」行，表头跨页重复")
    long_lines = "\n".join("教师：讲解第 %d 个要点，示例句 There is a museum." % index
                           for index in range(1, 30))
    result = build_plan({"text": "课题：长内容\n一、教学过程\n（一）导入（5 分钟）\n" + long_lines,
                         "outputPath": out("long-section")})
    document = load(result["outputPath"])
    procedure = find_table(document, "环节")
    labels = [procedure.cell(index, 0).text for index in range(1, len(procedure.rows))]
    check(any(label.endswith("（续）") for label in labels),
          "超长环节拆出「（续）」行，续页不再只有空白的环节列", labels[:6])
    check(len(procedure.rows) > len(result["sections"]) + 1,
          "统计按环节算（table.rows），表格行数按拆出来的行算",
          (result["table"]["rows"], result["sections"], len(procedure.rows)))
    header_trpr = procedure.rows[0]._tr.find(qn("w:trPr"))
    check(header_trpr is not None and header_trpr.find(qn("w:tblHeader")) is not None,
          "表头行标了「跨页重复」（续页有「环节/教师活动/学生活动/时间」）")
    children = [child.tag.rsplit("}", 1)[-1] for child in list(document.element.body)]
    pairs = [children[index] + children[index + 1] for index in range(len(children) - 1)]
    check("tbltbl" not in pairs,
          "相邻两张表之间垫了空段落（紧挨着写会被当成一张表，表头重复就失效）", pairs[:8])

    print("")
    if FAILS:
        print("失败 " + str(len(FAILS)) + " 项：")
        for label in FAILS:
            print(" - " + label)
        sys.exit(1)
    print("全部通过 ✅")


if __name__ == "__main__":
    main()