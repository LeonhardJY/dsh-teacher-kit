#!/usr/bin/env python
"""模板反推自检：拿"老师自己的旧卷子"当输入，看能不能读出规则并用它造新卷。"""

import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

import docx  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402

from teacher_kit.commands import TABLE, docx_new, inspect, template_diff, template_make, template_model  # noqa: E402
from teacher_kit.errors import EngineError  # noqa: E402

FAILS = []


def check(condition, label, detail=""):
    if condition:
        print("  ✓ " + label)
    else:
        FAILS.append(label)
        print("  ✗ " + label + (" — " + str(detail) if detail else ""))


def styled(paragraph, text, size, bold=False, align=None):
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    if align is not None:
        paragraph.alignment = align
    return run


def make_exam(path, left_margin=27.0):
    document = docx.Document()
    section = document.sections[0]
    section.orientation = 1  # landscape
    section.page_width = Mm(368)
    section.page_height = Mm(260)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(24)
    section.left_margin = Mm(left_margin)
    section.right_margin = Mm(left_margin)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = section._sectPr.makeelement(qn("w:cols"), {})
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "2")

    section.header.paragraphs[0].text = "XX 中学 2025—2026 学年第二学期 七年级英语"
    footer = section.footer.paragraphs[0]
    footer.text = "第 X 页 共 Y 页　　使用班级：______　　出卷老师：______"
    # 真页码域：检查器要能认出"第X页 共Y页"是域而不是死文本
    field_run = footer.add_run()
    element = field_run._element
    begin = element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = element.makeelement(qn("w:instrText"), {})
    instr.text = " PAGE "
    end = element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    element.append(begin)
    element.append(instr)
    element.append(end)

    styled(document.add_paragraph(), "七年级英语期中测试卷", 16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    styled(document.add_paragraph(), "（满分 100 分，时间 90 分钟）", 11, align=WD_ALIGN_PARAGRAPH.CENTER)
    styled(document.add_paragraph(), "一、单项选择（20 分）", 14, bold=True)
    styled(document.add_paragraph(), "1. — How do you get to school? — ______", 12)
    styled(document.add_paragraph(), "A. By bike", 12)
    styled(document.add_paragraph(), "B. On foot", 12)
    styled(document.add_paragraph(), "（1）He goes to school ______ bus.", 12)
    styled(document.add_paragraph(), "二、完形填空（15 分）", 14, bold=True)
    styled(document.add_paragraph(), "21. My father ______ to work every day.", 12)
    document.save(path)


def main():
    work = tempfile.mkdtemp(prefix="teacher-kit-tpl-")
    exam_a = os.path.join(work, "上学期期末卷.docx")
    exam_b = os.path.join(work, "上学期月考卷.docx")
    make_exam(exam_a)
    make_exam(exam_b, left_margin=30.0)

    print("[1] 单份：读出纸张与版心")
    profile = template_model({"paths": [exam_a]})
    paper = profile["paper"]
    check(paper["size"] == "8K", "认出 8 开", paper["size"])
    check(paper["widthMm"] == 368 and paper["heightMm"] == 260, "368×260mm", paper)
    check(paper["orientation"] == "landscape", "横向", paper["orientation"])
    check(paper["columns"] == 2, "两栏", paper["columns"])
    check(abs(paper["marginMm"]["left"] - 27) < 1, "左边距 27mm", paper["marginMm"])
    check(profile["confidence"] == "medium", "单份无冲突 → medium（样本只有一份，不敢说 high）", profile["confidence"])
    check(profile["conflicts"] == [], "没有冲突", profile["conflicts"])

    print("[2] 页眉页脚：抽成占位符")
    header = " / ".join(profile["header"]["template"])
    check("XX 中学" in header, "页眉有校名", header)
    check("{{schoolYear}}" in header, "学年变成 {{schoolYear}}", header)
    footer = profile["footer"]["template"]
    check("{{page}}" in footer and "{{pages}}" in footer, "页码变成 {{page}}/{{pages}}", footer)
    check("使用班级：{{blank}}" in footer, "班级留白变成 {{blank}}", footer)
    check(profile["footer"]["hasPageField"] is True, "认出真页码域", profile["footer"])

    print("[3] 字体字号与题号层级")
    roles = profile["roles"]
    check(roles["title"]["font"] == "宋体", "标题字体 宋体", roles["title"])
    check(roles["title"]["sizePt"] == 16, "标题字号 16pt", roles["title"])
    check(roles["title"]["bold"] is True, "标题加粗", roles["title"])
    check(roles["section"]["sizePt"] == 14, "大题 14pt", roles["section"])
    check(roles["question"]["sizePt"] == 12, "小题 12pt", roles["question"])
    check(roles["option"]["count"] >= 2, "认出选项行", roles.get("option"))
    numbering = profile["numbering"]
    check(numbering["sections"] == "一、二", "大题层级 一、二", numbering["sections"])
    check(numbering["questions"] == "1、21", "小题编号 1、21", numbering["questions"])
    check(numbering["subquestions"] == "1", "括号小题 1", numbering["subquestions"])

    print("[4] 反推结果能直接用来造新卷（闭环）")
    page = profile["buildPage"]
    check(page["size"] == "8K" and page["columns"] == 2 and page["orientation"] == "landscape", "buildPage 可直传", page)
    new_exam = os.path.join(work, "用反推规则造的新卷.docx")
    built = docx_new({
        "outputPath": new_exam,
        "page": page,
        "blocks": [
            {"type": "paragraph", "align": "center", "runs": [{"text": "七年级英语第二学期期末测试卷", "bold": True, "sizePt": 16}]},
            {"type": "paragraph", "text": "一、单项选择（20 分）"},
            {"type": "paragraph", "text": "1. — Where is the library? — ______"},
        ],
    })
    check(built["page"]["widthMm"] == 368 and built["page"]["columns"] == 2, "新卷是 8 开两栏", built["page"])
    facts = inspect({"path": new_exam})
    check(abs(facts["primary"]["pageWidthMm"] - 368) < 1, "读回 368mm", facts["primary"]["pageWidthMm"])
    check(facts["primary"]["columns"] == 2, "读回两栏", facts["primary"]["columns"])
    check(abs(facts["primary"]["marginLeftMm"] - 27) < 1, "页边距跟着反推走 27mm", facts["primary"]["marginLeftMm"])

    print("[5] 两份不一致：列成冲突，不替老师拍板")
    merged = template_model({"paths": [exam_a, exam_b]})
    same_twice = template_model({"paths": [exam_a, exam_a]})
    check(same_twice["confidence"] == "high", "两份一致 → high", same_twice["confidence"])
    conflicts = merged["conflicts"]
    check(len(conflicts) >= 1, "有冲突被列出", conflicts)
    left_conflict = next((c for c in conflicts if c["field"] == "paper.margin.left"), None)
    check(left_conflict is not None, "冲突点是左边距", [c["field"] for c in conflicts])
    if left_conflict:
        values = sorted(item["value"] for item in left_conflict["options"])
        check(values == [27.0, 30.0], "两个取值都保留下来", left_conflict["options"])
        check(len(left_conflict["options"][0]["sources"]) == 1, "每个取值都指出出处", left_conflict["options"])
    check(merged["confidence"] == "low", "有冲突 → low", merged["confidence"])
    check(len(merged["sources"]) == 2, "两份来源都记下", merged["sources"])

    print("[6] 两份对比（这学期 vs 上学期）")
    profile_b = template_model({"paths": [exam_b]})
    diff = template_diff({"before": profile, "after": profile_b})
    fields = [c["field"] for c in diff["changes"]]
    check(diff["changed"] >= 1, "有差异", diff["changed"])
    check("paper.marginMm.left" in fields, "差异定位到左边距", fields[:8])
    check(all("before" in c and "after" in c for c in diff["changes"]), "每处差异都有前后值", diff["changes"][:2])

    print("[7] 错误形状")
    try:
        template_model({"paths": []})
        check(False, "空输入 → paths-required")
    except EngineError as error:
        check(error.code == "paths-required", "空输入 → paths-required", error.code)
    try:
        template_model({"paths": [os.path.join(work, "没有这份.docx")]})
        check(False, "文件不存在 → file-not-found")
    except EngineError as error:
        check(error.code == "file-not-found", "文件不存在 → file-not-found", error.code)
    txt = os.path.join(work, "x.txt")
    open(txt, "w", encoding="utf-8").write("hi")
    try:
        template_model({"paths": [txt]})
        check(False, "txt → unsupported-source")
    except EngineError as error:
        check(error.code == "unsupported-source", "txt → unsupported-source", error.code)

    print("[8] 一学期后的闭环：旧卷子→空白模板→新卷子")
    made = template_make({"source": exam_a})
    check(os.path.isfile(made["outputPath"]), "生成模板文件", made["outputPath"])
    template_doc = docx.Document(made["outputPath"])
    header_text = " / ".join(p.text for p in template_doc.sections[0].header.paragraphs)
    footer_text = " / ".join(p.text for p in template_doc.sections[0].footer.paragraphs)
    check("2025—2026" not in header_text, "模板里不再有旧学年", header_text)
    check("{{schoolYear}}" in header_text, "模板页眉留了学年占位符", header_text)
    check("{{page}}" in footer_text and "{{pages}}" in footer_text, "模板页脚留了页码占位符", footer_text)
    check("{{className}}" in footer_text and "{{teacherName}}" in footer_text,
          "班级/出卷人留了专用占位符（面板填了就印上）", footer_text)
    body_text = "\n".join(p.text for p in template_doc.paragraphs)
    check("{{examTitle}}" in body_text, "正文留标题占位符", body_text)
    check("How do you get to school" not in body_text, "旧题目已清空（不给老师留垃圾）", body_text[:60])
    check(template_doc.sections[0].page_width.mm > 300, "模板还是 8 开横向", template_doc.sections[0].page_width.mm)

    filled_exam = os.path.join(work, "下学期期中卷.docx")
    built2 = docx_new({
        "outputPath": filled_exam,
        "baseTemplate": made["outputPath"],
        "values": {"schoolYear": "2026—2027 学年", "examTitle": "七年级英语第二学期期中测试卷",
                   "totalScore": "100", "minutes": "90", "blank": "七(3)班",
                   "className": "七(3)班", "teacherName": "王老师"},
        "page": made["buildPage"],
        "blocks": [
            {"type": "paragraph", "align": "center", "runs": [{"text": "一、单项选择（20 分）", "bold": True, "sizePt": 14}]},
            {"type": "paragraph", "text": "1. — What time do you go to school? — ______"},
        ],
    })
    check(built2["placeholdersFilled"] >= 1, "造卷时填了占位符", built2["placeholdersFilled"])
    check(not [w for w in built2["warnings"] if w["code"] == "unfilled-placeholder"], "没有漏填的占位符", built2["warnings"])
    final = docx.Document(filled_exam)
    final_header = " / ".join(p.text for p in final.sections[0].header.paragraphs)
    final_footer = " / ".join(p.text for p in final.sections[0].footer.paragraphs)
    check("2026—2027" in final_header, "新卷页眉是新学年", final_header)
    check("2025—2026" not in final_header, "旧学年没跟过来", final_header)
    check("七(3)班" in final_footer and "王老师" in final_footer, "班级和老师都填进去了", final_footer)
    footer_xml = final.sections[0].footer.paragraphs[0]._p.xml
    check("fldSimple" in footer_xml and "NUMPAGES" in footer_xml, "页码是真域（打印才对）", footer_xml[:200])
    final_body = "\n".join(p.text for p in final.paragraphs)
    check("What time do you go to school" in final_body, "新题目写进去了", final_body[:80])

    print("[9] 多节试卷反推模板：8 开两栏不能被末尾那张 A4 答题卡抢走")
    # 真机缺陷（persona P0-1）：8 开双栏的 sectPr 藏在段落的 pPr 里，清正文时被一起删，
    # 模板退回 body 末尾那节（A4 答题卡）→ 套模板出卷变 A4 单栏
    sectioned = os.path.join(work, "多节源卷.docx")
    body_lines = ["七年级英语期末测试", "", "（满分 40 分，时间 60 分钟）", ""]
    body_lines += ["一、单项选择（共 4 小题，每小题 2 分，共 8 分）"]
    for index in range(1, 5):
        body_lines += [str(index) + ". She ___ to school every day.", "A. go", "B. goes", "C. going", "D. went", "【答案】B"]
    body_lines += ["二、书面表达（32 分）", "5. 写一段话介绍你的朋友。", "【答案】略"]
    TABLE["exam:build"]({
        "text": "\n".join(body_lines), "title": "多节源卷", "outputPath": sectioned,
        "page": {"size": "8K", "orientation": "landscape", "columns": 2},
        "options": {"answerCard": "card", "answerSheet": "end"},
    })
    sectioned_doc = docx.Document(sectioned)
    check(len(sectioned_doc.sections) >= 3, "源卷本来是多节的（标题/正文/答题卡）", len(sectioned_doc.sections))

    made_multi = template_make({"source": sectioned, "outputPath": os.path.join(work, "多节模板.docx"), "overwrite": True})
    paper = made_multi["paper"]
    check(paper.get("size") == "8K", "模板纸型是 8 开（不是末尾那张 A4）", paper)
    check(abs(float(paper.get("widthMm") or 0) - 368) <= 2, "模板纸宽还是 368mm", paper.get("widthMm"))
    check(int(paper.get("columns") or 1) == 2, "模板还是两栏", paper.get("columns"))
    check(made_multi["buildPage"].get("columns") == 2, "buildPage 跟着实测纸型", made_multi["buildPage"])
    multi_doc = docx.Document(made_multi["outputPath"])
    check(len(multi_doc.sections) == 1, "模板正文压成一节（不留分节符）", len(multi_doc.sections))
    multi_cols = multi_doc.sections[0]._sectPr.find(qn("w:cols"))
    check(multi_cols is not None and multi_cols.get(qn("w:num")) == "2", "读回来还是两栏",
          multi_cols.get(qn("w:num")) if multi_cols is not None else None)

    applied = os.path.join(work, "套多节模板.docx")
    applied_result = TABLE["exam:build"]({
        "text": "\n".join(body_lines), "title": "套模板复用小测", "outputPath": applied,
        "baseTemplate": made_multi["outputPath"], "options": {"minutes": 60},
    })
    applied_doc = docx.Document(applied)
    applied_cols = [section._sectPr.find(qn("w:cols")) for section in applied_doc.sections]
    widths = [round(section.page_width.mm) for section in applied_doc.sections]
    check(all(width == 368 for width in widths), "套模板后每节都是 8 开宽（没掉回 A4）", widths)
    check(any(cols is not None and cols.get(qn("w:num")) == "2" for cols in applied_cols),
          "套模板后正文还是两栏", [cols.get(qn("w:num")) if cols is not None else None for cols in applied_cols])

    # 真机第二轮缺陷：多节压成一节后"链接到上一节"的页眉页脚断链 —— 模板里没有页脚，
    # 套模板出卷 preflight 报「没有页脚」（95 分），老师印出来翻页就掉号
    tpl_sect = multi_doc.sections[0]
    tpl_footer_refs = tpl_sect._sectPr.findall(qn("w:footerReference"))
    check(len(tpl_footer_refs) >= 1, "模板自己带着页脚引用（不是链接到上一节）", len(tpl_footer_refs))
    tpl_footer_text = " ".join(p.text for p in tpl_sect.footer.paragraphs)
    check("{{page}}" in tpl_footer_text or "{{blank}}" in tpl_footer_text,
          "模板页脚里的年份/页码换成了占位符", tpl_footer_text[:60])
    applied_footer = " ".join(p.text for p in applied_doc.sections[0].footer.paragraphs)
    check("第" in applied_footer and "页" in applied_footer, "套模板出卷后正文有页脚", applied_footer[:60])
    applied_codes = [item.get("code") for item in (applied_result.get("preflight") or {}).get("findings", [])]
    check("footer-missing" not in applied_codes, "套模板出卷不再报「没有页脚」", applied_codes)

    # 页脚里的"使用班级/出卷老师"要能吃到面板填的值：模板里必须留 {{className}} 这种占位符
    check("{{className}}" in tpl_footer_text and "{{teacherName}}" in tpl_footer_text,
          "模板页脚留了班级/老师的专用占位符", tpl_footer_text[:80])
    named = os.path.join(work, "套模板带班级.docx")
    TABLE["exam:build"]({
        "text": "\n".join(body_lines), "title": "套模板带班级", "outputPath": named,
        "baseTemplate": made_multi["outputPath"],
        "values": {"className": "七(3)班", "teacherName": "王老师"},
    })
    named_footer = " ".join(p.text for p in docx.Document(named).sections[0].footer.paragraphs)
    check("七(3)班" in named_footer and "王老师" in named_footer,
          "套模板出卷时班级/老师印进页脚", named_footer[:80])

    print("")
    if FAILS:
        print(str(len(FAILS)) + " 项失败 ❌")
        for item in FAILS:
            print("  - " + item)
        return 1
    print("全部通过 ✅")
    return 0


if __name__ == "__main__":
    sys.exit(main())
