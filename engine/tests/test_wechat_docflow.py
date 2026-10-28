#!/usr/bin/env python
"""公众号推文 / 公文流转（v0.15 新增）。跑 python test_wechat_docflow.py。"""

import os
import sys
import tempfile

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)

from teacher_kit.commands import TABLE  # noqa: E402
from teacher_kit.docflow import docflow_build  # noqa: E402
from teacher_kit.errors import EngineError  # noqa: E402
from teacher_kit.wechat import wechat_check, wechat_render  # noqa: E402

FAILS = []


def check(condition, label, detail=''):
    if condition:
        print('  ✓ ' + label)
    else:
        FAILS.append(label)
        print('  ✗ ' + label + (' — ' + str(detail) if detail else ''))


def expect_error(fn, code, label):
    try:
        fn()
    except EngineError as error:
        check(error.code == code, label, '拿到的是 ' + error.code)
        return
    check(False, label, '没有抛错')


def codes(result):
    return [item.get('code') for item in (result.get('warnings') or [])]


POST = {
    'title': '我校开展防溺水主题教育',
    'digest': '七年级全体学生参加，现场演练救生衣穿法。',
    'blocks': [
        {'type': 'lead', 'text': '为了让每个孩子知道落水后怎么办，学校把这一次的安全课搬到了河边。'},
        {'type': 'heading', 'text': '一、现场演练'},
        {'type': 'text', 'text': '七月的午后，七年级四个班分批到河堤，由体育老师示范救生衣的穿法和抛绳的要领。'},
        {'type': 'text', 'text': '每班抽十名学生实操，其余学生在安全线内观看并记笔记。'},
        {'type': 'heading', 'text': '二、下一步安排'},
        {'type': 'list', 'items': ['各班交一份家庭防溺水承诺', '暑假前再组织一次演练']},
    ],
    'author': '临河区第三中学',
    'signature': '临河区第三中学德育处',
}


def docflow_payload(workdir):
    return {
        'title': '关于开展 2026 年中小学生防溺水专项行动的通知',
        'unit': '临河区第三中学',
        'from': '临河区教育局',
        'docNo': '临教发〔2026〕18 号',
        'receivedAt': '2026-07-02',
        'copies': '3',
        'secret': '内部',
        'urgency': '急件',
        'dept': '德育处',
        'deadline': '2026-07-10',
        'handler': '王主任',
        'registryNo': '2026-收-042',
        'summary': '教育局要求各校在暑期前完成防溺水专题教育，并报送开展情况。',
        'items': ['各班开展一次专题班会', '组织一次现场演练', '7 月 10 日前报送总结'],
        'opinion': '拟请教务处牵头，德育处配合，7 月 8 日前完成专题教育和现场演练，7 月 10 日前将总结报教育局安全科。',
        'leaderNote': '同意，务必落实到每个班级。',
        'steps': [
            {'name': '登记', 'dept': '办公室', 'owner': '李老师', 'due': '07-02', 'note': '已登记'},
            {'name': '拟办', 'dept': '德育处', 'owner': '王主任', 'due': '07-03', 'note': '出方案'},
            {'name': '批示', 'dept': '校领导', 'owner': '张校长', 'due': '07-04', 'note': ''},
            {'name': '承办', 'dept': '各年级组', 'owner': '年级组长', 'due': '07-08', 'note': ''},
            {'name': '归档', 'dept': '办公室', 'owner': '李老师', 'due': '07-10', 'note': ''},
        ],
        'reply': {
            'to': '临河区教育局',
            'blocks': [
                '贵局《关于开展 2026 年中小学生防溺水专项行动的通知》（临教发〔2026〕18 号）收悉。',
                '我校于 7 月 3 日召开专题会议部署，7 月 6 日组织七年级全体学生开展现场演练，各班同步召开主题班会。',
                '现将开展情况报上，请审阅。',
            ],
            'signature': '临河区第三中学',
            'date': '2026年7月9日',
        },
        'outputPath': os.path.join(workdir, '公文办理单.docx'),
    }


def main():
    workdir = tempfile.mkdtemp(prefix='tk-wechat-')

    print('— 公众号：体检')
    dirty = wechat_check({
        'title': '我校在本次全市中小学生防溺水专项行动中取得全市第一的优异成绩并获评最佳组织奖',
        'digest': '',
        'text': '这是第一段。\n\n' + ('这一段很长，' * 60) + '\n\n最后一段。',
    })
    found = codes(dirty)
    check('title-too-long' in found, '标题过长会被指出来', found)
    check('digest-missing' in found, '没写摘要会被指出来', found)
    check('paragraph-too-long' in found, '整段太长（手机上一屏放不下）会被指出来', found)
    check('absolute-claim' in found, '「全市第一」「最佳」这类绝对化说法会被指出来', found)
    check('no-signoff' in found, '没落款会被指出来', found)
    check(dirty['stats']['paragraphs'] >= 3, '分段数统计正确', dirty['stats'])
    check(dirty['stats']['readSeconds'] >= 1, '给出预计阅读秒数', dirty['stats'])
    claim = [item for item in dirty['warnings'] if item.get('code') == 'absolute-claim']
    check(bool(claim) and claim[0].get('words'), '绝对化说法会点名是哪几个词',
          claim[0].get('words') if claim else None)

    print('— 公众号：干净稿子不该报警')
    clean = wechat_check(POST)
    check('title-too-long' not in codes(clean), '正常长度标题不报警', codes(clean))
    check('absolute-claim' not in codes(clean), '没有绝对化说法不报警', codes(clean))
    check(clean['title'] == POST['title'], '标题原样返回', clean['title'])

    print('— 公众号：出稿')
    result = wechat_render(dict(POST, outputPath=os.path.join(workdir, '推文')))
    check(os.path.isfile(result['htmlPath']), '出了可粘贴的 HTML', result.get('htmlPath'))
    check(os.path.isfile(result['outputPath']), '出了审稿用的 docx', result.get('outputPath'))
    html = open(result['htmlPath'], encoding='utf-8').read()
    check('<link' not in html and '<script' not in html, 'HTML 不带外链（公众号编辑器只吃内联）')
    check('#1f5fa9' in html, '默认用学校蓝主题', result.get('theme'))
    check('临河区第三中学德育处' in html, '落款进了成品', result.get('signature'))
    check('防溺水' in html, '正文内容进了成品')
    check(result['stats']['paragraphs'] >= 3 and result['stats']['bodyChars'] > 100,
          '出稿也带分段与字数统计', result['stats'])
    document = docx.Document(result['outputPath'])
    body = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    check(POST['title'] in body, 'docx 里有标题')
    check('出卷老师' not in body, 'docx 不印试卷那套页脚')

    warm = wechat_render(dict(POST, outputPath=os.path.join(workdir, '暖色'),
                              options={'theme': 'warm'}))
    warm_html = open(warm['htmlPath'], encoding='utf-8').read()
    check('#c2410c' in warm_html, '换主题生效（暖色）')

    expect_error(lambda: wechat_render({'title': '空稿'}),
                 'empty-post', '没有正文 → empty-post 而不是出一份空文件')

    print('— 公文：全套')
    full = docflow_build(docflow_payload(workdir))
    check(os.path.isfile(full['outputPath']), '出了办理单 docx', full.get('outputPath'))
    check(full['sections'] == ['收文登记表', '来文要点', '拟办意见', '领导批示', '承办流转单', '回复稿'],
          '六段齐全（带回复稿）', full['sections'])
    check(full['hasReply'] is True, 'hasReply 为真')
    check(full['steps'] == 5, '流转环节数对', full['steps'])
    check(len(full['warnings']) == 0, '要素齐全时不报警', codes(full))
    check(full['unit'] == '临河区第三中学', '本单位来自 payload', full.get('unit'))

    document = docx.Document(full['outputPath'])
    body = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    footer = '\n'.join(paragraph.text for paragraph in document.sections[0].footer.paragraphs)
    check('公文办理单' in body, '标题进正文')
    check('出卷老师' not in body and '使用班级' not in body, '公文上不印试卷页脚')
    check('临河区第三中学公文办理单' in footer, '页脚是本单位的公文办理单', footer.strip())
    section = document.sections[0]
    check(abs(section.page_width.cm - 21) < 0.2 and abs(section.page_height.cm - 29.7) < 0.2,
          'A4 竖向', (round(section.page_width.cm, 1), round(section.page_height.cm, 1)))
    check(abs(section.top_margin.mm - 37) < 0.6, '上边距按国标 37mm 左右',
          round(section.top_margin.mm, 1))
    check(len(document.tables) >= 2, '登记表 + 流转单两张表', len(document.tables))
    registry = document.tables[0]
    flat = '\n'.join(cell.text for row in registry.rows for cell in row.cells)
    check('临教发〔2026〕18 号' in flat, '登记表里有发文字号')
    check('2026-收-042' in flat, '登记表里有登记号')

    print('— 公文：缺项提醒')
    bare = docflow_build({'title': '关于报送暑假值班表的通知',
                          'outputPath': os.path.join(workdir, '空表.docx')})
    found = codes(bare)
    check('field-missing' in found, '要素缺失会提醒', found)
    check('opinion-empty' in found, '没写拟办意见会提醒', found)
    check('steps-empty' in found, '没排流转环节会提醒', found)
    check(bare['sections'] == ['收文登记表', '来文要点', '拟办意见', '领导批示', '承办流转单'],
          '没给回复稿就只有五段', bare['sections'])
    check(bare['hasReply'] is False, '没给回复稿 hasReply 为假')
    check(all(item.get('level') in ('info', 'warn', 'error') for item in bare['warnings']),
          '每条提醒都有级别')

    urgent = docflow_build({'title': '关于防汛的紧急通知', 'urgency': '特急',
                            'outputPath': os.path.join(workdir, '急件.docx')})
    urgent_codes = codes(urgent)
    check('urgent-no-deadline' in urgent_codes, '急件没写期限 → 明确报错级提醒', urgent_codes)
    urgent_item = [item for item in urgent['warnings'] if item.get('code') == 'urgent-no-deadline']
    check(bool(urgent_item) and urgent_item[0]['level'] == 'error', '这条是 error 级',
          urgent_item[0]['level'] if urgent_item else None)

    short_reply = docflow_build({
        'title': '关于报送暑假值班表的通知',
        'reply': {'blocks': ['收到。']},
        'outputPath': os.path.join(workdir, '短回复.docx'),
    })
    short_codes = codes(short_reply)
    check('reply-too-short' in short_codes, '回复稿太短会提醒', short_codes)
    check('reply-no-to' in short_codes, '回复稿缺主送机关会提醒', short_codes)

    expect_error(lambda: docflow_build({}), 'title-required', '没标题 → title-required')
    expect_error(lambda: docflow_build({'title': '关于 X 的通知', 'baseTemplate': os.path.join(workdir, '没有这份.docx')}),
                 'template-missing', '模板不存在 → template-missing')

    print('— 命令表接线')
    check(callable(TABLE.get('wechat:check')), 'wechat:check 在命令表里')
    check(callable(TABLE.get('wechat:render')), 'wechat:render 在命令表里')
    check(callable(TABLE.get('docflow:build')), 'docflow:build 在命令表里')

    print()
    if FAILS:
        print('失败 ' + str(len(FAILS)) + ' 项：')
        for label in FAILS:
            print(' - ' + label)
        sys.exit(1)
    print('全部通过 ✅')


if __name__ == '__main__':
    main()
